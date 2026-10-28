"""Preflight checks that mirror CET's own QC rules (cet_spec.sql): blanks that would silently blank a result, and out-of-range values."""
from __future__ import annotations

import pytest

from recet import preflight
from conftest import make_measure_csv


def _codes(path, severity):
    return {i.code for i in preflight.check_inputs(path).issues if i.severity == severity}


@pytest.mark.parametrize("column", ["NTGRkWh", "NumUnits", "EUL_Yrs", "UnitkWh1stBaseline"])
def test_a_blank_the_arithmetic_cannot_do_without_is_a_blocking_error(tmp_path, column):
    """A blank NTGRkWh used to give ElecBen = 0 and a blank (NaN) TSB and TRC cost with nothing flagged."""
    p = make_measure_csv(tmp_path / "m.csv", **{column: ""})
    assert "blank_required_value" in _codes(p, "error")


def test_every_field_cet_requires_must_have_a_value(tmp_path):
    """CET marks the numeric Measure fields NotEmpty: a blank one is an error (the pricing uses them as filed), single baseline or not. The fields CET does not
    require may be blank."""
    for col in preflight.NEEDED_VALUES:
        assert "blank_required_value" in _codes(make_measure_csv(tmp_path / f"{col}.csv", **{col: ""}), "error"), col
    ok = make_measure_csv(tmp_path / "ok.csv", UnitkWhIOUWater1stBaseline="", UnitDirectInstallLab="", UnitTaxCredits="", MeasInflation="")
    assert "blank_required_value" not in _codes(ok, "error")


@pytest.mark.parametrize("column,value", [("NTGRkWh", "1.6"), ("NTGRkWh", "0"), ("EUL_Yrs", "41"), ("RUL_Yrs", "31"), ("InstallationRatekWh", "1.2"),
                                          ("RealizationRatekWh", "2.5"), ("MeasInflation", "1.5"), ("UnitMeaCost1stBaseline", "-5")])
def test_values_outside_cets_ranges_are_warnings_not_errors(tmp_path, column, value):
    p = make_measure_csv(tmp_path / "m.csv", **{column: value, **({"EUL_Yrs": "35"} if column == "RUL_Yrs" else {})})   # an RUL above the EUL is an error
    rep = preflight.check_inputs(p)
    assert "outside_cet_range" in {i.code for i in rep.issues if i.severity == "warning"}
    assert rep.ok                                                  # the tool still prices it


def test_an_rul_above_the_eul_or_a_negative_life_is_an_error(tmp_path):
    assert "rul_above_eul" in _codes(make_measure_csv(tmp_path / "m.csv", RUL_Yrs="5", EUL_Yrs="4"), "error")
    assert "rul_not_below_eul" in _codes(make_measure_csv(tmp_path / "eq.csv", RUL_Yrs="4", EUL_Yrs="4"), "warning")     # equal: a warning, priced single-baseline
    assert "negative_life" in _codes(make_measure_csv(tmp_path / "n1.csv", EUL_Yrs="-1"), "error")
    assert "negative_life" in _codes(make_measure_csv(tmp_path / "n2.csv", RUL_Yrs="-0.5"), "error")


def test_a_clean_row_raises_none_of_them(tmp_path):
    codes = _codes(make_measure_csv(tmp_path / "m.csv"), "warning") | _codes(tmp_path / "m.csv", "error")
    assert not codes & {"blank_required_value", "negative_life", "rul_above_eul", "outside_cet_range", "rul_not_below_eul"}


def _rows(path, rows):
    """Write a Measure.csv of several rows (each a dict of overrides on the conftest measure)."""
    lines = []
    for i, over in enumerate(rows):
        make_measure_csv(path, **{"CEInputID": f"T{i + 1}", **over})
        text = path.read_text(encoding="utf-8").splitlines()
        lines = lines or [text[0]]
        lines.append(text[1])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def test_issues_say_which_cells_they_are_about(tmp_path):
    """The web editor highlights the cells a check names: (table, data row of the file, column)."""
    p = _rows(tmp_path / "m.csv", [{}, {"NTGRkWh": "2"}, {"AC_Elec_Impact_Profile": "NO_SUCH_SHAPE"}])
    rep = preflight.check_inputs(p).to_dict()
    cells = {i["code"]: i["cells"] for i in rep["issues"]}
    assert cells["outside_cet_range"] == [["Measure", 1, "NTGRkWh"]]
    assert ["Measure", 2, "AC_Elec_Impact_Profile"] in cells["no_electric_profile"]

    assert cells and rep["issues"][0].get("impact") is not None

    p = _rows(tmp_path / "d.csv", [{}, {"NumUnits": "x"}])
    (tmp_path / "d.csv").write_text((tmp_path / "d.csv").read_text().replace("\nT2|", "\nT1|"))
    cells = {i["code"]: i["cells"] for i in preflight.check_inputs(p).to_dict()["issues"]}
    assert cells["non_numeric"] == [["Measure", 1, "NumUnits"]]
    assert cells["duplicate_id"] == [["Measure", 0, "CEInputID"], ["Measure", 1, "CEInputID"]]


def test_unpriced_and_loosely_matched_records_say_which_and_how_much_energy(tmp_path, monkeypatch):
    """no_electric_profile, loose_profile_match and no_gas_profile name their records and total the first-year net kWh / therms those records carry
    (NumUnits x unit savings x NTG x installation x realization rates), in the report, its text and preflight_findings.xlsx."""
    from openpyxl import load_workbook
    from recet.pipeline import real_impact_profiles
    p = _rows(tmp_path / "m.csv", [
        {"AC_Elec_Impact_Profile": "NO_SUCH_SHAPE", "NumUnits": "2", "UnitkWh1stBaseline": "500", "NTGRkWh": "0.8"},          # 800 kWh, unpriced
        {"AC_Elec_Impact_Profile": "LOOSE_SHAPE", "AC_Elec_Target_Sector": "Odd", "UnitkWh1stBaseline": "300", "RealizationRatekWh": "0.5"},   # 150 kWh, loose
        {"UnitTherm1stBaseline": "40", "NumUnits": "3", "AC_Gas_Impact_Profile": "NO_SUCH_GAS_SHAPE"},                                       # 120 therms, unpriced
    ])
    monkeypatch.setattr(real_impact_profiles, "is_available", lambda: True)
    monkeypatch.setattr(real_impact_profiles, "_catalog_hourly", lambda: object())
    monkeypatch.setattr(real_impact_profiles, "match_profile_id",
                        lambda key, pa, ts: ("1", True) if key == "TEST_SHAPE" else (("2", False) if key == "LOOSE_SHAPE" else (None, False)))
    rep = preflight.check_inputs(p)
    by = {i.code: i for i in rep.issues}
    assert by["no_electric_profile"].impact == {preflight.IMPACT_KWH: 800.0} and by["no_electric_profile"].cells == [("Measure", 0, "AC_Elec_Impact_Profile")]
    loose = by["loose_profile_match"]
    assert loose.count == 1 and loose.impact == {preflight.IMPACT_KWH: 150.0} and loose.examples == ["LOOSE_SHAPE | PGE | Odd"]
    assert loose.cells == [("Measure", 1, "AC_Elec_Target_Sector")]
    assert by["no_gas_profile"].impact == {preflight.IMPACT_THERMS: 120.0} and by["no_gas_profile"].cells == [("Measure", 2, "AC_Gas_Impact_Profile")]
    assert "(150 first-year net kWh)" in rep.text() and "(120 first-year net therms)" in rep.text()
    ws = load_workbook(rep.write_xlsx(tmp_path / "f.xlsx", measure_ids=["T1", "T2", "T3"]))["Findings"]
    rows = {r[1]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    head = [c.value for c in ws[1]]
    assert rows["loose_profile_match"][head.index("First-year net kWh")] == 150 and rows["no_gas_profile"][head.index("First-year net therms")] == 120
    cells = [r for r in load_workbook(tmp_path / "f.xlsx")["Cells"].iter_rows(min_row=2, values_only=True) if r[1] == "loose_profile_match"]
    assert cells == [("INFO", "loose_profile_match", "Measure.csv", 3, "T2", "AC_Elec_Target_Sector")]
