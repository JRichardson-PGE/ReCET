"""End-to-end runs on a one-measure job with a synthetic ACC and a flat user profile: the expected values are computed by hand."""
from __future__ import annotations

import pandas as pd
import pytest

from recet import overrides, versions
from recet.jobs import RunConfig, run_job
from conftest import expected_flat_value, make_electric_acc, make_flat_profile_csv, make_measure_csv


def _cfg(tmp_path, **kw):
    base = dict(measures=str(make_measure_csv(tmp_path / "Measure.csv")), out_dir=str(tmp_path / "out"), run_name="t", acc_version="2026",
                acc_electric_file=str(make_electric_acc(tmp_path / "acc.parquet")), profiles_8760_file=str(make_flat_profile_csv(tmp_path / "p.csv")),
                emissions=False)
    base.update(kw)
    return RunConfig(**base)


def _rate(acc_key="2026"):
    from recet.pipeline import discount_settings
    return discount_settings.get_discount_rate_qtr("PGE", versions.BUILTIN[acc_key].acc)


def test_supplied_acc_and_profile_drive_the_result(tmp_path):
    r = run_job(_cfg(tmp_path), log=lambda s: None)
    f = r.frame.iloc[0]
    assert f["ElecPricingStatus"] == "priced: hourly profile" and f["ProfileSource"] == "user profile"
    assert f["ElecBen"] == pytest.approx(expected_flat_value(1000, 0.100, _rate()), rel=1e-12)      # $100/MWh x 1,000 kWh, discounted by quarter
    assert f["TotalSystemBenefit"] == pytest.approx(f["ElecBen"])
    assert f["Used_ACCVersion"].startswith("2026 ACC")
    assert (r.run_dir / "run_settings.json").exists() and (r.run_dir / "preflight.txt").exists()


def test_the_input_check_is_written_as_a_workbook(tmp_path):
    from openpyxl import load_workbook
    r = run_job(_cfg(tmp_path, measures=str(make_measure_csv(tmp_path / "M.csv", NTGRkWh="2")), preflight_only=True), log=lambda s: None)
    wb = load_workbook(r.files["preflight_findings"], read_only=True)
    assert wb.sheetnames == ["Summary", "Findings", "Cells"]
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Result"].startswith("no blocking problems") and summary["measures"] == "1"
    findings = list(wb["Findings"].iter_rows(min_row=2, values_only=True))
    assert ("WARNING", "outside_cet_range") in [(f[0], f[1]) for f in findings]
    assert [f[0] for f in findings] == sorted((f[0] for f in findings), key=["ERROR", "WARNING", "INFO"].index)    # errors first
    cells = list(wb["Cells"].iter_rows(min_row=2, values_only=True))
    assert ("WARNING", "outside_cet_range", "Measure.csv", 2, "T1", "NTGRkWh") in cells      # line 2 = the first record; its CEInputID


def test_replacement_uses_the_supplied_acc_not_the_builtin(tmp_path):
    cheap = run_job(_cfg(tmp_path, acc_electric_file=str(make_electric_acc(tmp_path / "a1.parquet", energy=20, ghg=0)), run_name="a"), log=lambda s: None)
    dear = run_job(_cfg(tmp_path, acc_electric_file=str(make_electric_acc(tmp_path / "a2.parquet", energy=200, ghg=0)), run_name="b"), log=lambda s: None)
    assert dear.frame.iloc[0]["ElecBen"] == pytest.approx(10 * cheap.frame.iloc[0]["ElecBen"], rel=1e-12)


def test_new_acc_label_is_defined_by_supplied_files(tmp_path):
    r = run_job(_cfg(tmp_path, acc_version="2030-draft", run_name="new"), log=lambda s: None)
    assert r.frame.iloc[0]["Used_ACCVersion"] == "Custom ACC '2030-draft'"
    assert r.frame.iloc[0]["ElecBen"] == pytest.approx(expected_flat_value(1000, 0.100, _rate("2026")), rel=1e-12)   # borrows the newest E3 version


def test_unknown_label_without_files_is_an_error():
    with pytest.raises(ValueError, match="not one the package ships"):
        versions.resolve_acc("nope")


def test_market_effects_scale_benefits_and_costs_by_ntg_plus_005(tmp_path):
    m = make_measure_csv(tmp_path / "M2.csv", NTGRkWh="0.8", NTGRkW="0.8", NTGRCost="0.8", UnitMeaCost1stBaseline="100")
    a = run_job(_cfg(tmp_path, measures=str(m), run_name="m0"), log=lambda s: None).frame.iloc[0]
    b = run_job(_cfg(tmp_path, measures=str(m), market_effects=0.05, run_name="m5"), log=lambda s: None).frame.iloc[0]
    assert b["ElecBen"] == pytest.approx(a["ElecBen"] * (0.8 + 0.05) / 0.8, rel=1e-12)
    assert b["TotalSystemBenefit"] == pytest.approx(a["TotalSystemBenefit"] * 0.85 / 0.8, rel=1e-12)
    # the one option also applies to costs (CET's MECost): with no incentives GrossMeasCostPV = GrossParticipantCostPV = G. As CET exports them,
    # NetParticipantCostPV stays 0.8 G and NetParticipantCostAdjustedPV (which the TRC cost uses) becomes 0.8 G + 0.05 G
    g = a["GrossParticipantCostPV"]
    assert g > 0 and a["NetParticipantCostPV"] == pytest.approx(0.8 * g, rel=1e-12) == b["NetParticipantCostPV"]
    assert b["MarkEffectPlusExcessIncPV"] == pytest.approx(0.05 * g, rel=1e-12)
    assert b["NetParticipantCostAdjustedPV"] == pytest.approx(0.85 * g, rel=1e-12)
    assert b["TRCCost"] - a["TRCCost"] == pytest.approx(0.05 * g, rel=1e-9)

def test_components_add_up_to_the_total(tmp_path):
    r = run_job(_cfg(tmp_path, components=True, run_name="c"), log=lambda s: None)
    f = r.frame.iloc[0]
    assert f["ElecValue_energy"] == pytest.approx(expected_flat_value(1000, 0.060, _rate()), rel=1e-12)
    assert f["ElecValue_ghg_adder"] == pytest.approx(expected_flat_value(1000, 0.040, _rate()), rel=1e-12)
    elec = [c for c in r.frame.columns if c.startswith("ElecValue_")]
    assert r.frame[elec].sum(axis=1).iloc[0] == pytest.approx(f["ElecBen"] - f["ElecSupplyCost"], rel=1e-12)
    assert (r.run_dir / "components_summary_acc2026.csv").exists()


def test_portfolio_8760_sums_to_the_measures_kwh(tmp_path):
    r = run_job(_cfg(tmp_path, portfolio_8760=True, run_name="p"), log=lambda s: None)
    prof = pd.read_csv(r.files["portfolio_8760"])
    assert len(prof) == 8760
    assert prof["net_kwh_2028"].sum() == pytest.approx(1000.0, rel=1e-9)                 # Q1 claim, 1-year life -> the full 1,000 kWh in 2028
    assert prof["net_kwh_2028"].nunique() == 1                                            # flat user profile


def test_preflight_reports_unpriced_measure(tmp_path):
    cfg = _cfg(tmp_path, profiles_8760_file=None, run_name="pf", preflight_only=True)
    r = run_job(cfg, log=lambda s: None)
    assert "no_electric_profile" in (r.run_dir / "preflight.txt").read_text() or "no_profile_library" in (r.run_dir / "preflight.txt").read_text()


def test_preflight_blocks_bad_input(tmp_path):
    bad = tmp_path / "bad.csv"
    make_measure_csv(bad, IOU_AC_Territory="XYZ")
    with pytest.raises(RuntimeError, match="blocking"):
        run_job(_cfg(tmp_path, measures=str(bad), run_name="bad"), log=lambda s: None)


def test_no_societal_blanks_scb_sch(tmp_path):
    f = run_job(_cfg(tmp_path, include_societal=False, run_name="ns"), log=lambda s: None).frame.iloc[0]
    assert pd.isna(f["SCBCost"]) and pd.isna(f["ElecBen_SB"]) and not pd.isna(f["TRCCost"])


def test_invalid_supplied_acc_stops_before_pricing(tmp_path):
    bad = pd.read_parquet(make_electric_acc(tmp_path / "ok.parquet")).drop(columns=["total_avoided_cost"])
    bad.to_parquet(tmp_path / "bad.parquet")
    with pytest.raises(overrides.OverrideError, match="not usable"):
        run_job(_cfg(tmp_path, acc_electric_file=str(tmp_path / "bad.parquet"), run_name="badacc"), log=lambda s: None)


def test_an_excel_job_over_the_limit_writes_every_table_as_zipped_csv(tmp_path):
    """excel_max_rows: a job with more measures (1 here, limit 0) writes every table as excel_fallback (zipped CSV by default), and says so."""
    import json
    lines = []
    r = run_job(_cfg(tmp_path, output_format="xlsx", excel_max_rows=0, components=True, run_name="xl"), log=lines.append)
    tables = {k: v.name for k, v in r.files.items() if k not in ("preflight", "preflight_findings", "run_settings")}   # the findings workbook is always xlsx
    assert tables and all(n.endswith(".csv.zip") for n in tables.values()), tables
    assert any("every table" in n and "zipped CSV" in n for n in r.notes)
    assert any("every table of this run is written as zipped CSV" in str(line) for line in lines)
    settings = json.loads((r.run_dir / "run_settings.json").read_text())
    assert settings["notes"] == r.notes and settings["config"]["output_format"] == "csv.zip"
    small = run_job(_cfg(tmp_path, output_format="xlsx", run_name="xl2"), log=lambda s: None)                     # within the default limit: Excel
    assert small.files["measure_level"].suffix == ".xlsx" and small.files["portfolio_summary"].suffix == ".xlsx" and not small.notes


def test_a_job_with_gas_profiles_named_but_no_gas_savings_runs(tmp_path):
    """pandas 3 reads a filled AC_Gas_Impact_Profile column as str, and an empty str Series sums to "" (run_poc.run_gas_benefit's count of matched profiles)."""
    cfg = _cfg(tmp_path, measures=str(make_measure_csv(tmp_path / "M.csv", AC_Gas_Impact_Profile="Res_Furnace")))
    assert run_job(cfg, log=lambda s: None).frame.iloc[0]["GasBen"] == 0
