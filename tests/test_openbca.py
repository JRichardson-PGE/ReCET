"""OpenBCA hosted by ReCET (recet/openbca): finding a checkout, staging a run's copy of it, the command line. These need no OpenBCA install; the one
test that runs OpenBCA itself (comparison C1 on its demo workbooks) runs only when RECET_OPENBCA_HOME names a synced checkout."""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

from recet.openbca import run


def _fake_checkout(root: Path) -> run.Checkout:
    for d in run.PROJECT_DIRS:
        (root / d).mkdir(parents=True)
    (root / "core" / "config.yaml").write_text("gateways: {}\n")
    (root / "core" / ".cache").mkdir()
    (root / "core" / ".cache" / "x").write_text("stale")
    (root / "excel_input_parsing" / "input_templates").mkdir()
    (root / "excel_input_parsing" / "input_templates" / run.PROGRAM_INPUT).write_bytes(b"someone else's job")
    return run.Checkout(root, Path(sys.executable))


def test_find_checkout_needs_a_checkout_and_its_interpreter(tmp_path, monkeypatch):
    monkeypatch.delenv("RECET_OPENBCA_HOME", raising=False)
    monkeypatch.delenv("RECET_OPENBCA_PYTHON", raising=False)
    with pytest.raises(FileNotFoundError, match="RECET_OPENBCA_HOME"):
        run.find_checkout()
    with pytest.raises(FileNotFoundError, match="not an OpenBCA checkout"):
        run.find_checkout(tmp_path)
    _fake_checkout(tmp_path / "ob")
    with pytest.raises(FileNotFoundError, match="uv sync"):
        run.find_checkout(tmp_path / "ob")
    got = run.find_checkout(tmp_path / "ob", python=sys.executable)
    assert got.home == (tmp_path / "ob").resolve() and got.python == Path(sys.executable)


def test_stage_copies_the_projects_and_puts_the_workbooks_under_openbcas_names(tmp_path):
    co = _fake_checkout(tmp_path / "ob")
    (tmp_path / "mine.xlsm").write_bytes(b"program input")
    (tmp_path / "cfg.xlsx").write_bytes(b"configuration")
    ws = run.stage(co, tmp_path / "run" / "openbca", tmp_path / "mine.xlsm", tmp_path / "cfg.xlsx")
    templates = ws / "excel_input_parsing" / "input_templates"
    assert (templates / run.PROGRAM_INPUT).read_bytes() == b"program input"
    assert (templates / run.CONFIGURATION).read_bytes() == b"configuration"
    assert (ws / "core" / "config.yaml").exists() and (ws / "engine.py").exists()
    assert not (ws / "core" / ".cache").exists()  # no SQLmesh state carried in from the checkout
    # the checkout itself is never written to
    assert (tmp_path / "ob" / "excel_input_parsing" / "input_templates" / run.PROGRAM_INPUT).read_bytes() == b"someone else's job"


def test_engine_is_not_imported_by_recet():
    """engine.py runs in OpenBCA's interpreter; importing recet.openbca must not need SQLmesh or DuckDB."""
    import recet.openbca.run  # noqa: F401
    assert "sqlmesh" not in sys.modules


def test_cli_reports_a_missing_checkout(monkeypatch, capsys):
    from recet import cli
    monkeypatch.delenv("RECET_OPENBCA_HOME", raising=False)
    with pytest.raises(FileNotFoundError):
        cli.main(["openbca", "compare", "a.db", "b.db"])


@pytest.mark.skipif(not os.environ.get("RECET_OPENBCA_HOME"), reason="needs RECET_OPENBCA_HOME (a synced OpenBCA checkout)")
def test_c1_hosted_run_of_the_demo_matches_a_second_run(tmp_path):
    """C1 in small: two hosted runs of OpenBCA's demo workbooks give the same tables (the full C1 compares with OpenBCA's own `make run-openbca-model`)."""
    co = run.find_checkout()
    demo = co.home / "demo_input_files"
    a = run.run_openbca(demo / run.PROGRAM_INPUT, demo / run.CONFIGURATION, tmp_path, "a", checkout=co)
    b = run.run_openbca(demo / run.PROGRAM_INPUT, demo / run.CONFIGURATION, tmp_path, "b", checkout=co)
    assert "openbca_results_summary_by_id" in a.files and "openbca_jst_ratio" in a.files
    assert run.compare_dbs(a.run_dir / "openbca" / "openbca.db", b.run_dir / "openbca" / "openbca.db", co)


def test_a_relative_checkout_gives_an_absolute_interpreter(tmp_path, monkeypatch):
    """The engine runs with cwd = the run's workspace, so a relative interpreter path would point nowhere."""
    _fake_checkout(tmp_path / "ob")
    py = tmp_path / "ob" / ".venv" / "bin" / "python"
    py.parent.mkdir(parents=True)
    py.write_text("")
    monkeypatch.chdir(tmp_path)
    got = run.find_checkout("ob", python="ob/.venv/bin/python")
    assert got.python.is_absolute() and got.python == py


def _tiny_acc(path: Path) -> Path:
    """An ACC electric file in ReCET's layout ($/MWh, hours 1..8760), two IOU-zones, two years, TRC and SCT Base."""
    import numpy as np
    import pandas as pd
    rows = []
    for iou, cz in (("PG&E", "CZ12"), ("SCE", "CZ6")):
        for test in ("TRC", "SCT Base"):
            for year in (2026, 2027):
                h = np.arange(1, 8761)
                rows.append(pd.DataFrame({"iou": iou, "climate_zone": cz, "cost_test": test, "year": year, "hour_of_year": h,
                                          "hour": (h - 1) % 24, "month": np.minimum((h - 1) // 730 + 1, 12),
                                          "total_avoided_cost": 50.0 + (h % 97) * 1.0007 + (year - 2026) * 3 + (test != "TRC") * 10}))
    pd.concat(rows).to_parquet(path, index=False)
    return path


def test_ca_avoided_costs_are_the_accs_total_in_dollars_per_kwh(tmp_path):
    import numpy as np
    from recet.openbca import ca_config
    acc = _tiny_acc(tmp_path / "acc.parquet")
    got = ca_config.avoided_costs_ts(acc, [("PG&E", "CZ12")], (2026, 2027))
    assert len(got) == 2 * 8760 and set(got["avoided_cost_subset"]) == {"PG&E CZ12"}
    assert got["hour_of_year"].min() == 0 and got["hour_of_year"].max() == 8759
    assert got["hour_of_day"].min() == 1 and got["hour_of_day"].max() == 24
    first = got[(got.year == 2026) & (got.hour_of_year == 0)]["avoided_cost_value"].item()
    assert np.isclose(first, (50.0 + 1 * 1.0007) / 1000)
    with pytest.raises(ValueError, match="no TRC rows"):
        ca_config.avoided_costs_ts(acc, [("SDG&E", "CZ7")], (2026, 2027))
    # C2: OpenBCA's table (its values rounded to $0.001/kWh, as its NUMERIC cast does) against the ACC
    back = got.assign(avoided_cost_value=ca_config.rounded_like_openbca(got["avoided_cost_value"].to_numpy()))
    c2 = ca_config.check_against_acc(acc, back)
    assert len(c2) == 2 and (c2["difference"].abs() < 1e-9).all() and (c2["openbca"] != c2["acc"]).any()


def test_ca_build_writes_every_jurisdiction_table(tmp_path):
    import pandas as pd
    from recet.openbca import ca_config
    files = ca_config.build(_tiny_acc(tmp_path / "acc.parquet"), [("SCE", "CZ6")], (2026, 2027), "PAC", 0.0766, tmp_path / "jur")
    assert set(files) == set(ca_config.TABLES) == set(run.JURISDICTION_MODELS)
    for table, (_, columns) in run.JURISDICTION_MODELS.items():
        assert list(pd.read_parquet(files[table]).columns) == list(columns)
    g = pd.read_parquet(files["global_parameters"]).iloc[0]
    assert g["cost_treatment"] == "USI" and g["inflation_rate"] == 0 and g["discount_rate"] == 0.0766
    with pytest.raises(ValueError):
        ca_config.global_parameters("SCT", 0.07, 2026)


def test_program_input_workbook_has_openbcas_layout(tmp_path):
    import numpy as np
    import pandas as pd
    from openpyxl import load_workbook
    from recet.openbca import workbook
    m = pd.DataFrame([{"Unique ID": "M0", "Program Name": "P1", "Unit Quantity": 2, "Adj. Gross Annual Electric Savings (kWh)": np.float64(100.0),
                       "Net-to-gross Ratio": np.nan, "Electric Savings Load Shape": "Flat"}])
    p = pd.DataFrame([{"Program Name": "P1", "Program Year": 2026, "Program Admin Costs ($)": 5.0}])
    path = workbook.write_program_input(tmp_path / "pi.xlsx", m, p, {"Flat": np.full(8760, 1 / 8760)})
    wb = load_workbook(path, read_only=True)
    rows = list(wb["Measure Inputs"].iter_rows(values_only=True))
    assert list(rows[3]) == workbook.MEASURE_HEADERS
    rec = dict(zip(workbook.MEASURE_HEADERS, rows[4]))
    assert rec["Unique ID"] == "M0" and rec["Adj. Gross Annual Electric Savings (kWh)"] == 100.0 and rec["Net-to-gross Ratio"] is None
    assert list(list(wb["Program Inputs"].iter_rows(values_only=True))[2]) == workbook.PROGRAM_HEADERS
    shapes = list(wb["Electric Load Shapes"].iter_rows(values_only=True))
    assert shapes[1] == ("Hour", "Flat") and len(shapes) == 2 + 8760 and shapes[-1][0] == 8760
    with pytest.raises(ValueError, match="not OpenBCA"):
        workbook.write_program_input(tmp_path / "x.xlsx", m.assign(Bogus=1), None, {"Flat": np.zeros(8760)})
    with pytest.raises(ValueError, match="8,760"):
        workbook.write_program_input(tmp_path / "y.xlsx", m, None, {"Flat": np.zeros(10)})
