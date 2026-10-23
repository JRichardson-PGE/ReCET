"""Run parameters, statewide allocation factors and the Excel job workbook: each supplied input must change exactly what it says, and a workbook job must equal the
same job given as files."""
from __future__ import annotations

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from recet import overrides, templates, workbook
from recet.jobs import RunConfig, run_job
from recet.pipeline import statewide_allocation
from conftest import expected_flat_value, make_electric_acc, make_flat_profile_csv, make_measure_csv

QUIET = dict(log=lambda s: None)


def _cfg(tmp_path, **kw):
    base = dict(measures=str(make_measure_csv(tmp_path / "Measure.csv")), out_dir=str(tmp_path / "out"), run_name="t", acc_version="2026",
                acc_electric_file=str(make_electric_acc(tmp_path / "acc.parquet")), profiles_8760_file=str(make_flat_profile_csv(tmp_path / "p.csv")),
                emissions=False)
    base.update(kw)
    return RunConfig(**base)


def _run_params(path, **rows):
    """rows: PA -> (annual rate, annual SCT rate)."""
    pd.DataFrame([{"PA": pa, "DiscountRateAnnual": a, "DiscountRateAnnual_SCT": s, "CO2Gas": 0.00531} for pa, (a, s) in rows.items()]).to_csv(path, index=False)
    return path


# ---------------------------------------------------------------------------------------------------------------- run parameters
def test_run_parameters_replace_the_discount_rates(tmp_path):
    f = run_job(_cfg(tmp_path, run_parameters_file=str(_run_params(tmp_path / "rp.csv", PGE=(0.12, 0.02)))), **QUIET).frame.iloc[0]
    assert f["ElecBen"] == pytest.approx(expected_flat_value(1000, 0.100, 0.12 / 4), rel=1e-12)          # the quarterly rate defaults to annual / 4
    assert f["ElecBen_SB"] == pytest.approx(expected_flat_value(1000, 0.100, 0.02 / 4), rel=1e-12)


def test_run_parameters_leave_unlisted_pas_alone(tmp_path):
    base = run_job(_cfg(tmp_path, run_name="a"), **QUIET).frame.iloc[0]
    other = run_job(_cfg(tmp_path, run_name="b", run_parameters_file=str(_run_params(tmp_path / "rp.csv", SCE=(0.30, 0.30)))), **QUIET).frame.iloc[0]
    assert other["ElecBen"] == pytest.approx(base["ElecBen"], rel=1e-14)


def test_run_parameters_validation(tmp_path):
    with pytest.raises(overrides.OverrideError, match="fraction"):
        overrides.load_run_parameters(_run_params(tmp_path / "bad.csv", PGE=(7.66, 0.05)))
    pd.DataFrame({"PA": ["PGE"], "DiscountRateAnnual": [0.08]}).to_csv(tmp_path / "short.csv", index=False)
    with pytest.raises(overrides.OverrideError, match="missing column"):
        overrides.load_run_parameters(tmp_path / "short.csv")
    t, notes = overrides.load_run_parameters(templates.write_template("run-parameters", tmp_path / "tpl.csv", "2026"))
    assert list(t.index) == ["PGE", "SCE", "SDGE", "SCG"] and not notes and t.loc["PGE", "DiscountRateAnnual"] == pytest.approx(0.0766)


# ---------------------------------------------------------------------------------------------------------------- SW allocation factors
def _sw(path, year=2028, program="SW_TEST", pge=0.25, sce=0.75):
    row = {"Year": year, "ProgramID": program, "PercentElectric": 1.0}
    for iou, e in (("PGE", pge), ("SDGE", 0.0), ("SCE", sce), ("SCG", 0.0)):
        row[f"{iou}ElectricFactor"], row[f"{iou}GasFactor"] = e, 0.0
    pd.DataFrame([row]).to_csv(path, index=False)
    return path


def test_sw_factors_split_a_statewide_program(tmp_path):
    meas = make_measure_csv(tmp_path / "Measure_sw.csv", PrgID="SW_TEST")
    plain = run_job(_cfg(tmp_path, measures=str(meas), run_name="a"), **QUIET).frame.set_index("CEInputID")
    split = run_job(_cfg(tmp_path, measures=str(meas), run_name="b", allocate_statewide=True, sw_factors_file=str(_sw(tmp_path / "sw.csv"))),
                    **QUIET).frame.set_index("CEInputID")
    assert set(split.index) == {"PGE_T1", "SCE_T1"}
    assert split.loc["PGE_T1", "ElecBen"] == pytest.approx(0.25 * plain.loc["T1", "ElecBen"], rel=1e-12)
    assert split.loc["SCE_T1", "ElecBen"] == pytest.approx(0.75 * plain.loc["T1", "ElecBen"], rel=1e-12)
    assert split["ElecBen"].sum() == pytest.approx(plain.loc["T1", "ElecBen"], rel=1e-12)


def test_sw_factors_replace_and_extend_the_builtin_table(tmp_path):
    builtin = statewide_allocation.load_factors()
    key = builtin.index[0]
    user, _ = overrides.load_sw_factors(_sw(tmp_path / "sw.csv", year=key[0], program=key[1], pge=1.0, sce=0.0))
    statewide_allocation.set_user_factors(user)
    try:
        merged = statewide_allocation.load_factors()
    finally:
        statewide_allocation.set_user_factors(None)
    assert len(merged) == len(builtin) and merged.loc[key, "PGEElectricFactor"] == 1.0 and merged.loc[key, "PGEWeightedFactor"] == 1.0
    user2, _ = overrides.load_sw_factors(_sw(tmp_path / "sw2.csv", program="SW_NEW"))
    statewide_allocation.set_user_factors(user2)
    try:
        assert len(statewide_allocation.load_factors()) == len(builtin) + 1
    finally:
        statewide_allocation.set_user_factors(None)


def test_sw_factors_validation(tmp_path):
    p = _sw(tmp_path / "sw.csv", pge=0.5, sce=0.3)
    _t, notes = overrides.load_sw_factors(p)
    assert any("do not sum to 1" in n for n in notes)
    df = pd.read_csv(p).drop(columns="PercentElectric")
    df.to_csv(tmp_path / "nopct.csv", index=False)
    with pytest.raises(overrides.OverrideError, match="WeightedFactor"):
        overrides.load_sw_factors(tmp_path / "nopct.csv")


# ---------------------------------------------------------------------------------------------------------------- the job workbook
def _workbook_from(tmp_path, measure_csv, run_params=None, programcost=None):
    p = templates.write_template("workbook", tmp_path / "job.xlsx", "2026")
    wb = load_workbook(p)

    def cell(v):                                     # numbers as numbers, the way a user's Excel holds them
        try:
            x = float(v)
        except ValueError:
            return v or None
        return int(x) if x.is_integer() else x

    def fill(ws, df):
        ws.delete_rows(2, ws.max_row)
        for r in df.itertuples(index=False):
            ws.append([cell(v) for v in r])
    hdr = [c.value for c in wb["Measure"][1]]
    m = pd.read_csv(measure_csv, sep="|", dtype=str, keep_default_na=False)
    fill(wb["Measure"], m.reindex(columns=hdr, fill_value=""))
    wb["ProgramCost"].delete_rows(2, wb["ProgramCost"].max_row)
    if run_params is not None:
        ws = wb["RunParameters"]
        ws.delete_rows(2, ws.max_row)
        ws.append(["PGE", run_params, run_params / 4, 0.0503, 0.012575, 2024, 0.00531])
    wb.save(p)
    return p


def test_workbook_job_equals_the_file_job(tmp_path):
    files = run_job(_cfg(tmp_path, run_name="files"), **QUIET).frame
    wbp = _workbook_from(tmp_path, tmp_path / "Measure.csv")
    book = run_job(_cfg(tmp_path, run_name="book", measures="", workbook=str(wbp)), **QUIET).frame
    for c in ("ElecBen", "ElecBen_SB", "TRCCost", "TotalSystemBenefit", "ElecPricingStatus", "AC_Climate_Zone", "ClaimYearQuarter"):
        assert files.iloc[0][c] == book.iloc[0][c], c


def test_workbook_sheets_feed_the_run(tmp_path):
    wbp = _workbook_from(tmp_path, make_measure_csv(tmp_path / "Measure_wb.csv"), run_params=0.12)
    f = run_job(_cfg(tmp_path, measures="", workbook=str(wbp)), **QUIET).frame.iloc[0]
    assert f["ElecBen"] == pytest.approx(expected_flat_value(1000, 0.100, 0.12 / 4), rel=1e-12)


def test_workbook_reads_cells_as_cet_text(tmp_path):
    wbp = templates.write_template("workbook", tmp_path / "job.xlsx", "2026")
    out = workbook.extract(wbp, tmp_path / "x")
    m = pd.read_csv(out["measures"], sep="|", dtype=str, keep_default_na=False)
    assert m.loc[0, "AC_Climate_Zone"] == "13" and m.loc[0, "ClaimYearQuarter"] == "2028Q3" and m.loc[0, "NumUnits"] == "100"
    assert set(out) == {"measures", "programcost", "run_parameters_file", "sw_factors_file", "programs_file"}      # empty profile sheets count as not given
    assert workbook._text(13.0) == "13" and workbook._text(0.6) == "0.6" and workbook._text(None) == "" and workbook._text(True) == "1"


def test_workbook_without_measures_is_an_error(tmp_path):
    p = templates.write_template("workbook", tmp_path / "job.xlsx", "2026")
    wb = load_workbook(p)
    wb["Measure"].delete_rows(2, wb["Measure"].max_row)
    wb.save(p)
    with pytest.raises(workbook.WorkbookError, match="no Measure sheet"):
        workbook.extract(p, tmp_path / "x")


# ---------------------------------------------------------------------------------------------------------------- Excel output
def test_xlsx_output_holds_the_same_numbers_and_text(tmp_path):
    from recet import outputs
    r = run_job(_cfg(tmp_path, output_format="xlsx", portfolio_8760=True), **QUIET)
    for key in ("measure_level", "portfolio_summary", "portfolio_8760"):
        assert r.files[key].suffix == ".xlsx", key
    back = outputs.read_table(r.files["measure_level"])
    f = r.frame
    assert list(back.columns) == list(f.columns)
    assert back.loc[0, "ElecBen"] == pytest.approx(f.loc[0, "ElecBen"], rel=1e-12) and back.loc[0, "CEInputID"] == "T1"
    wb = load_workbook(r.files["measure_level"], read_only=True)
    ws = wb.worksheets[0]
    assert ws.title == "measure_level_acc2026" and next(ws.iter_rows(max_row=1, values_only=True))[0] == f.columns[0]
    wb.close()


def test_xlsx_text_starting_with_equals_is_not_a_formula(tmp_path):
    from recet import outputs
    p = outputs.write_table(pd.DataFrame({"a": ["=1+1", "plain"], "b": [1.5, np.nan]}), tmp_path / "t", "xlsx")
    wb = load_workbook(p)
    ws = wb.active
    assert ws["A2"].value == "=1+1" and ws["A2"].data_type == "s" and ws["B3"].value is None
