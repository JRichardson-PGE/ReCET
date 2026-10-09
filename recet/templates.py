"""Templates for every kind of user-supplied input (and the CET Measure.csv header and an Excel job workbook), plus an exporter that turns an existing library profile
into an editable file.

  profiles-8760   long CSV: end_use_name, pa, target_sector, climate_zone, hour_of_year, impact, represents_year      (overrides.load_user_profiles)
  profiles-gas    CSV: Profile, Month, Impact                                                                          (overrides.load_user_gas_profiles)
  acc-electric    CSV/Parquet: one row per (iou, climate_zone, cost_test, year, hour_of_year) with the component columns   (overrides.check_electric_acc)
  acc-gas         CSV/Parquet: one row per (utility, class, cost_test, scc, year) with monthly market/td and annual components (overrides.check_gas_acc)
  run-parameters  CSV: one row per PA with the built-in run parameters of the chosen ACC version (discount rates, base year, gas CO2)   (overrides.load_run_parameters)
  sw-factors      CSV: the built-in statewide allocation factor table, to edit                                                   (overrides.load_sw_factors)
  programs        CSV: the built-in program metadata (CEDARS Programs list), to edit                                             (programs.load_programs)
  workbook        XLSX: a job workbook -- Instructions, Measure, ProgramCost, RunParameters, SWAllocation, Programs, GasProfiles, Profiles8760   (workbook.extract)
  measure         the CET Measure.csv header (pipe-delimited)

`recet template KIND --out FILE` and the GUI's "Save template" buttons call `write_template`. Where the reference data is available the ACC templates are real slices
of the built-in ACC (so the values are realistic and the layout is exactly what the run reads); otherwise they are filled with zeros. The full format description is
`recet/docs/INPUT_FORMATS.md` (also written next to any template by `write_all`).
"""
from __future__ import annotations

import shutil
from pathlib import Path

import numpy as np
import pandas as pd

from . import overrides, paths, versions

KINDS = ("workbook", "profiles-8760", "profiles-gas", "acc-electric", "acc-gas", "run-parameters", "sw-factors", "programs", "measure")
DOCS_DIR = Path(__file__).resolve().parent / "docs"

_CET_MEASURE_HEADER = [
    "CEInputID", "PrgID", "ClaimYearQuarter", "Sector", "DeliveryType", "BldgType", "E3ClimateZone", "E3GasSavProfile", "E3GasSector", "E3MeaElecEndUseShape",
    "E3TargetSector", "MeasAppType", "MeasCode", "MeasDescription", "MeasImpactType", "FuelID", "MeasureID", "TechGroup", "TechType", "UseCategory", "UseSubCategory",
    "WaterUse", "PreDesc", "StdDesc", "SourceDesc", "Version", "NormUnit", "NumUnits", "UnitkW1stBaseline", "UnitkWh1stBaseline", "UnitTherm1stBaseline",
    "UnitkW2ndBaseline", "UnitkWh2ndBaseline", "UnitTherm2ndBaseline", "UnitkWhIOUWater1stBaseline", "UnitkWhIOUWater2ndBaseline", "UnitMeaCost1stBaseline",
    "UnitMeaCost2ndBaseline", "UnitDirectInstallLab", "UnitDirectInstallMat", "UnitEndUserRebate", "UnitIncentiveToOthers", "NTG_ID", "NTGRkW", "NTGRkWh", "NTGRTherm",
    "NTGRCost", "EUL_ID", "EUL_Yrs", "RUL_ID", "RUL_Yrs", "GSIA_ID", "RealizationRatekW", "RealizationRatekWh", "RealizationRateTherm", "InstallationRatekW",
    "InstallationRatekWh", "InstallationRateTherm", "Residential_Flag", "Upstream_Flag", "IOU_AC_Territory", "UnitGasInfraBens", "UnitRefrigCosts", "UnitRefrigBens",
    "UnitMiscCosts", "MiscCostsDesc", "UnitMiscBens", "MiscBensDesc", "UnitTaxCredits", "RateScheduleElec", "RateScheduleGas", "CombustionType", "MeasInflation", "Comments",
]


def _electric_template(acc_key: str) -> pd.DataFrame:
    """One year of one zone (PG&E CZ13, TRC) from the built-in ACC if it can be read, else zeros in the same layout."""
    p = paths.ref_path(versions.BUILTIN[acc_key].electric_key)
    if p is not None:
        try:
            first = pd.read_parquet(p, columns=["year"], filters=[("iou", "==", "PG&E"), ("climate_zone", "==", "CZ13"), ("cost_test", "==", "TRC")])["year"].min()
            df = pd.read_parquet(p, filters=[("iou", "==", "PG&E"), ("climate_zone", "==", "CZ13"), ("cost_test", "==", "TRC"), ("year", "==", int(first))])
            df = df.drop(columns=[c for c in ("dr_total_ghg",) if c in df.columns])
            for c in ("iou", "climate_zone", "cost_test", "price_zone"):
                if c in df.columns:
                    df[c] = df[c].astype(str)
            return df.reset_index(drop=True)
        except Exception:
            pass
    year = 2026
    df = pd.DataFrame({"iou": "PG&E", "climate_zone": "CZ13", "cost_test": "TRC", "year": year, "hour_of_year": np.arange(1, 8761),
                       "datetime": pd.date_range(f"{year}-01-01", periods=8760, freq="h")})
    for c in ("ghg_cap_and_invest", "ghg_adder", "air_quality_adder", "energy", "generation_capacity", "transmission", "distribution", "avoided_as_procurement",
              "losses", "methane_leakage", "total_avoided_cost", "emission_rate_tonnes_per_kwh"):
        df[c] = 0.0
    return df


def _gas_template(acc_key: str) -> pd.DataFrame:
    p = paths.package_data(*versions.BUILTIN[acc_key].gas_rel)
    df = pd.read_parquet(p)
    first = int(df["year"].min())
    df = df[(df["utility"] == "PG&E") & (df["class"] == "Residential") & (df["year"].isin([first, first + 1]))]
    return df.reset_index(drop=True)


_PROGRAMCOST_HEADER = ["PrgID", "PrgYear", "ClaimYearQuarter", "AdminCostsOverheadAndGA", "AdminCostsOther", "MarketingOutreach", "DIActivity", "DIInstallation",
                       "DIHardwareAndMaterials", "DIRebateAndInspection", "EMV", "UserInputIncentive", "OnBillFinancing", "CostsRecoveredFromOtherSources", "PA"]

# One realistic example row for the workbook template (a PG&E commercial lighting measure on a TOU end use); the user replaces it.
_EXAMPLE_MEASURE = {
    "CEInputID": "EXAMPLE_1", "PrgID": "EXAMPLE_PRG", "ClaimYearQuarter": "2028Q3", "Sector": "Commercial", "DeliveryType": "DnDeemed", "BldgType": "Com",
    "E3ClimateZone": "13", "E3GasSavProfile": "Annual", "E3GasSector": "Commercial", "E3MeaElecEndUseShape": "1 = Commercial Indoor Lighting",
    "E3TargetSector": "COMMERCIAL", "MeasAppType": "NR", "MeasDescription": "Example: LED troffer retrofit (replace this row)", "MeasImpactType": "Deem",
    "FuelID": "EE-Elec-Only", "NormUnit": "Each", "NumUnits": "100", "UnitkW1stBaseline": "0.03", "UnitkWh1stBaseline": "120", "UnitTherm1stBaseline": "-0.5",
    "UnitkW2ndBaseline": "0", "UnitkWh2ndBaseline": "0", "UnitTherm2ndBaseline": "0", "UnitkWhIOUWater1stBaseline": "0", "UnitkWhIOUWater2ndBaseline": "0",
    "UnitMeaCost1stBaseline": "85", "UnitMeaCost2ndBaseline": "0", "UnitDirectInstallLab": "0", "UnitDirectInstallMat": "0", "UnitEndUserRebate": "25",
    "UnitIncentiveToOthers": "0", "NTGRkW": "0.6", "NTGRkWh": "0.6", "NTGRTherm": "0.6", "NTGRCost": "0.6", "EUL_Yrs": "15", "RUL_Yrs": "0",
    "RealizationRatekW": "1", "RealizationRatekWh": "1", "RealizationRateTherm": "1", "InstallationRatekW": "1", "InstallationRatekWh": "1",
    "InstallationRateTherm": "1", "Residential_Flag": "0", "Upstream_Flag": "0", "IOU_AC_Territory": "PGE", "UnitGasInfraBens": "0", "UnitRefrigCosts": "0",
    "UnitRefrigBens": "0", "UnitMiscCosts": "0", "UnitMiscBens": "0", "UnitTaxCredits": "0", "CombustionType": "Small Boilers (<100 MMBtu/hr Heat Input):Uncontrolled", "MeasInflation": "0",
}
_EXAMPLE_PROGRAMCOST = {"PrgID": "EXAMPLE_PRG", "PrgYear": "2028", "ClaimYearQuarter": "2028Q4", "AdminCostsOverheadAndGA": "1500", "AdminCostsOther": "0",
                        "MarketingOutreach": "500", "DIActivity": "0", "DIInstallation": "0", "DIHardwareAndMaterials": "0", "DIRebateAndInspection": "0", "EMV": "0",
                        "UserInputIncentive": "0", "OnBillFinancing": "0", "CostsRecoveredFromOtherSources": "0", "PA": "PGE"}


# Public names (the input editor, editor.py, builds its tables from these)
MEASURE_COLUMNS = _CET_MEASURE_HEADER
PROGRAMCOST_COLUMNS = _PROGRAMCOST_HEADER
EXAMPLE_MEASURE = _EXAMPLE_MEASURE
EXAMPLE_PROGRAMCOST = _EXAMPLE_PROGRAMCOST


def run_parameters_template(acc_key: str = "2026") -> pd.DataFrame:
    """The built-in run parameters (CET E3Settings) of an ACC version's E3 tables, one row per IOU, in the run-parameters file layout."""
    from .pipeline import discount_settings
    version = versions.BUILTIN[acc_key].cet_version if acc_key in versions.BUILTIN else max(discount_settings.known_versions())
    t = discount_settings._table()
    t = t[(t["ACC Version"] == version) & t["PA"].isin(["PGE", "SCE", "SDGE", "SCG"])].set_index("PA").loc[["PGE", "SCE", "SDGE", "SCG"]].reset_index()
    return t[["PA", "DiscountRateAnnual", "DiscountRateQtr", "DiscountRateAnnual_SCT", "DiscountRateQtr_SCT", "BaseYear", "CO2Gas"]]


def sw_factors_template() -> pd.DataFrame:
    from .pipeline import statewide_allocation
    return pd.read_csv(statewide_allocation.FACTORS_CSV)


def programs_template() -> pd.DataFrame:
    from . import programs
    return programs.builtin()


def write_workbook_template(out_path: str | Path, acc_key: str = "2026") -> Path:
    """An Excel job workbook: an Instructions sheet, the data sheets with their headers, one example row in Measure and ProgramCost, and the built-in run parameters
    and SW factors filled in (so leaving those sheets as they are changes nothing)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from . import workbook as wbmod

    head_font, body_font = Font(name="Arial", bold=True, color="FFFFFF", size=10), Font(name="Arial", size=10)
    head_fill, example_fill = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    lines = [
        ("ReCET - job workbook", True),
        ("Use this workbook instead of separate files: choose it as 'Job workbook' in the window, or run  recet run --workbook FILE.xlsx", False),
        ("", False),
        ("Sheet", True),
        ("Measure        REQUIRED. CET Measure.csv columns, one measure per row. Replace the yellow example row.", False),
        ("ProgramCost    Optional. CET ProgramCost.csv columns. Replace or delete the yellow example row.", False),
        ("RunParameters  Optional. Discount rates, base year and gas CO2 per PA. Filled in with the built-in values for ACC " + acc_key
         + "; a listed PA replaces the built-in row, so edit only the values you want to change. Rates are fractions (0.0766 = 7.66%).", False),
        ("SWAllocation   Optional. Statewide allocation factors, filled in with the built-in table. Used only when 'Allocate statewide programs' is on; "
         "a (Year, ProgramID) row replaces the built-in one and a new one is added.", False),
        ("Programs       Optional. Program metadata (the CEDARS Programs list), filled in with the built-in table. A PrgID replaces the built-in row and a new one "
         "is added; it adds Prg_ columns to the results and drives the 'program segment' summary rows and the Exclude_From_CE option.", False),
        ("GasProfiles    Optional. Gas monthly profiles (Profile, Month 1-12, Impact). A name in the built-in library replaces it.", False),
        ("Profiles8760   Optional. 8,760 electric profiles in long format (8,760 rows per profile and zone; Excel fits about 119). For more, use a separate file.", False),
        ("", False),
        ("Rules", True),
        ("Keep the header row and the sheet names. A sheet with only its header row is treated as not given; other sheets (like this one) are ignored.", False),
        ("Whole numbers are read without decimals, so a climate zone typed as 13 is read as '13'. No cell may contain the character '|'.", False),
        ("The ACC files are too large for Excel: supply them as separate files if you need them.", False),
        ("A file chosen separately in the window (for example a run-parameters CSV) wins over the matching sheet.", False),
        ("Full formats: Help > Documentation > Input formats.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ins.cell(i, 1, text)
        c.font = Font(name="Arial", size=12 if i == 1 else 10, bold=bold)
    ins.column_dimensions["A"].width = 150

    def sheet(name, columns, rows, example=False):
        ws = wb.create_sheet(name)
        for j, col in enumerate(columns, start=1):
            c = ws.cell(1, j, col)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(str(col)) + 2))
        for i, r in enumerate(rows, start=2):
            for j, col in enumerate(columns, start=1):
                v = r.get(col, "")
                if isinstance(v, str) and v != "":
                    try:
                        v = float(v) if any(ch in v for ch in ".eE") else int(v)
                    except ValueError:
                        pass
                elif not isinstance(v, str) and pd.isna(v):
                    v = None
                c = ws.cell(i, j, v)
                c.font = body_font
                if example:
                    c.fill = example_fill
        ws.freeze_panes = "A2"
        return ws

    meas = sheet("Measure", _CET_MEASURE_HEADER, [_EXAMPLE_MEASURE], example=True)
    for col in ("CEInputID", "PrgID", "ClaimYearQuarter", "E3ClimateZone"):           # keep these as text even when they look like numbers
        j = _CET_MEASURE_HEADER.index(col) + 1
        meas.cell(2, j).number_format = "@"
        meas.cell(2, j).value = str(_EXAMPLE_MEASURE[col])
    sheet("ProgramCost", _PROGRAMCOST_HEADER, [_EXAMPLE_PROGRAMCOST], example=True)
    rp = run_parameters_template(acc_key)
    sheet("RunParameters", list(rp.columns), rp.to_dict("records"))
    sw = sw_factors_template()
    sheet("SWAllocation", list(sw.columns), sw.to_dict("records"))
    pg = programs_template()
    sheet("Programs", list(pg.columns), pg.to_dict("records"))
    sheet("GasProfiles", ["Profile", "Month", "Impact"], [])
    sheet("Profiles8760", ["end_use_name", "pa", "target_sector", "climate_zone", "hour_of_year", "impact", "represents_year"], [])
    assert set(wbmod.SHEETS) <= set(wb.sheetnames)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def write_template(kind: str, out_path: str | Path, acc_key: str = "2026") -> Path:
    """Write one template file (format chosen by the file extension for tables: .csv or .parquet)."""
    if kind not in KINDS:
        raise ValueError(f"unknown template {kind!r}; choose from {KINDS}")
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    if kind == "profiles-8760":
        df = pd.DataFrame({
            "end_use_name": "EXAMPLE_END_USE (replace: use the measure's E3MeaElecEndUseShape exactly)",
            "pa": "PGE", "target_sector": "Residential", "climate_zone": "ALL",
            "hour_of_year": np.arange(1, 8761), "impact": 1.0 / 8760, "represents_year": 2018})
    elif kind == "profiles-gas":
        df = pd.DataFrame({"Profile": "EXAMPLE_GAS_PROFILE (replace: use the measure's E3GasSavProfile exactly)", "Month": np.arange(1, 13), "Impact": 1.0 / 12})
    elif kind == "acc-electric":
        df = _electric_template(acc_key)
    elif kind == "acc-gas":
        df = _gas_template(acc_key)
    elif kind == "run-parameters":
        df = run_parameters_template(acc_key)
    elif kind == "sw-factors":
        df = sw_factors_template()
    elif kind == "programs":
        df = programs_template()
    elif kind == "workbook":
        return write_workbook_template(out, acc_key)
    else:
        out.write_text("|".join(_CET_MEASURE_HEADER) + "\r\n", encoding="utf-8")
        return out
    if out.suffix.lower() == ".parquet":
        df.to_parquet(out, index=False)
    else:
        df.to_csv(out, index=False)
    return out


def write_all(out_dir: str | Path, acc_key: str = "2026") -> list[Path]:
    """Every template plus INPUT_FORMATS.md into out_dir."""
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)
    written = [write_template("workbook", d / "template_job_workbook.xlsx", acc_key),
               write_template("run-parameters", d / "template_run_parameters.csv", acc_key),
               write_template("sw-factors", d / "template_sw_allocation_factors.csv"),
               write_template("programs", d / "template_programs.csv"),
               write_template("profiles-8760", d / "template_electric_8760_profiles.csv"),
               write_template("profiles-gas", d / "template_gas_monthly_profiles.csv"),
               write_template("acc-electric", d / "template_acc_electric.csv", acc_key),
               write_template("acc-gas", d / "template_acc_gas.csv", acc_key),
               write_template("measure", d / "template_Measure_header.csv")]
    doc = DOCS_DIR / "INPUT_FORMATS.md"
    if doc.exists():
        shutil.copy(doc, d / "INPUT_FORMATS.md")
        written.append(d / "INPUT_FORMATS.md")
    return written


def export_library_profile(end_use_name: str, pa: str, target_sector: str, out_path: str | Path) -> Path:
    """Write the CEDARS library's 8,760 profile for (end use, IOU, sector) -- all its climate zones -- as an editable long-format file. Edit it and supply it as a
    user 8,760 profile: because the keys match, it replaces the library profile."""
    from .pipeline import real_impact_profiles as rip
    pid, exact = rip.match_profile_id(end_use_name, overrides.norm_pa(pa) or pa, target_sector)
    if pid is None:
        raise ValueError(f"no library profile for end use {end_use_name!r}, IOU {pa!r}, sector {target_sector!r}")
    meta = rip.profile_meta(pid)
    rows = []
    for zone, shape in rip._profile_shapes_by_zone(pid).items():
        rows.append(pd.DataFrame({"end_use_name": end_use_name, "pa": overrides.norm_pa(pa) or pa, "target_sector": target_sector, "climate_zone": zone,
                                  "hour_of_year": np.arange(1, 8761), "impact": shape, "represents_year": meta["represents_year"]}))
    df = pd.concat(rows, ignore_index=True)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(out, index=False) if out.suffix.lower() == ".parquet" else df.to_csv(out, index=False)
    return out
