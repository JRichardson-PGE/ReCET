"""Writing an OpenBCA Program Input workbook (the job), laid out as OpenBCA's template so its own parser reads it unchanged.

Only the sheets the parser reads are written: Program Inputs (header on row 3), Measure Inputs (header on row 4; columns C:G are the five free labels),
Electric Load Shapes (title row, then "Hour" and one column per load shape, 8,760 rows) and, when given, Natural Gas Load Shapes. Values only, no
formulas: OpenBCA reads cached values, which a workbook written here would not have.
"""
from __future__ import annotations

from pathlib import Path
from typing import Mapping

import numpy as np
import pandas as pd

# OpenBCA v1.0.4's Measure Inputs header, as in its template (spacing included: the parser's clean_header turns these into its column names).
MEASURE_HEADERS = [
    "Unique ID", "Program Name", "Label 1", "Label 2", "Label 3", "Label 4", "Label 5", "Measure ID", "Project ID", "Measure Name",
    "Avoided Cost Subset ", "Start Year", "Start Quarter", "Measure-Specific Discount Rate", "Measure Unit", "Unit Quantity",
    "Estimated Useful Life (Years)", "Net-to-gross Ratio", "Admin Cost Upfront ($)", "Admin Cost Annual ($/Year)",
    "Utility Direct Investment in DERs Upfront($)", "Avoided Credit and Collection Cost Annual ($/Year)", "Utility Incentive Upfront  ($)",
    "Utility Incentive Annual ($/Year)", "Incremental Cost Upfront  ($)", "Incremental Cost Annual ($/Year)",
    "Host Customer Transaction Cost Upfront ($)", "Host Customer Interconnection Cost Upfront ($)", "Host Customer Tax Incentive Upfront ($)",
    "Electric Savings Load Shape", "Adj. Gross Annual Electric Savings (kWh)", "Adj. Gross Coincident Peak Savings (kW)",
    "Natural Gas Savings Load Shape", "Adj. Gross Annual Natural Gas Savings (MMBTU)", "Adj. Gross Annual Propane Savings (MMBTU)",
    "Adj. Gross Annual Oil Savings (MMBTU)", "Adj. Gross Annual Diesel Savings (MMBTU)", "Adj. Gross Annual Wood Savings (MMBTU)",
    "Host Customer Non-Energy Impacts Upfront ($)", "Host Customer Non-Energy Impacts - Low-Income Upfront ($)", "Change in Host Customer Resilience",
    "Change in Societal Resilience", "CUSTOM 1 Value Stream Name", "CUSTOM 1 Annual Savings", "CUSTOM 2 Value Stream Name", "CUSTOM 2 Annual Savings",
    "CUSTOM 3 Value Stream Name", "CUSTOM 3 Annual Savings", "CUSTOM 4 Value Stream Name", "CUSTOM 4 Annual Savings", "CUSTOM 5 Value Stream Name",
    "CUSTOM 5 Annual Savings",
]
PROGRAM_HEADERS = ["Program Name", "Program Year", "Program Admin Costs ($)", "Program Utility Incentive ($)",
                   "Program Performance Incentive to Utility ($)", "Program Federal Incentive ($)"]


def _cell(v):
    if v is None or (isinstance(v, float) and np.isnan(v)) or v is pd.NA:
        return None
    return v.item() if isinstance(v, np.generic) else v


def write_program_input(path: str | Path, measures: pd.DataFrame, programs: pd.DataFrame | None, electric_shapes: Mapping[str, np.ndarray],
                        gas_shapes: Mapping[str, np.ndarray] | None = None) -> Path:
    """`measures` / `programs` have MEASURE_HEADERS / PROGRAM_HEADERS columns (missing ones are left blank); each shape is 8,760 hourly values
    (electric) or 12 monthly values (gas), normalized or not as OpenBCA allows."""
    from openpyxl import Workbook
    unknown = set(measures.columns) - set(MEASURE_HEADERS)
    if unknown:
        raise ValueError(f"not OpenBCA Measure Inputs columns: {sorted(unknown)}")
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Program Inputs")
    ws.append(["Step 1."]), ws.append([])
    ws.append(PROGRAM_HEADERS)
    for row in (programs.reindex(columns=PROGRAM_HEADERS).itertuples(index=False) if programs is not None else []):
        ws.append([_cell(v) for v in row])
    ws = wb.create_sheet("Measure Inputs")
    ws.append(["Step 2."]), ws.append([]), ws.append([])
    ws.append(MEASURE_HEADERS)
    for row in measures.reindex(columns=MEASURE_HEADERS).itertuples(index=False):
        ws.append([_cell(v) for v in row])
    names = list(electric_shapes)
    ws = wb.create_sheet("Electric Load Shapes")
    ws.append(["Loadshape Mapping - Time Series Hourly"])
    ws.append(["Hour", *names])
    cols = [np.asarray(electric_shapes[n], dtype=float) for n in names]
    if any(len(c) != 8760 for c in cols):
        raise ValueError("an electric load shape without 8,760 values")
    for h in range(8760):
        ws.append([h + 1, *(float(c[h]) for c in cols)])
    if gas_shapes:
        ws = wb.create_sheet("Natural Gas Load Shapes")
        ws.append(["Loadshape Mapping - Time Series Monthly"])
        gnames = list(gas_shapes)
        ws.append(["Month", *gnames])
        for m in range(12):
            ws.append([m + 1, *(float(np.asarray(gas_shapes[n])[m]) for n in gnames)])
    path = Path(path)
    wb.save(path)
    return path
