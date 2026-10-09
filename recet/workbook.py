"""A job as one Excel workbook instead of separate files.

Sheets (names are matched ignoring case, spaces and underscores; any other sheet, such as `Instructions`, is ignored):

  Measure         required   the CET Measure.csv columns, one measure per row
  ProgramCost     optional   the CET ProgramCost.csv columns
  RunParameters   optional   run parameters (discount rates, base year, gas CO2) per PA        -> overrides.load_run_parameters
  SWAllocation    optional   statewide allocation factors                                      -> overrides.load_sw_factors
  GasProfiles     optional   gas monthly profiles (Profile, Month, Impact)                     -> overrides.load_user_gas_profiles
  Programs        optional   program metadata (CEDARS Programs list layout)                    -> programs.load_programs
  Profiles8760    optional   8,760 electric profiles (long format; Excel holds about 119 profile x zone blocks at most)   -> overrides.load_user_profiles

A sheet with a header row but no data rows counts as not given. The ACC files are not workbook sheets (an hourly electric ACC is millions of rows): supply them as
separate files. `extract()` writes each sheet that has data as the file the rest of the program already reads (Measure and ProgramCost as CET's pipe-delimited
text, the others as CSV), so a workbook run and a run on the equivalent files are the same run.
"""
from __future__ import annotations

import datetime as dt
import math
import re
from pathlib import Path

import pandas as pd

SHEETS = {  # canonical sheet name -> (RunConfig field, file written, kind)
    "Measure": ("measures", "Measure.csv", "cet"),
    "ProgramCost": ("programcost", "ProgramCost.csv", "cet"),
    "RunParameters": ("run_parameters_file", "run_parameters.csv", "csv"),
    "SWAllocation": ("sw_factors_file", "sw_allocation_factors.csv", "csv"),
    "GasProfiles": ("profiles_gas_file", "gas_monthly_profiles.csv", "csv"),
    "Programs": ("programs_file", "programs.csv", "csv"),
    "Profiles8760": ("profiles_8760_file", "electric_8760_profiles.csv", "csv"),
}
_KEY = {re.sub(r"[\s_]", "", k).lower(): k for k in SHEETS}


class WorkbookError(ValueError):
    pass


def _text(v) -> str:
    """A cell as CET text: blank for empty, whole numbers without '.0' (so a zone 13 stays '13'), other numbers in their shortest exact form."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    return str(v)                      # text exactly as typed (the CSV route keeps trailing spaces too)


def _sheet_frame(ws) -> pd.DataFrame | None:
    rows = ws.iter_rows(values_only=True)
    header = None
    for r in rows:
        if r is not None and any(c is not None and str(c).strip() for c in r):
            header = [_text(c).strip() for c in r]
            break
    if header is None:
        return None
    while header and not header[-1]:
        header.pop()
    if not header:
        return None
    if any(not h for h in header):
        raise WorkbookError(f"sheet {ws.title!r}: blank column name in the header row (column {header.index('') + 1})")
    data = []
    for r in rows:
        vals = [_text(c) for c in (r or ())][:len(header)]
        vals += [""] * (len(header) - len(vals))
        if any(vals):
            data.append(vals)
    if not data:
        return None
    return pd.DataFrame(data, columns=header)


def sheet_names(path: str | Path) -> dict[str, str]:
    """{canonical name: the workbook's own sheet name} for the sheets this program reads."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return {_KEY[re.sub(r"[\s_]", "", n).lower()]: n for n in wb.sheetnames if re.sub(r"[\s_]", "", n).lower() in _KEY}
    finally:
        wb.close()


def extract(path: str | Path, out_dir: str | Path) -> dict[str, Path]:
    """Write every data sheet of the workbook as its input file in `out_dir`. Returns {RunConfig field: file}. Raises WorkbookError if there is no Measure data."""
    from openpyxl import load_workbook
    p = Path(path)
    if not p.exists():
        raise WorkbookError(f"workbook not found: {p}")
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(p, read_only=True, data_only=True)       # data_only: a formula cell gives its last calculated value
    files: dict[str, Path] = {}
    try:
        names = {re.sub(r"[\s_]", "", n).lower(): n for n in wb.sheetnames}
        for key, canon in _KEY.items():
            if key not in names:
                continue
            df = _sheet_frame(wb[names[key]])
            if df is None:
                continue
            field, fname, kind = SHEETS[canon]
            target = out / fname
            if kind == "cet":
                bad = [c for c in df.columns if df[c].str.contains("|", regex=False).any()] + [c for c in df.columns if "|" in c]
                if bad:
                    raise WorkbookError(f"sheet {names[key]!r}: CET files are '|'-delimited, so no cell may contain '|' (column(s) {sorted(set(bad))})")
                lines = ["|".join(df.columns)] + ["|".join(r) for r in df.itertuples(index=False, name=None)]
                target.write_text("\n".join(lines) + "\n", encoding="utf-8")
            else:
                df.to_csv(target, index=False)
            files[field] = target
    finally:
        wb.close()
    if "measures" not in files:
        raise WorkbookError(f"{p.name}: no Measure sheet with data (sheets found: {wb.sheetnames})")
    return files
