"""When a queued run will start and finish (BRD BR-06.2: estimated completion for every run, a queued one included).

A queued run has not reported any progress, so its time is estimated from its size: the measure rows of its Measure.csv (or its workbook's Measure sheet),
counted once and kept in its folder as `estimate.json`. The seconds per million measures come from the stage seconds the progress bar is weighted by
(recet/progress.py: the 1,000,000-measure server run, for the run's output format and its SCT / emissions options); a comparison is two runs, a preflight-only
run only its input check. That figure is scaled by how this server's recent runs compared with their own estimates (`calibration`), so a faster or slower
machine corrects itself after a few runs.

`queue` then walks the queue in the order the workers claim it (input checks first, then the oldest run): each worker is free when its running run's time left
runs out (the run's own progress estimate, else its size estimate less the time it has run), and each queued run starts on the first free worker. Runs
queued later can still go ahead of it (an input check), so the times are an estimate as of now.
"""
from __future__ import annotations

import json
import statistics
import time
from pathlib import Path

ESTIMATE_FILE = "estimate.json"
OVERHEAD_SECONDS = 10.0            # a run's fixed cost (starting Python, loading the ACC) whatever its size
UNKNOWN_SECONDS = 60.0             # a run whose size cannot be read
CALIBRATION_RUNS = 10              # the recent finished runs the calibration looks at
CALIBRATION_MIN_MEASURES = 20_000  # smaller runs are mostly overhead: they say little about the per-measure rate
CALIBRATION_RANGE = (0.25, 4.0)


def _count_measures(cfg: dict) -> int | None:
    from ..pipeline import delimited
    m = cfg.get("measures")
    if m and Path(m).is_file() and delimited.detect_sep(m):
        return delimited.count_rows(m)
    w = cfg.get("workbook")
    if w and Path(w).is_file():
        try:
            from openpyxl import load_workbook
            wb = load_workbook(w, read_only=True)
            ws = next((wb[n] for n in wb.sheetnames if n.strip().lower() == "measure"), None)
            n = max((ws.max_row or 1) - 1, 0) if ws is not None else None
            wb.close()
            return n
        except Exception:
            return None
    return None


def _saved(folder: Path) -> dict | None:
    try:
        return json.loads((folder / ESTIMATE_FILE).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def known_size(folder: Path) -> dict | None:
    """The size kept in estimate.json, without counting anything: the runs list shows it for many runs at once, and a run queued before estimates were
    kept would otherwise have its whole Measure file counted on a page view."""
    return _saved(folder)


def size(folder: Path) -> dict | None:
    """The run's size and the options that change its time: {"measures", "preflight_only", "compare", "include_societal", "emissions", "output_format", "filing"},
    read from its config.json and inputs once and kept in estimate.json. None when the run has no config."""
    p = folder / ESTIMATE_FILE
    saved = _saved(folder)
    if saved is not None:
        return saved
    try:
        cfg = json.loads((folder / "config.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    out = {"measures": _count_measures(cfg), "preflight_only": bool(cfg.get("preflight_only")), "compare": bool(cfg.get("compare_with")),
           "include_societal": cfg.get("include_societal", True), "emissions": cfg.get("emissions", True), "output_format": cfg.get("output_format") or "csv.zip",
           "filing": bool(cfg.get("filing_validation"))}
    try:
        p.write_text(json.dumps(out), encoding="utf-8")
    except OSError:
        pass
    return out


def model_seconds(sz: dict | None) -> float:
    """The run's time from its size alone (before calibration)."""
    from .. import progress
    if not sz or sz.get("measures") is None:
        return UNKNOWN_SECONDS
    secs = dict(progress.SECONDS_IN_PARTS)
    secs["file"] = progress.FILE_SECONDS.get(sz.get("output_format") or "csv.zip", progress.FILE_SECONDS["csv.zip"])
    f_price, f_write = progress.OPTION_FACTORS[(bool(sz.get("include_societal", True)), bool(sz.get("emissions", True)))]
    secs["pricing"] *= f_price
    secs["writing"] *= f_write
    if sz.get("filing"):                            # CEDARS filing checks lengthen preflight
        secs["preflight"] += progress.FILING_PREFLIGHT_SECONDS
    if sz.get("preflight_only"):
        secs = {k: v for k, v in secs.items() if k in ("preparing", "splitting", "preflight")}
    per_million = sum(secs.values())
    one = OVERHEAD_SECONDS + per_million * sz["measures"] / 1e6
    return 2 * one if sz.get("compare") and not sz.get("preflight_only") else one


def calibration(db, runs_dir: Path) -> float:
    """Median of actual / estimated time over this server's recent finished full runs of some size (1.0 when there are none)."""
    rows = db.all("SELECT id, started_at, finished_at FROM runs WHERE status = 'finished' AND preflight_only = 0 AND started_at IS NOT NULL "
                  "AND finished_at IS NOT NULL ORDER BY finished_at DESC LIMIT 50")
    ratios = []
    for r in rows:
        sz = _saved(runs_dir / r["id"])                 # only runs whose size was kept: never count an old run's rows here
        if not sz or (sz.get("measures") or 0) < CALIBRATION_MIN_MEASURES:
            continue
        ratios.append((r["finished_at"] - r["started_at"]) / model_seconds(sz))
        if len(ratios) >= CALIBRATION_RUNS:
            break
    if not ratios:
        return 1.0
    lo, hi = CALIBRATION_RANGE
    return min(hi, max(lo, statistics.median(ratios)))


def queue(db, runs_dir: Path, workers: int, running_left: dict[str, float | None], now: float | None = None) -> dict[str, dict]:
    """{queued run id: {"starts_in", "finishes_in", "seconds"}} in seconds from now. `running_left` gives each running run's time left from its progress
    report (None when it has not reported one yet)."""
    now = time.time() if now is None else now
    k = calibration(db, runs_dir)
    slots: list[float] = []
    for r in db.all("SELECT id, started_at FROM runs WHERE status = 'running' ORDER BY started_at"):
        left = running_left.get(r["id"])
        if left is None:
            left = k * model_seconds(size(runs_dir / r["id"])) - (now - (r["started_at"] or now))
        slots.append(max(float(left), 0.0))
    slots = sorted(slots)[:max(workers, 1)]
    slots += [0.0] * (max(workers, 1) - len(slots))
    out = {}
    for r in db.all("SELECT id FROM runs WHERE status = 'queued' ORDER BY preflight_only DESC, created_at"):
        secs = k * model_seconds(size(runs_dir / r["id"]))
        i = min(range(len(slots)), key=slots.__getitem__)
        start = slots[i]
        slots[i] = start + secs
        out[r["id"]] = {"starts_in": round(start), "finishes_in": round(start + secs), "seconds": round(secs)}
    return out


def _span(s: float) -> str:
    return f"about {max(1, round(s / 60))} min" if s >= 90 else f"about {max(5, round(s / 5) * 5)} s"


def describe(e: dict | None) -> str:
    """'starts in about 6 min, done in about 13 min (estimated)' for the run page and the runs list."""
    if not e:
        return ""
    start = "starts next" if e["starts_in"] < 5 else f"starts in {_span(e['starts_in'])}"
    return f"{start}, done in {_span(e['finishes_in'])} (estimated)"
