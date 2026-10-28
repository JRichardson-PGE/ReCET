"""Building and writing a run's result files.

measure-level table -- one row per measure, column groups left to right:
  Used_ACCVersion                  which ACC version the run priced against
  <raw Measure.csv columns>        exactly as in the input file (text preserved)
  Used_Year ... Scaled_Profile_Flag   the few derived values a reader cannot see in the inputs (DERIVED_OUTPUTS): the claim year and quarter used, the annual
                                   net kWh / therms each baseline is priced on (with installation and realization rates), the fuel-substitution flag and
                                   whether UnitkWh came from a scaled (kWh-per-unit) profile
  ElecPricingStatus, ProfileSource, GasPricingStatus, GasProfileSource, WaterPricingStatus
                                   how each fuel was priced, from which profile, or that it was UNPRICED (a 0 benefit alone cannot tell "no savings" from "could not price")
  <outputs>                        benefits (electric, gas, water energy, other), supply costs, participant / bill-impact PVs, Gross and NoAdmin variants, TRC/PAC/RIM/SCB/SCH costs and ratios,
                                   societal (_SB/_SH) parts, program-cost weighting, CO2/NOx/PM10 (net/gross, first-year and lifecycle)
  ElecValue_* / GasValue_*         (components mode) each ACC component's discounted lifetime value; they sum to ElecBen - ElecSupplyCost / GasBen - GasSupplyCost
  CreditedPA                       the PA the measure is credited to: the `PA` of its program in the ProgramCost input (`SW` for an unallocated statewide program;
                                   with statewide allocation, the IOU each allocated share went to); blank without a ProgramCost input
  Prg_MatchedOn, Prg_<column>      the program's metadata from the CEDARS Programs list (programs.py), matched on the measure's PrgID as priced

The per-record checks (CEDARS filing checks, the ACC check) are not columns of this table: they are in the preflight findings and filing_validation files.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd

from .pipeline.delimited import read_cet_csv

# the loader's derived columns that are written (the loader names them as written); its other columns restate an input and are not written
DERIVED_OUTPUTS = ["Used_Year", "Used_Quarter", "AC_Net_kWh_1st_Baseline", "AC_Net_kWh_2nd_Baseline", "AC_Net_Therms_1st_Baseline", "AC_Net_Therms_2nd_Baseline",
                   "AC_Net_kWh_Water_1st_Baseline", "AC_Net_kWh_Water_2nd_Baseline", "Fuel_Sub_Flag", "Scaled_Profile_Flag"]
DROP_FROM_RESULT = {"CEInputID", "PrgID", "NTGRkWh_plus_MEBens", "NTGRTherm_plus_MEBens"}   # restate an input, or a working value of the cost tests

# the columns compared between two runs / against CET, and the totals in the summary
KEY_OUTPUTS = ["ElecBen", "GasBen", "WaterEnergyBen", "OtherBen", "TotalSystemBenefit", "ElecSupplyCost", "WaterEnergyCost", "TRCCost", "PACCost", "RIMCost", "SCBCost", "SCHCost",
               "TRCBen", "PACBen", "RIMBen", "SCBBen", "SCHBen", "TRCRatio", "PACRatio", "RIMRatio", "SCBRatio", "SCHRatio", "PCTBen", "PCTCost", "PCTRatio", "FirstYearNetElecCO2", "FirstYearNetGasCO2", "LifecycleBillReducElec", "LifecycleBillReducGas"]
SUMMARY_TOTALS = ["ElecBen", "GasBen", "WaterEnergyBen", "OtherBen", "ElecSupplyCost", "GasSupplyCost", "WaterEnergyCost", "TotalSystemBenefit", "TRCCost", "PACCost", "RIMCost", "SCBCost", "SCHCost",
                  "TRCCostNoAdmin", "PACCostNoAdmin", "WeightedProgramCost", "LifecycleBillReducElec", "LifecycleBillReducGas", "PCTBen", "PCTCost", "FirstYearNetElecCO2", "FirstYearNetGasCO2"]


def read_raw(measure_csv: str | Path, allocate_statewide: bool = False) -> pd.DataFrame:
    """The Measure.csv columns as text (blank stays blank); with statewide allocation, the allocated rows (numbers are then written as computed)."""
    from .templates import with_defaults
    p = Path(measure_csv)
    if not allocate_statewide:
        return with_defaults(read_cet_csv(p, dtype=str, keep_default_na=False, low_memory=False))
    from .pipeline.statewide_allocation import allocate_measure_inputs
    df, _report = allocate_measure_inputs(with_defaults(read_cet_csv(p, low_memory=False)))
    return df.astype(object).where(df.notna(), "").astype(str)


def credited_pa(measures: pd.DataFrame, programcost_csv, allocate_statewide: bool = False) -> pd.Series:
    """Per measure (aligned with `measures`), the `PA` of its program in the ProgramCost input, exactly as filed -- so an unallocated statewide program shows `SW`.
    With statewide allocation the ProgramCost rows are allocated the same way the measures are, so a statewide program's share shows the IOU it was allocated to
    (its program row's PA is that IOU) and every other program shows its filed PA. Matched on (PrgID, claim year = PrgYear), else on PrgID alone; blank when there
    is no ProgramCost input or no row for the program."""
    blank = pd.Series("", index=measures.index, dtype=object)
    if not programcost_csv:
        return blank
    if allocate_statewide:                          # depends on the allocation factors too (a run can supply its own): not cached
        by_year, by_prg = _credited_maps.__wrapped__(str(programcost_csv), True, 0, 0)
    else:
        st = Path(programcost_csv).stat()
        by_year, by_prg = _credited_maps(str(programcost_csv), False, st.st_size, st.st_mtime_ns)
    vals = [by_year.get((p, int(y)), by_prg.get(p, "")) for p, y in zip(measures["PrgID"].astype(str).str.strip(), measures["Used_Year"])]
    return pd.Series(vals, index=measures.index, dtype=object)


@lru_cache(maxsize=4)
def _credited_maps(path: str, allocate_statewide: bool, size: int, mtime_ns: int) -> tuple[dict, dict]:
    """credited_pa's (PrgID, PrgYear) -> PA and PrgID -> PA maps of one ProgramCost file, read once per file and version of it (a run asks once per part).
    The maps are shared: do not change them."""
    pc = read_cet_csv(path, dtype=str, keep_default_na=False)
    if allocate_statewide:
        from .pipeline.statewide_allocation import allocate_program_costs
        pc, _ = allocate_program_costs(pc)
    pc = pc.assign(PrgID=pc["PrgID"].astype(str).str.strip(), PA=pc["PA"].astype(str).str.strip(),
                   PrgYear=pd.to_numeric(pc["PrgYear"], errors="coerce").astype("Int64"))
    by_year = {(p, int(y)): a for p, y, a in zip(pc["PrgID"], pc["PrgYear"], pc["PA"]) if pd.notna(y)}
    by_prg = pc.drop_duplicates("PrgID").set_index("PrgID")["PA"].to_dict()
    return by_year, by_prg


def _aligned_raw(raw: pd.DataFrame, measures: pd.DataFrame) -> pd.DataFrame:
    if not raw["CEInputID"].is_unique:
        raise ValueError("CEInputID is not unique")
    raw = raw.set_index("CEInputID", drop=False).reindex(measures["CEInputID"]).reset_index(drop=True)
    if raw["CEInputID"].isna().any() or not (raw["CEInputID"].to_numpy() == measures["CEInputID"].to_numpy()).all():
        raise ValueError("raw Measure.csv rows and loaded measures do not line up on CEInputID")
    return raw


def assemble_measure_level(raw: pd.DataFrame, measures: pd.DataFrame, result: pd.DataFrame, acc_label: str | pd.Series,
                           credited: pd.Series | None = None) -> pd.DataFrame:
    """One row per measure, in the loader's order. `raw` may have more rows than `measures` (a subset run); it is aligned on CEInputID."""
    raw = _aligned_raw(raw, measures)
    result = result.set_index("CEInputID").reindex(measures["CEInputID"]).reset_index()
    derived = measures[DERIVED_OUTPUTS].reset_index(drop=True)
    out = result.drop(columns=[c for c in DROP_FROM_RESULT if c in result]).reset_index(drop=True)
    frame = pd.concat([raw.reset_index(drop=True), derived, out], axis=1).copy()   # copy: one block, so the two columns added below do not warn
    frame.insert(0, "Used_ACCVersion", acc_label if isinstance(acc_label, str) else pd.Series(acc_label).to_numpy())   # one label, or one per measure (mixed ACC versions)
    frame["CreditedPA"] = credited.to_numpy() if credited is not None else ""
    return frame


def attach_programs(frame: pd.DataFrame, prg_ids) -> pd.DataFrame:
    """The program metadata columns (programs.lookup) appended to the measure-level table; `prg_ids` is the measures' PrgID as priced, in the frame's order."""
    from . import programs
    meta = programs.lookup(prg_ids)
    meta.index = frame.index
    return pd.concat([frame, meta], axis=1)


def filing_table(filing: pd.DataFrame, rule_set: str) -> pd.DataFrame:
    """The filing_validation table a run writes: every record that fails a CEDARS upload check, with the rule texts."""
    from . import filing_rules
    rs = filing_rules.rule_set(rule_set)
    text = {r.id: r.text for r in rs.rules}
    out = filing.rename(columns={"id": "record"}).copy()
    out["rules_failed"] = out["FilingRules"].map(lambda ids: " | ".join(f"{i}: {text.get(i, '')}" for i in ids.split(", ") if i))
    out.insert(0, "rule_set", rs.label)
    return out


EXCLUDE_FLAG = "Prg_Exclude_From_CE"


SUMMARY_RATIOS = ("TRC_ratio", "PAC_ratio", "RIM_ratio", "SCB_ratio", "SCH_ratio", "PCT_ratio", "TRC_NoAdmin_ratio", "PAC_NoAdmin_ratio")


RATIO_TESTS = ("TRC", "PAC", "RIM", "SCB", "SCH")     # each writes its ratios' numerator as <test>Ben (PCT's is PCTBen, an output of its own)


def _numerators(get) -> dict[str, object]:
    """{test: the numerator of its ratios} for RATIO_TESTS: the total of the measure-level <test>Ben column, or (a table written before 0.1.53, which has no
    such column) the same sum of benefits built from its parts."""
    z = lambda x: 0.0 if x is None else x

    def benefits(suffix: str = ""):
        parts = [get(f"{c}{suffix}") for c in ("ElecBen", "GasBen", "WaterEnergyBen")] + [get("OtherBen")]
        return None if any(p is None for p in parts) else sum(parts)

    b, b_sb, b_sh = benefits(), benefits("_SB"), benefits("_SH")
    tax, incr = z(get("TaxCredits")), z(get("LifecycleBillIncrElec")) + z(get("LifecycleBillIncrGas"))
    plus = lambda x, y: None if x is None else x + y
    built = {"TRC": plus(b, tax), "PAC": b, "RIM": plus(b, incr), "SCB": b_sb, "SCH": b_sh}
    out = {}
    for t in RATIO_TESTS:
        v = get(f"{t}Ben")
        out[t] = built[t] if v is None else v
    return out


def _ratio_parts(get) -> list[tuple[str, object, object]]:
    """(ratio name, numerator, denominator) for the cost-test ratios, with `get(column)` returning a total (a number, or a Series for per-row ratios) or None."""
    n = _numerators(get)
    return [*((t, n[t], get(f"{t}Cost")) for t in RATIO_TESTS), *((f"{t}_NoAdmin", n[t], get(f"{t}CostNoAdmin")) for t in RATIO_TESTS),
            ("PCT", get("PCTBen"), get("PCTCost"))]


def group_ratios(g: pd.DataFrame) -> dict[str, float]:
    """The cost-test ratios of a group of measures, as CET reports them on a run receipt: the summed BENEFITS (supply costs are already in the costs, not netted
    from the numerator) over the summed cost; TRC adds tax credits and RIM the bill increases to the numerator; PCT is PCTBen over PCTCost. A ratio whose inputs are missing or whose cost is 0
    is left out. (Until 2026-09-21 the numerator was TotalSystemBenefit, which nets the supply costs out of the numerator as well as counting them in the cost.)
    Used by the portfolio summary and the Results tab's views, so both give the same ratios."""
    get = lambda c: float(g[c].sum()) if c in g and g[c].notna().any() else None
    return {f"{name}_ratio": num / den for name, num, den in _ratio_parts(get) if num is not None and den}


def group_numerators(g: pd.DataFrame) -> dict[str, float]:
    """The numerators of a group's ratios (group_ratios), under the measure-level names TRCBen ... SCHBen; one whose inputs are missing is left out. So each
    <test>_ratio of a summary row is its <test>Ben over its <test>Cost (or <test>CostNoAdmin)."""
    get = lambda c: float(g[c].sum()) if c in g and g[c].notna().any() else None
    return {f"{t}Ben": v for t, v in _numerators(get).items() if v is not None}


def row_ratios(df: pd.DataFrame) -> pd.DataFrame:
    """The same ratios for each row on its own (one measure each), vectorised; NaN where the cost is 0 or an input is missing."""
    get = lambda c: df[c].astype(float) if c in df and df[c].notna().any() else None
    out = pd.DataFrame(index=df.index)
    for name, num, den in _ratio_parts(get):
        if num is not None and den is not None:
            out[f"{name}_ratio"] = num / den.where(den != 0)
    return out


def summarize(frame: pd.DataFrame, apply_exclude_from_ce: bool = False) -> pd.DataFrame:
    """Portfolio totals and ratios, overall and by IOU territory (PA), by credited PA and by program segment. With `apply_exclude_from_ce`, measures of programs
    whose Exclude_From_CE flag is true are left out of every row and shown in one last row of their own (group `excluded`)."""
    def block(g: pd.DataFrame, label: str, group: str) -> dict:
        return _summary_row(g, label, group, len(g))
    excluded = None
    if apply_exclude_from_ce and EXCLUDE_FLAG in frame:
        from . import programs
        flag = programs.is_true(frame[EXCLUDE_FLAG]).to_numpy()
        excluded, frame = frame[flag], frame[~flag]
    rows = [block(frame, "all measures", "portfolio")]
    for col, group in (("IOU_AC_Territory", "IOU territory (PA)"), ("CreditedPA", "credited PA"), ("Prg_ProgramSegment", "program segment")):
        if col in frame and not (col != "IOU_AC_Territory" and (frame[col].astype(str) == "").all()):
            for key, g in frame.groupby(frame[col].astype(str).replace("", "(blank)"), sort=True):
                rows.append(block(g, str(key), group))
    if excluded is not None:
        rows.append(block(excluded, "Exclude_From_CE programs (left out of the rows above)", "excluded"))
    return pd.DataFrame(rows)


class SummaryParts:
    """summarize() over a table that arrives in parts (a chunked run): each part is reduced to per-group sums as it comes, so the whole table is never held.
    `result()` gives the rows summarize(whole table) gives, to rounding (the sums are taken in a different order)."""
    GROUPS = (("IOU_AC_Territory", "IOU territory (PA)"), ("CreditedPA", "credited PA"), ("Prg_ProgramSegment", "program segment"))
    EXCLUDED = ("excluded", "Exclude_From_CE programs (left out of the rows above)")

    def __init__(self, apply_exclude_from_ce: bool = False):
        self.apply_exclude = apply_exclude_from_ce
        self.sums: dict[tuple[str, str], pd.DataFrame] = {}      # (group, key) -> one row of column sums (NaN where the column was NaN throughout)
        self.counts: dict[tuple[str, str], int] = {}
        self.present: dict[str, bool] = {}                       # group column -> in the table
        self.nonblank: dict[str, bool] = {}                      # group column -> some row is not blank
        self.cols: list[str] | None = None

    def _add(self, key: tuple[str, str], sums, n: int) -> None:
        part = pd.DataFrame([np.asarray(sums, dtype=float)], columns=self.cols)
        prev = self.sums.get(key)
        self.sums[key] = part if prev is None else prev.add(part, fill_value=0.0)      # NaN only where both are NaN
        self.counts[key] = self.counts.get(key, 0) + n

    def add(self, frame: pd.DataFrame) -> None:
        if self.cols is None:
            self.cols = [c for c in frame.columns if c in summary_input_columns()]
        excluded = None
        if self.apply_exclude and EXCLUDE_FLAG in frame:
            from . import programs
            flag = programs.is_true(frame[EXCLUDE_FLAG]).to_numpy()
            excluded, frame = frame[flag], frame[~flag]
        x = frame[self.cols].astype(float)                       # the summed columns once; each grouping is one vectorised groupby (it looped over the groups)
        self._add(("portfolio", "all measures"), x.sum(min_count=1), len(x))
        for col, group in self.GROUPS:
            self.present.setdefault(col, col in frame)
            if col in frame:
                keys = frame[col].astype(str)
                self.nonblank[col] = self.nonblank.get(col, False) or bool((keys != "").any())
                k = keys.replace("", "(blank)").to_numpy()
                sums = x.groupby(k, sort=False).sum(min_count=1)
                counts = pd.Series(k).value_counts(sort=False)
                for key in sums.index:
                    self._add((group, str(key)), sums.loc[key].to_numpy(), int(counts[key]))
        if excluded is not None and len(excluded):
            self._add(self.EXCLUDED, excluded[self.cols].astype(float).sum(min_count=1), len(excluded))
        elif excluded is not None:
            self.counts.setdefault(self.EXCLUDED, 0)

    def result(self) -> pd.DataFrame:
        def block(key):
            g = self.sums.get(key, pd.DataFrame(columns=self.cols or []))
            return _summary_row(g, key[1], key[0], self.counts[key])
        rows = [block(("portfolio", "all measures"))]
        for col, group in self.GROUPS:
            if self.present.get(col) and (col == "IOU_AC_Territory" or self.nonblank.get(col)):
                rows += [block(k) for k in sorted(k for k in self.sums if k[0] == group)]
        if self.EXCLUDED in self.counts:
            rows.append(block(self.EXCLUDED))
        return pd.DataFrame(rows)


def summary_input_columns() -> set[str]:
    """Every column summarize() reads: the totals and each ratio's parts."""
    used: set[str] = set()

    def get(c):
        used.add(c)
        return 1.0
    _ratio_parts(get)
    return used | set(SUMMARY_TOTALS)


def _summary_row(g: pd.DataFrame, label: str, group: str, n: int) -> dict:
    row = {"group": group, "value": label, "measures": n}
    for c in SUMMARY_TOTALS:
        if c in g and g[c].notna().any():
            row[c] = float(g[c].sum())
    row.update(group_numerators(g))
    for name, value in group_ratios(g).items():
        if name in SUMMARY_RATIOS:
            row[name] = value
    return row


def compare_frames(a: pd.DataFrame, b: pd.DataFrame, label_a: str, label_b: str, key: str = "CEInputID") -> tuple[pd.DataFrame, pd.DataFrame]:
    """(per-measure comparison, portfolio summary) of two runs of the same job. Columns: <col>_<label_a>, <col>_<label_b>, <col>_delta, <col>_pct."""
    cols = [c for c in KEY_OUTPUTS if c in a and c in b]
    j = a[[key, "PrgID", "MeasDescription", "IOU_AC_Territory", "ElecPricingStatus", "GasPricingStatus"] + cols].merge(
        b[[key, "ElecPricingStatus", "GasPricingStatus"] + cols], on=key, suffixes=(f"_{label_a}", f"_{label_b}"), how="outer")
    out = j[[key, "PrgID", "MeasDescription", "IOU_AC_Territory"]].copy()
    for c in cols:
        va, vb = j[f"{c}_{label_a}"], j[f"{c}_{label_b}"]
        out[f"{c}_{label_a}"], out[f"{c}_{label_b}"] = va, vb
        out[f"{c}_delta"] = vb - va
        out[f"{c}_pct"] = ((vb - va) / va.abs()).where(va != 0)
    out[f"ElecPricingStatus_{label_a}"], out[f"ElecPricingStatus_{label_b}"] = j[f"ElecPricingStatus_{label_a}"], j[f"ElecPricingStatus_{label_b}"]
    out[f"GasPricingStatus_{label_a}"], out[f"GasPricingStatus_{label_b}"] = j[f"GasPricingStatus_{label_a}"], j[f"GasPricingStatus_{label_b}"]
    rows = []
    for c in cols:
        if c.endswith("Ratio"):
            continue
        ta, tb = float(j[f"{c}_{label_a}"].sum()), float(j[f"{c}_{label_b}"].sum())
        rows.append({"metric": c, label_a: ta, label_b: tb, "delta": tb - ta, "pct": (tb - ta) / abs(ta) if ta else np.nan})
    return out, pd.DataFrame(rows)


CET_COMPARED = ("TotalSystemBenefit", "ElecBen", "GasBen", "WaterEnergyBen", "WaterEnergyBen_SB", "WaterEnergyBen_SH", "WaterEnergyBenGross", "OtherBen",
                "ElecSupplyCost", "GasSupplyCost", "WaterEnergyCost", "TRCCost", "PACCost", "RIMCost", "SCBCost", "SCHCost",
                "TRCRatio", "PACRatio", "RIMRatio", "FirstYearNetElecCO2", "FirstYearNetGasCO2", "LifecycleBillReducElec", "LifecycleBillReducGas")


def read_cet_outputs(cet_outputs: str | Path) -> pd.DataFrame:
    """The columns of a CET *_outputs.csv that compare_to_cet uses, under ReCET's names (CET_ID named CEInputID; CET's NetElecCO2, BillReducElec ...
    named as ReCET writes them, fieldnames.OUTPUT_TO_CET)."""
    from .fieldnames import OUTPUT_TO_CET, cet_outputs_to_recet
    cols = {"CET_ID", "CEInputID", *CET_COMPARED, *(OUTPUT_TO_CET.get(c, c) for c in CET_COMPARED)}
    cet = read_cet_csv(cet_outputs, low_memory=False, usecols=lambda c: str(c).strip() in cols).rename(columns={"CET_ID": "CEInputID"})
    return cet_outputs_to_recet(cet)


def compare_to_cet_part(frame: pd.DataFrame, cet: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, tuple[float, float]], int]:
    """(per-measure rows, {metric: (POC total, CET total)} for the summed metrics, measures matched) of `frame` against `cet` (read_cet_outputs)."""
    cols = [c for c in CET_COMPARED if c in cet and c in frame]
    j = frame[["CEInputID", "PrgID", "MeasDescription"] + cols].merge(cet[["CEInputID"] + cols], on="CEInputID", suffixes=("_poc", "_cet"), how="inner")
    per = j[["CEInputID", "PrgID", "MeasDescription"]].copy()
    sums = {}
    for c in cols:
        per[f"{c}_poc"], per[f"{c}_cet"] = j[f"{c}_poc"], j[f"{c}_cet"]
        per[f"{c}_diff"] = j[f"{c}_poc"] - j[f"{c}_cet"]
        per[f"{c}_ratio"] = (j[f"{c}_poc"] / j[f"{c}_cet"]).where(j[f"{c}_cet"] != 0)
        if not c.endswith("Ratio"):
            sums[c] = (float(j[f"{c}_poc"].sum()), float(j[f"{c}_cet"].sum()))
    return per, sums, len(j)


def cet_comparison_summary(sums: dict[str, tuple[float, float]], matched: int, measures: int, cet_rows: int) -> pd.DataFrame:
    summary = pd.DataFrame([{"metric": c, "POC": tp, "CET": tc, "POC/CET": tp / tc if tc else np.nan} for c, (tp, tc) in sums.items()])
    summary.attrs["matched"] = f"{matched:,} of {measures:,} measures matched by CEInputID ({cet_rows:,} in the CET file)"
    return summary


def compare_to_cet(frame: pd.DataFrame, cet_outputs: str | Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(per-measure, summary) of the POC against a real CET run's *_outputs.csv, matched on CEInputID (= CET_ID). Only meaningful for a job run the way the CET
    run was (same First Year, same ACC version, statewide programs allocated or not to match)."""
    cet = read_cet_outputs(cet_outputs)
    per, sums, matched = compare_to_cet_part(frame, cet)
    return per, cet_comparison_summary(sums, matched, len(frame), len(cet))


FORMATS = ("csv", "csv.zip", "parquet", "xlsx")
PARQUET_ROW_GROUP = 200_000  # rows per row group when a run in parts joins its parquet parts (TableWriter); RunConfig.parquet_row_group
CSV_BATCH_ROWS = 10_000      # rows formatted at a time when CSV is written from parquet (50,000 held 2.3 GB more at 1M rows; 10,000 holds 0.4 GB, as fast)
RESULTS_DIR = "_results"     # a run's measure-level table as parquet when its output is CSV, kept for the Results tables (results_source)


def table_path(path: Path, fmt: str) -> Path:
    """The file a table named `path` (no suffix) is written to in format `fmt`: <name>.csv / .parquet / .xlsx, or <name>.csv.zip (a zip holding <name>.csv)."""
    return path.with_name(path.name + ".csv.zip") if fmt == "csv.zip" else path.with_suffix("." + fmt)


def deflater() -> tuple[object, str]:
    """(a function giving a fresh raw-deflate compressor, its name). ISA-L level 1 when the optional `isal` package is installed (about 6x faster than zlib
    level 6, files about 11% larger: the 1,000,000-row measure-level CSV zips to 862 MB instead of 777 MB), else Python's zlib at level 6. Either is ordinary
    deflate, readable by every zip tool."""
    import zlib
    try:
        from isal import isal_zlib
        return (lambda: isal_zlib.compressobj(1, isal_zlib.DEFLATED, -15)), "ISA-L level 1"
    except ImportError:
        return (lambda: zlib.compressobj(6, zlib.DEFLATED, -15)), "zlib level 6"


def _csv_ready(batch):
    """A record batch with the columns pyarrow would write differently from pandas made as pandas writes them: booleans as True / False (pyarrow writes
    true / false) and timestamps as 2018-01-01 00:00:00 (pyarrow adds nanoseconds). Numbers are pyarrow's shortest round-trip text, so 1.0 is written 1 and
    small values may be written without an exponent (7.04e-05 as 0.0000704...): the same numbers when read back, different text from pandas' to_csv."""
    import pyarrow as pa
    import pyarrow.compute as pc
    cols, changed = [], False
    for f, col in zip(batch.schema, batch.columns):
        if pa.types.is_boolean(f.type):
            col, changed = pc.if_else(col, "True", "False"), True
        elif pa.types.is_timestamp(f.type):                  # whole seconds first: strftime's %S carries the unit's fraction
            col, changed = pc.strftime(pc.cast(col, pa.timestamp("s", f.type.tz), safe=False), format="%Y-%m-%d %H:%M:%S"), True
        cols.append(col)
    return pa.RecordBatch.from_arrays(cols, names=batch.schema.names) if changed else batch


def write_csv(batches, path: Path, zipped: bool) -> Path:
    """Write record batches (an iterable) as one CSV, UTF-8 with a BOM, at `path`; `zipped`: inside a zip at `path` holding <name without .zip>, dated now.
    pyarrow formats each batch (about 8x faster than pandas' to_csv) on a second thread while this one compresses and writes the previous one, so the
    formatting and the compression overlap; only two batches are held at a time. The uncompressed CSV never reaches the disk when zipped."""
    import queue
    import threading
    import pyarrow as pa
    import pyarrow.csv as pacsv
    q: queue.Queue = queue.Queue(maxsize=2)
    failed: list[BaseException] = []
    stop = threading.Event()

    def produce():
        try:
            header = True
            for b in batches:
                if stop.is_set():
                    return
                buf = pa.BufferOutputStream()
                pacsv.write_csv(pa.Table.from_batches([_csv_ready(b)]), buf, pacsv.WriteOptions(include_header=header, quoting_style="needed"))
                header = False
                q.put(buf.getvalue())
        except BaseException as exc:                        # handed to the writing thread, which raises it
            failed.append(exc)
        finally:
            q.put(None)

    th = threading.Thread(target=produce, name="recet-csv", daemon=True)
    th.start()

    def drain(out) -> None:
        out.write(b"\xef\xbb\xbf")                          # BOM, so Excel reads UTF-8
        while (chunk := q.get()) is not None:
            out.write(chunk)

    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        if zipped:
            import time
            import zipfile
            info = zipfile.ZipInfo(path.name[: -len(".zip")], date_time=time.localtime()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            make, _name = deflater()
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as z, z.open(info, "w", force_zip64=True) as m:
                if hasattr(m, "_compressor"):                 # zipfile's own deflater replaced by the faster one (same deflate stream format)
                    m._compressor = make()
                drain(m)
        else:
            with open(path, "wb") as f:
                drain(f)
    except BaseException:
        stop.set()
        while th.is_alive():                                  # let the formatting thread finish so nothing is left writing
            try:
                q.get(timeout=0.1)
            except queue.Empty:
                pass
        path.unlink(missing_ok=True)
        raise
    th.join()
    if failed:
        path.unlink(missing_ok=True)
        raise failed[0]
    return path


def write_csv_from_parquet(parquet: Path, path: Path, zipped: bool, on_progress=None, columns: list[str] | None = None) -> Path:
    """The parquet table at `parquet` as CSV (write_csv), read CSV_BATCH_ROWS rows at a time. Kept small on purpose, measured on the 1,000,000-row
    measure-level table: read single-threaded through a 1 MB buffer with no read-ahead (pyarrow otherwise pre-reads whole row groups), and with pyarrow's
    system allocator for the duration (its default, mimalloc, keeps freed memory), the conversion holds about 0.4 GB, against 2.3 GB with the defaults and
    50,000-row batches, in the same time (about 18 s with ISA-L on the laptop). `on_progress(share)` is called after each batch is read (share 0-1).
    `columns` writes only those columns, in that order (an output template)."""
    import pyarrow as pa
    import pyarrow.parquet as pq
    previous = pa.default_memory_pool()
    pa.set_memory_pool(pa.system_memory_pool())
    try:
        pf = pq.ParquetFile(parquet, pre_buffer=False, buffer_size=1 << 20)
        batches = pf.iter_batches(batch_size=CSV_BATCH_ROWS, use_threads=False, columns=columns)
        if columns is not None:
            batches = (_select(b, columns) for b in batches)
        if on_progress is not None:
            batches = _counted(batches, max(1, -(-pf.metadata.num_rows // CSV_BATCH_ROWS)), on_progress)
        return write_csv(batches, path, zipped)
    finally:
        pa.set_memory_pool(previous)


def _select(batch, columns: list[str]):
    """The batch's `columns`, in that order (iter_batches gives them in the file's order)."""
    import pyarrow as pa
    return pa.RecordBatch.from_arrays([batch.column(c) for c in columns], names=list(columns))


def _parquet_select(src: Path, dst: Path, columns: list[str], row_group_rows: int) -> None:
    """A parquet file holding only `columns` of `src`, in that order, one row group at a time."""
    import pyarrow as pa
    import pyarrow.parquet as pq
    pf = pq.ParquetFile(src)
    schema = pa.schema([pf.schema_arrow.field(c) for c in columns])
    with pq.ParquetWriter(dst, schema) as w:
        for b in pf.iter_batches(batch_size=row_group_rows, columns=columns, use_threads=False):
            w.write_table(pa.Table.from_batches([_select(b, columns)], schema=schema), row_group_size=row_group_rows)


def _counted(items, total: int, on_progress):
    """`items`, calling on_progress(share done) after each one is taken."""
    for i, x in enumerate(items):
        yield x
        on_progress(min(1.0, (i + 1) / total))


def _frame_batches(df: pd.DataFrame):
    import pyarrow as pa
    t = pa.Table.from_pandas(df, preserve_index=False)
    return t.to_batches(max_chunksize=CSV_BATCH_ROWS) or [pa.RecordBatch.from_pylist([], schema=t.schema)]


def _remove(path: Path) -> None:
    """Delete a temporary file, retrying for a moment if Windows still holds it (an antivirus scanner briefly locks newly written files); one that
    still cannot be deleted is left for the run folder's cleanup rather than failing the run."""
    import time
    for attempt in range(10):
        try:
            path.unlink(missing_ok=True)
            return
        except PermissionError:
            time.sleep(0.1 * (attempt + 1))


def results_source(folder: Path) -> Path | None:
    """The file a run folder's Results tables read: the kept parquet of a CSV run (RESULTS_DIR), else its measure-level file; None if there is neither."""
    folder = Path(folder)
    kept = sorted((folder / RESULTS_DIR).glob("measure_level_*.parquet"))
    if kept:
        return kept[0]
    files = sorted(folder.glob("measure_level_*.*"))
    return files[0] if files else None


EXCEL_MAX_ROWS = 1_048_575          # Excel's own limit: rows below the header
EXCEL_ROW_LIMIT = 50_000            # ReCET's limit for an xlsx table (RunConfig.excel_max_rows): writing runs at about 400 rows/s for the 344-column
                                    # measure-level table, so 50,000 rows take about 2 minutes and make about a 100 MB workbook; 1,000,000 would take 40 minutes


@dataclass
class ExcelLimit:
    """How many rows an xlsx table may have, the format a longer table is written in instead (csv; the web app uses csv.zip), and where to say so."""
    max_rows: int = EXCEL_ROW_LIMIT
    fallback: str = "csv"
    notify: Callable[[str], None] | None = None

    def format_for(self, fmt: str, rows: int, name: str) -> str:
        if fmt != "xlsx" or rows <= min(self.max_rows, EXCEL_MAX_ROWS):
            return fmt
        fallback = self.fallback if self.fallback in FORMATS and self.fallback != "xlsx" else "csv"
        msg = (f"{name}: {rows:,} rows is more than the {min(self.max_rows, EXCEL_MAX_ROWS):,}-row limit for Excel output, so it is written as "
               f"{'zipped CSV' if fallback == 'csv.zip' else fallback.upper() if fallback == 'csv' else fallback} instead")
        if self.notify is not None:
            self.notify(msg)
        else:
            import warnings
            warnings.warn(msg, stacklevel=3)
        return fallback


def write_table(df: pd.DataFrame, path: Path, fmt: str = "csv", excel: ExcelLimit | None = None) -> Path:
    """Write `df` as <path with suffix from fmt> (table_path). csv is UTF-8 with a BOM so Excel opens it directly, formatted by pyarrow (write_csv); csv.zip
    is the same CSV inside a zip (what the web app writes: about a fifth of the size); xlsx is one formatted sheet (see _write_xlsx). An xlsx table longer than
    `excel.max_rows` (default EXCEL_ROW_LIMIT) is written in `excel.fallback` (default csv) instead, and `excel.notify` is told."""
    fmt = (excel or ExcelLimit()).format_for(fmt, len(df), path.name)
    path = table_path(path, fmt)
    path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "parquet":
        df.to_parquet(path, index=False)
    elif fmt in ("csv", "csv.zip"):
        write_csv(_frame_batches(df), path, zipped=fmt == "csv.zip")
    elif fmt == "xlsx":
        _write_xlsx(df, path)
    else:
        raise ValueError(f"unknown output format {fmt!r}; choose from {FORMATS}")
    return path


def _write_xlsx(df: pd.DataFrame, path: Path) -> None:
    """One sheet named after the table: a bold, shaded, frozen header with filters, columns sized to their content, numbers as numbers and text as text (a text
    cell starting with '=' stays text, never a formula; characters Excel cannot store are dropped)."""
    x = _XlsxStream(path, df, len(df))
    x.append(df)
    x.close()


class _XlsxStream:
    """_write_xlsx for a table written in parts: the header, column widths (from the first part) and filter range (from `total_rows`) are set up front, rows are
    appended part by part (openpyxl's write-only mode keeps them on disk, not in memory)."""

    def __init__(self, path: Path, first: pd.DataFrame, total_rows: int):
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        self.path = path
        self.wb = Workbook(write_only=True)
        ws = self.ws = self.wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", path.stem)[:31])
        ws.freeze_panes = "A2"
        cols = list(first.columns)
        for j, c in enumerate(cols, start=1):
            s = first[c]
            sample = s.head(200).astype(str).str.len().max() if len(s) else 0
            sample = 0 if pd.isna(sample) else int(sample)
            ws.column_dimensions[get_column_letter(j)].width = max(8, min(45, max(len(str(c)), sample) + 2))
        if cols:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{total_rows + 1}"
        head_font, head_fill = Font(name="Arial", bold=True, color="FFFFFF", size=10), PatternFill("solid", fgColor="1F3864")
        header = []
        for c in cols:
            cell = WriteOnlyCell(ws, value=str(c))
            cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(vertical="center")
            header.append(cell)
        ws.append(header)

    def append(self, df: pd.DataFrame) -> None:
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        ws = self.ws

        def convert(v):
            if v is None:
                return None
            if isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
                if v.startswith("="):
                    cell = WriteOnlyCell(ws, value=v)
                    cell.data_type = "s"
                    return cell
                return v
            if isinstance(v, (np.floating, float)):
                return None if np.isnan(v) or np.isinf(v) else float(v)
            if isinstance(v, (np.integer,)):
                return int(v)
            if isinstance(v, (np.bool_, bool)):
                return bool(v)
            if isinstance(v, pd.Timestamp):
                return None if pd.isna(v) else v.to_pydatetime()
            if v is pd.NA or v is pd.NaT:
                return None
            return v
        for row in df.itertuples(index=False, name=None):
            ws.append([convert(v) for v in row])

    def close(self) -> None:
        self.wb.save(self.path)


class TableWriter:
    """write_table for a table that arrives in parts (a chunked run), holding one part at a time. `total_rows` (known before the first part) decides, as
    write_table does, whether an xlsx table fits in Excel. Every part must have the first part's columns (in any order). xlsx rows are streamed. parquet parts
    are written as separate files and joined at close() under one schema (a column blank in one part takes the type it has elsewhere), gathered into row groups
    of `row_group_rows` rows. Bigger groups compress much better (a 1,000,000-row measure-level table: 666 MB in 20,000-row groups, 302 MB in 100,000, 253 MB
    in 200,000) but each is held in memory while it is written (about 0.4 GB per 100,000 rows of the measure-level table), so the size is a setting:
    RunConfig.parquet_row_group.

    csv and csv.zip are written PARQUET FIRST: the parts are joined into a parquet file as above and that file is then written out as CSV by pyarrow (write_csv;
    a 1,000,000-row table in about 30 s, against about 260 s for pandas). `keep_parquet` keeps that parquet in the run's RESULTS_DIR (the Results tables read
    it: results_source); otherwise it is deleted.

    `select` (an output template, BRD BR-06.10) writes only those columns, in that order, to the file; a column the table does not have is left out
    (`left_out`). The table then always goes through parquet, kept in full when `keep_parquet` (so the Results tables still have every column), and the file is
    written from it."""

    def __init__(self, path: Path, fmt: str, total_rows: int, row_group_rows: int = PARQUET_ROW_GROUP, keep_parquet: bool = False,
                 excel: ExcelLimit | None = None, select: list[str] | None = None):
        if fmt not in FORMATS:
            raise ValueError(f"unknown output format {fmt!r}; choose from {FORMATS}")
        self.excel = excel
        fmt = (excel or ExcelLimit()).format_for(fmt, total_rows, path.name)
        self.fmt, self.total, self.base = fmt, total_rows, path
        self.row_group_rows = max(int(row_group_rows or PARQUET_ROW_GROUP), 1)
        self.path = table_path(path, fmt)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if select:
            from .output_templates import current
            select = current(select)                 # a template saved before 0.1.50 names some columns by their old names
        self.select = select or None
        self.via_parquet = fmt in ("csv", "csv.zip") or self.select is not None
        self.keep_parquet = keep_parquet
        self.left_out: list[str] = []
        self.parquet = (path.parent / RESULTS_DIR / f"{path.name}.parquet") if self.via_parquet else self.path
        self.cols: list[str] | None = None
        self.parts: list[Path] = []
        self.xlsx: _XlsxStream | None = None
        self.rows = 0

    def write(self, df: pd.DataFrame) -> None:
        if self.cols is None:
            self.cols = list(df.columns)
        elif list(df.columns) != self.cols:
            if set(df.columns) != set(self.cols):
                raise ValueError(f"{self.path.name}: a part has different columns: extra {sorted(set(df.columns) - set(self.cols))}, "
                                 f"missing {sorted(set(self.cols) - set(df.columns))}")
            df = df[self.cols]
        if self.fmt == "xlsx" and self.select is None:
            if self.xlsx is None:
                self.xlsx = _XlsxStream(self.path, df, self.total)
            self.xlsx.append(df)
        else:
            self.parquet.parent.mkdir(parents=True, exist_ok=True)
            part = self.parquet.with_name(f"{self.base.name}.part{len(self.parts):04d}.parquet")
            df.to_parquet(part, index=False)       # with dictionary encoding: without it the write is ~0.2 s faster a part, but the join's peak is 0.4 GB higher
            self.parts.append(part)
        self.rows += len(df)

    def _join_parquet(self, on_progress=None) -> None:
        import pyarrow as pa
        import pyarrow.parquet as pq
        if not self.parts:
            self.parquet.parent.mkdir(parents=True, exist_ok=True)
            pd.DataFrame(columns=self.cols or []).to_parquet(self.parquet, index=False)
            return
        schemas = [pq.read_schema(p) for p in self.parts]
        schema = pa.unify_schemas([s.remove_metadata() for s in schemas], promote_options="permissive").with_metadata(schemas[0].metadata)
        plain = schema.remove_metadata()
        with pq.ParquetWriter(self.parquet, schema) as w:
            n, buf, rows = self.row_group_rows, [], 0
            for k, p in enumerate(self.parts):                   # parts gathered into full row groups of n rows (about one group in memory)
                if on_progress is not None:
                    on_progress(k / len(self.parts))
                buf.append(pq.read_table(p).replace_schema_metadata(None).cast(plain))
                rows += buf[-1].num_rows
                _remove(p)
                if rows >= n:
                    t = pa.concat_tables(buf)
                    while t.num_rows >= n:
                        w.write_table(t.slice(0, n).replace_schema_metadata(schema.metadata), row_group_size=n)
                        t = t.slice(n)
                    buf, rows = ([t] if t.num_rows else []), t.num_rows
            if buf:
                w.write_table(pa.concat_tables(buf).replace_schema_metadata(schema.metadata), row_group_size=n)
        for p in self.parts:
            _remove(p)

    def close(self, on_progress=None) -> Path:
        """Finish the file. `on_progress(joined, converted)` (each a share 0-1) is called as the parquet parts are joined and the CSV is written."""
        if self.fmt == "xlsx" and self.select is None:
            if self.xlsx is not None:
                self.xlsx.close()
            else:
                write_table(pd.DataFrame(columns=self.cols or []), self.base, self.fmt, self.excel)
            return self.path
        self._join_parquet(None if on_progress is None else lambda j: on_progress(j, 0.0))
        if on_progress is not None:
            on_progress(1.0, 0.0)
        if self.via_parquet:
            cols = self.output_columns
            if self.select is not None:
                self.left_out = [c for c in self.select if c not in (self.cols or [])]
            if self.fmt in ("csv", "csv.zip"):
                write_csv_from_parquet(self.parquet, self.path, zipped=self.fmt == "csv.zip",
                                       on_progress=None if on_progress is None else lambda c: on_progress(1.0, c),
                                       columns=cols if self.select is not None else None)
            elif self.fmt == "parquet":
                _parquet_select(self.parquet, self.path, cols, self.row_group_rows)
            else:                                                     # xlsx from the parquet, a row group at a time
                import pyarrow.parquet as pq
                x = None
                for b in pq.ParquetFile(self.parquet).iter_batches(batch_size=CSV_BATCH_ROWS, columns=cols, use_threads=False):
                    df = _select(b, cols).to_pandas()
                    x = x or _XlsxStream(self.path, df, self.rows)
                    x.append(df)
                if x is None:
                    write_table(pd.DataFrame(columns=cols), self.base, self.fmt, self.excel)
                else:
                    x.close()
            if not self.keep_parquet:
                self.parquet.unlink(missing_ok=True)
                try:
                    self.parquet.parent.rmdir()
                except OSError:
                    pass
        return self.path

    @property
    def output_columns(self) -> list[str]:
        """The columns the file holds, in order: the template's that the table has, or every column."""
        if self.select is None:
            return list(self.cols or [])
        return [c for c in self.select if c in (self.cols or [])]


def read_table(path: Path) -> pd.DataFrame:
    """Read a table written by write_table, whatever its format."""
    path = Path(path)
    if path.suffix == ".parquet":
        return pd.read_parquet(path)
    if path.suffix == ".xlsx":
        return pd.read_excel(path)
    if path.suffix == ".zip":                                 # csv.zip: the one CSV inside
        import zipfile
        with zipfile.ZipFile(path) as z:
            name = next(n for n in z.namelist() if n.lower().endswith(".csv"))
            with z.open(name) as f:
                return pd.read_csv(f, low_memory=False, encoding="utf-8-sig")
    return pd.read_csv(path, low_memory=False, encoding="utf-8-sig")
