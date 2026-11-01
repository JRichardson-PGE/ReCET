"""Editing a small job's Measure and ProgramCost tables inside the tool (the window's "Edit inputs" tab). No tkinter here, so it is testable and reusable.

The editor is meant for small jobs, such as an implementer's project of a few measures, where cell-by-cell editing is the natural way to work. Jobs with more than
`MAX_EDIT_ROWS` measures are refused: edit those offline (Excel or the job workbook template) and run them from the Job tab.

* `load_tables()` reads a job from the Job tab's inputs (a job workbook and/or the CET files; a CET file wins over its workbook sheet, as in a run) into two
  all-text DataFrames with every CET column, in CET's order.
* `ChoiceSource` gives each column's dropdown values: the CET value lists shipped in `data/value_lists/` (codes active for the claim year; in the web app with
  its admins' changes on top, `set_value_list_changes`), the IOUs the POC can
  price, and, for `AC_Elec_Impact_Profile` / `AC_Gas_Impact_Profile`, the profiles that actually exist for the row's IOU and target sector (library and supplied files).
* `check_tables()` returns one `CellIssue` per problem cell. It applies the preflight's own rules (`preflight.py`: required values, numbers, CET ranges, IOU,
  ClaimYearQuarter, duplicate IDs, the second baseline, gas sector, PrgID alignment), plus "not an active CET value" and "no profile, so not priced" warnings.
  The full preflight (`Full check` in the window) still runs on the saved files; this check is the instant per-cell version of it.
* `write_cet_files()` / `write_workbook()` save the tables; `prepare_run_inputs()` writes them for a run and returns the RunConfig changes.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from functools import lru_cache
from pathlib import Path

import numpy as np
import pandas as pd

from . import paths, preflight, templates

MAX_EDIT_ROWS = 500
MEASURE_COLUMNS = list(templates.MEASURE_COLUMNS)
PROGRAMCOST_COLUMNS = list(templates.PROGRAMCOST_COLUMNS)
PROGRAMCOST_NUMERIC = [c for c in PROGRAMCOST_COLUMNS if c not in ("PrgID", "PrgYear", "ClaimYearQuarter", "PA")]
MEASURE_NUMERIC = preflight.NUMERIC_COLUMNS + ["UnitkWhIOUWater1stBaseline", "UnitkWhIOUWater2ndBaseline"]
FLAG_COLUMNS = ("Residential_Flag", "Upstream_Flag")
END_USE_COLUMNS = ("AC_Elec_Impact_Profile", "AC_Elec_Impact_Profile_2nd_Baseline", "AC_Elec_Impact_Profile_Water", "AC_Elec_Impact_Profile_Parent",
                   "AC_Elec_Impact_Profile_Parent_2nd_Baseline")   # electric shapes: per-row dropdowns (IOU and target sector)
CYQ = re.compile(r"^\d{4}Q[1-4]$")

# Measure / ProgramCost column -> CET value-list file in data/value_lists (RUL_ID is not checked: CET's RUL IDs are not a published list)
VALUE_LISTS = {
    "Sector": "sector", "DeliveryType": "deliverytype", "BldgType": "bldgtype", "AC_Climate_Zone": "ac_climate_zone", "AC_Gas_Impact_Profile": "ac_gas_impact_profile",
    "AC_Gas_Target_Sector": "ac_gas_target_sector", "AC_Elec_Impact_Profile": "ac_elec_impact_profile", "AC_Elec_Target_Sector": "ac_elec_target_sector", "MeasAppType": "measapptype",
    "MeasImpactType": "measimpacttype", "FuelID": "fuelid", "TechGroup": "techgroup", "TechType": "techtype", "UseCategory": "usecategory",
    "UseSubCategory": "usesubcategory", "WaterUse": "wateruse", "Version": "version", "NormUnit": "normunit", "NTG_ID": "ntg", "EUL_ID": "eul",
    "GSIA_ID": "gsia", "RateScheduleElec": "ratescheduleelec", "RateScheduleGas": "rateschedulegas", "CombustionType": "combustiontype",
}
PROGRAMCOST_VALUE_LISTS = {"PA": "pa"}

# Defaults for a new, blank measure row. NTGR*, EUL_Yrs and the descriptive fields are left blank on purpose: they have no safe default.
NEW_MEASURE_DEFAULTS = {
    **{c: "0" for c in MEASURE_NUMERIC if c.startswith("Unit")},
    "NumUnits": "1", "RUL_Yrs": "0", "MeasInflation": "0", "Residential_Flag": "0", "Upstream_Flag": "0",
    **{f"{k}{f}": "1" for k in ("RealizationRate", "InstallationRate") for f in ("kW", "kWh", "Therm")},
}


# Display order in the editor (the files keep CET's order). Pinned on the left, then the columns that change a result (grouped as in INPUT_FIELDS.md 1.3-1.8), then
# the ones that change nothing (carried to the output only). "Changes a result" = read by the calculation or into a derived output column; the kW fields only feed
# the kW savings columns, so they come last in that group. Sector changes nothing; RateScheduleElec/Gas are not used (the schedule follows IOU and sector).
MEASURE_FROZEN = ["CEInputID", "MeasCode", "MeasDescription"]
MEASURE_CALC = [
    "PrgID", "ClaimYearQuarter", "ACCVersion", "IOU_AC_Territory", "AC_Climate_Zone", "AC_Elec_Target_Sector", "AC_Elec_Impact_Profile", "AC_Elec_Impact_Profile_2nd_Baseline",
    "AC_Elec_Impact_Profile_Water", "AC_Elec_Impact_Profile_Parent", "AC_Elec_Impact_Profile_Parent_2nd_Baseline", "AC_Gas_Target_Sector", "AC_Gas_Impact_Profile",
    "FuelID", "MeasImpactType", "CombustionType",
    "NumUnits", "UnitkWh1stBaseline", "UnitkWh2ndBaseline", "UnitTherm1stBaseline", "UnitTherm2ndBaseline", "UnitkWhIOUWater1stBaseline",
    "UnitkWhIOUWater2ndBaseline", "EUL_Yrs", "RUL_Yrs",
    "NTGRkWh", "NTGRTherm", "NTGRCost", "RealizationRatekWh", "RealizationRateTherm", "InstallationRatekWh", "InstallationRateTherm",
    "UnitMeaCost1stBaseline", "UnitMeaCost2ndBaseline", "UnitDirectInstallLab", "UnitDirectInstallMat", "UnitIncentiveToOthers", "UnitEndUserRebate", "MeasInflation",
    "UnitGasInfraBens", "UnitRefrigBens", "UnitMiscBens", "UnitRefrigCosts", "UnitMiscCosts", "UnitTaxCredits",
    "UnitkW1stBaseline", "UnitkW2ndBaseline", "NTGRkW", "RealizationRatekW", "InstallationRatekW",
]
MEASURE_NO_EFFECT = [c for c in templates.MEASURE_COLUMNS if c not in MEASURE_FROZEN + MEASURE_CALC]
PROGRAMCOST_FROZEN = ["PrgID"]
PROGRAMCOST_CALC = ["ClaimYearQuarter", "PA", "AdminCostsOverheadAndGA", "AdminCostsOther", "MarketingOutreach", "DIActivity", "DIInstallation",
                    "DIHardwareAndMaterials", "DIRebateAndInspection", "EMV", "UserInputIncentive", "CostsRecoveredFromOtherSources"]
PROGRAMCOST_NO_EFFECT = [c for c in templates.PROGRAMCOST_COLUMNS if c not in PROGRAMCOST_FROZEN + PROGRAMCOST_CALC]   # PrgYear, OnBillFinancing (not summed)


CALC_COLUMNS = {"Measure": frozenset(MEASURE_CALC), "ProgramCost": frozenset(PROGRAMCOST_CALC) | {"PrgID"}}   # PrgID shares program costs among measures


def affects_results(table: str, column: str) -> bool:
    """True if the column changes a result (a priced value or a derived output column); False if it is only carried to the output (or is the row's ID)."""
    return column in CALC_COLUMNS[table]


def display_layout(table: str, columns: list[str]) -> tuple[list[str], list[str], set[str]]:
    """(pinned columns, scrolling columns, columns with no effect on any result) for a table holding `columns`. Columns the input had beyond CET's go last.
    The pinned measure columns (ID, code, description) change no result, so they are in the third set too; the pinned PrgID of program costs does."""
    frozen, calc, rest = ((MEASURE_FROZEN, MEASURE_CALC, MEASURE_NO_EFFECT) if table == "Measure"
                          else (PROGRAMCOST_FROZEN, PROGRAMCOST_CALC, PROGRAMCOST_NO_EFFECT))
    extra = [c for c in columns if c not in frozen + calc + rest]
    pinned = [c for c in frozen if c in columns]
    scroll = [c for c in calc + rest if c in columns] + extra
    return pinned, scroll, {c for c in pinned + scroll if not affects_results(table, c)}


DISPLAY_DECIMALS = 6


def display_number(text, decimals: int = DISPLAY_DECIMALS) -> str:
    """How the editor shows a number: at most `decimals` decimal places, trailing zeros dropped ("0.600000" -> "0.6", "15.0" -> "15"). A value too small
    to show that way gets 3 significant figures ("0.0000004" -> "4e-07"). Text that is not a number is shown as it is. Display only: the stored value is
    what is edited, checked and saved."""
    s = "" if text is None else str(text).strip()
    if not s:
        return ""
    try:
        f = float(s)
    except ValueError:
        return s
    if f != f or f in (float("inf"), float("-inf")):
        return s
    out = f"{f:.{decimals}f}".rstrip("0").rstrip(".")
    if out in ("0", "-0") and f != 0:
        out = f"{f:.3g}"
    return "0" if out == "-0" else out


def numeric_columns(table: str) -> set[str]:
    return set(MEASURE_NUMERIC) if table == "Measure" else set(PROGRAMCOST_NUMERIC) | {"PrgYear"}


# ------------------------------------------------------------------------------------------------------------------------- what each column affects
# Shown as the column header's hover tooltip. (affects, how) per column, from CALCULATIONS.md / INPUT_FIELDS.md. A column with no entry here changes nothing.
_ELEC = "electric benefits and supply costs (TRC, SCT Base, SCT High)"
_GAS = "gas benefits and supply costs (TRC, SCT Base, SCT High)"
_REST = "total system benefit, and every cost-test cost and ratio"
_INCENT = "rebates and incentives (in the PAC, TRC, RIM and SCT costs and ratios)"
_KW = "only the reported kW savings columns (FirstYearNetKW ...)"
_KW_HOW = "Not priced: the 8,760 load shape carries the demand (T&D) value."
_PROGRAM = "program costs: WeightedProgramCost and the program-cost part of the TRC, PAC, RIM and SCT costs and ratios (not TSB, not the NoAdmin tests)"
_FUELSUB = ("whether negative fuel values count as supply costs or negative benefits, so the cost-test costs and ratios, and the bill-impact split "
            "(total system benefit is the same either way)")
EFFECTS = {
    "Measure": {
        "PrgID": ("the program-cost share (WeightedProgramCost) and so the program-cost part of the TRC, PAC, RIM and SCT costs and ratios; summary groupings "
                  "(credited PA, program segment); statewide allocation when it is on",
                  "Measures with the same PrgID share that program's ProgramCost costs in proportion to their positive benefits. It does not change the measure's own "
                  "benefits."),
        "ClaimYearQuarter": ("everything valued over time: all benefits and supply costs, participant cost, bill impacts, emissions, " + _REST,
                             "Sets when savings start, which ACC years apply, and how far each value is discounted to the First Year."),
        "IOU_AC_Territory": ("all benefits and supply costs, every present value (discount rate), bill impacts, emissions, " + _REST,
                             "Picks the IOU's avoided costs, discount rate, retail rates, gas CO2 rate and NOx/PM10 factors, and is part of the 8,760 profile match."),
        "AC_Climate_Zone": (_ELEC + ", water-energy value, electric emissions, " + _REST,
                          "Picks the ACC climate zone and the profile's shape for that zone; part of the NOx/PM10 key."),
        "AC_Elec_Target_Sector": ("electric benefits (through the profile match), the electric bill impact, electric NOx/PM10, " + _REST,
                           "Part of the 8,760 profile match; picks the electric retail rate schedule and the NOx/PM10 factors."),
        "AC_Elec_Impact_Profile": (_ELEC + ", water-energy value, electric emissions, " + _REST + " (nothing electric is priced if no profile matches)",
                                 "The 8,760-hour load shape the kWh savings are spread over before pricing; part of the NOx/PM10 key."),
        "AC_Elec_Impact_Profile_2nd_Baseline": ("the 2nd-baseline years of " + _ELEC + ", electric emissions, " + _REST,
                                    "Optional ReCET column (not a CET field): the load shape of the years after RUL_Yrs on an early replacement. Blank = AC_Elec_Impact_Profile."),
        "AC_Elec_Impact_Profile_Water": ("water-energy value, " + _REST,
                                      "Optional ReCET column (not a CET field): the load shape the water energy is spread over. Blank = AC_Elec_Impact_Profile."),
        "AC_Elec_Impact_Profile_Parent": ("the incremental columns (the record priced on this profile, subtracted from its own value); nothing else",
                                          "Optional ReCET column (not a CET field): a true-up's parent profile, the load shape of the claim it adjusts. Blank = no parent."),
        "AC_Elec_Impact_Profile_Parent_2nd_Baseline": ("the incremental columns of an early replacement's 2nd-baseline years; nothing else",
                                                       "Optional ReCET column (not a CET field): the parent's 2nd-baseline load shape. Blank = AC_Elec_Impact_Profile_Parent."),
        "ACCVersion": ("every avoided-cost value of the measure, its discount rates, bill impacts and emissions, " + _REST,
                       "Optional ReCET column (not a CET field): the ACC version this measure is priced on (2021, 2022, 2024, 2026). Blank = the run's version."),
        "AC_Gas_Target_Sector": (_GAS + ", the gas bill impact, " + _REST, "Picks the gas ACC class (Residential, Commercial, Core) and the gas retail rate schedule."),
        "AC_Gas_Impact_Profile": (_GAS + ", " + _REST + " (no gas value if the name matches nothing)",
                            "The 12-month shape the therm savings are spread over. A ...FuelSub... name on a residential measure adds behind-the-meter methane."),
        "FuelID": (_FUELSUB, "AllElec-New, FuelSub-ToElec and FuelSub-ToGas mark a fuel-substitution measure."),
        "MeasImpactType": (_FUELSUB, "A value containing 'FuelSub' also marks a fuel-substitution measure (see FuelID)."),
        "CombustionType": ("gas NOx emissions only", "Looks up the NOx rate (lb per therm) for the combustion type; blank or unknown gives 0."),
        "NumUnits": ("every quantity, benefit, cost and emission", "Multiplies every per-unit value."),
        "UnitkWh1stBaseline": (_ELEC + ", the electric bill impact, electric emissions, " + _REST, "Annual kWh savings per unit against the 1st baseline."),
        "UnitkWh2ndBaseline": ("the same as UnitkWh1stBaseline, for the years after RUL_Yrs", "Used only on early-replacement measures (RUL_Yrs > 0)."),
        "UnitTherm1stBaseline": (_GAS + ", the gas bill impact, gas emissions, " + _REST, "Annual therm savings per unit against the 1st baseline."),
        "UnitTherm2ndBaseline": ("the same as UnitTherm1stBaseline, for the years after RUL_Yrs", "Used only on early-replacement measures (RUL_Yrs > 0)."),
        "UnitkWhIOUWater1stBaseline": ("water-energy benefits or costs, " + _REST,
                                       "IOU energy embedded in the water saved, priced on the generation part of the electric avoided cost (no T&D)."),
        "UnitkWhIOUWater2ndBaseline": ("the same as UnitkWhIOUWater1stBaseline, for the years after RUL_Yrs", "Used only when RUL_Yrs > 0."),
        "EUL_Yrs": ("every lifetime value: benefits, supply costs, bill impacts, lifecycle emissions, " + _REST,
                    "Savings run for EUL years from the claim quarter, the last quarter prorated."),
        "RUL_Yrs": ("lifetime values (the switch to the 2nd baseline), the measure cost in the participant cost, bill impacts, emissions, " + _REST,
                    "Above 0 makes an early replacement: 1st-baseline savings until the RUL, 2nd-baseline after, and the measure cost uses UnitMeaCost2ndBaseline."),
        "NTGRkWh": ("net electric and water-energy values, other benefits and costs, tax credits, electric emissions, " + _REST,
                    "Net-to-gross for kWh (the Gross columns divide it back out). Required."),
        "NTGRTherm": ("net gas benefits and supply costs, gas emissions, the gas bill impact, " + _REST, "Net-to-gross for therms. Required."),
        "NTGRCost": ("the net participant cost, so the TRC and SCT costs and ratios", "Scales the participant cost. Required."),
        "RealizationRatekWh": (_ELEC + ", water-energy value, the electric bill impact, electric emissions, " + _REST, "Multiplies the kWh savings that are priced."),
        "InstallationRatekWh": (_ELEC + ", water-energy value, the electric bill impact, electric emissions, " + _REST, "Multiplies the kWh savings that are priced."),
        "RealizationRateTherm": (_GAS + ", the gas bill impact, gas emissions, " + _REST, "Multiplies the therm savings that are priced."),
        "InstallationRateTherm": (_GAS + ", the gas bill impact, gas emissions, " + _REST, "Multiplies the therm savings that are priced."),
        "UnitMeaCost1stBaseline": ("the participant cost (TRC and SCT costs and ratios) and the excess-incentive test",
                                   "Incremental measure cost; on an early replacement its difference from UnitMeaCost2ndBaseline is deferred to the RUL."),
        "UnitMeaCost2ndBaseline": ("the participant cost of early-replacement measures (TRC and SCT)", "Used only when RUL_Yrs > 0."),
        "UnitDirectInstallLab": (_INCENT + ", the participant cost and the excess-incentive test", "Counted as an incentive paid by the program."),
        "UnitDirectInstallMat": (_INCENT + ", the participant cost and the excess-incentive test", "Counted as an incentive paid by the program."),
        "UnitIncentiveToOthers": (_INCENT + ", the participant cost and the excess-incentive test", "Counted as an incentive paid by the program."),
        "UnitEndUserRebate": (_INCENT + " and the participant cost", "A rebate to the end user; not part of the excess-incentive test."),
        "MeasInflation": ("the participant cost of early-replacement measures", "Inflates the deferred measure cost when RUL_Yrs > 0; no effect otherwise."),
        "UnitGasInfraBens": ("other benefits, total system benefit, every ratio's numerator, program-cost shares", "Undiscounted per-unit benefit, net of NTGRkWh."),
        "UnitRefrigBens": ("other benefits, total system benefit, every ratio's numerator, program-cost shares", "Undiscounted per-unit benefit, net of NTGRkWh."),
        "UnitMiscBens": ("other benefits, total system benefit, every ratio's numerator, program-cost shares", "Undiscounted per-unit benefit, net of NTGRkWh."),
        "UnitRefrigCosts": ("other costs, total system benefit, every cost-test cost and ratio", "Undiscounted per-unit cost, net of NTGRkWh."),
        "UnitMiscCosts": ("other costs, total system benefit, every cost-test cost and ratio", "Undiscounted per-unit cost, net of NTGRkWh."),
        "UnitTaxCredits": ("the TRC ratio's numerator only", "Not part of total system benefit or of any cost."),
        "UnitkW1stBaseline": (_KW, _KW_HOW), "UnitkW2ndBaseline": (_KW, _KW_HOW), "NTGRkW": (_KW, _KW_HOW),
        "RealizationRatekW": (_KW, _KW_HOW), "InstallationRatekW": (_KW, _KW_HOW),
        # no effect, with a reason worth giving
        "CEInputID": (None, "The measure's key: it labels every output row and joins to a CET outputs file. Must be unique."),
        "Sector": (None, "Carried to the output only."),
        "RateScheduleElec": (None, "Not read: the electric rate schedule follows IOU_AC_Territory and AC_Elec_Target_Sector."),
        "RateScheduleGas": (None, "Not read: the gas rate schedule follows IOU_AC_Territory and AC_Gas_Target_Sector."),
        "NTG_ID": (None, "A reference ID; the NTGR* values are what is used."),
        "EUL_ID": (None, "A reference ID; EUL_Yrs is what is used."),
        "RUL_ID": (None, "A reference ID; RUL_Yrs is what is used."),
        "GSIA_ID": (None, "A reference ID; the installation rates are what is used."),
    },
    "ProgramCost": {
        "PrgID": ("which measures share this row's costs", "Must match the measures' PrgID (else the run stops)."),
        "ClaimYearQuarter": ("the discounting of this row's costs: " + _PROGRAM, "Its year is when the cost is incurred and discounted from."),
        "PA": ("the discount rate of this row's costs (" + _PROGRAM + "); the credited PA in the summary",
               "SW is treated as SCE unless statewide allocation is on."),
        **{c: (_PROGRAM, "Summed per program and shared among its measures in proportion to their positive benefits.")
           for c in ("AdminCostsOverheadAndGA", "AdminCostsOther", "MarketingOutreach", "DIActivity", "DIInstallation", "DIHardwareAndMaterials",
                     "DIRebateAndInspection", "EMV", "UserInputIncentive")},
        "CostsRecoveredFromOtherSources": (_PROGRAM, "Summed with the other costs: enter it as a negative to reduce the total, as CET does."),
        "PrgYear": (None, "No priced value: the run uses ClaimYearQuarter's year. Used only to pick the credited PA's row."),
        "OnBillFinancing": (None, "Not part of CET's program-cost sum."),
    },
}


def column_tooltip(table: str, column: str) -> str:
    """The hover text of a column header: whether it changes results, which outputs it affects, and how."""
    affects, how = EFFECTS[table].get(column, (None, ""))
    if affects_results(table, column):
        return f"{column}\nCHANGES RESULTS\n\nAffects: {affects}.\n\n{how}".rstrip()
    return f"{column}\nNO EFFECT ON RESULTS\n\n{how or 'Carried to the output unchanged.'}".rstrip()


class EditorError(ValueError):
    pass


# ------------------------------------------------------------------------------------------------------------------------- loading and saving
def _read_cet(path: str | Path) -> pd.DataFrame:
    p = Path(path)
    from .pipeline import delimited
    if delimited.detect_sep(p) is None:
        raise EditorError(f"{p.name} has no '|' or ',' in its header (use pipe- or comma-delimited text)")
    return delimited.read_cet_csv(p, dtype=str, keep_default_na=False, na_filter=False)


def conform(df: pd.DataFrame | None, columns: list[str]) -> pd.DataFrame:
    """All-text frame with every CET column in CET order (missing ones blank) followed by any extra columns the input had."""
    if df is None:
        return pd.DataFrame(columns=columns, dtype=str)
    df = df.fillna("").astype(str)
    extra = [c for c in df.columns if c not in columns]
    for c in columns:
        if c not in df.columns:
            df[c] = ""
    return df[columns + extra].reset_index(drop=True)


@dataclass
class LoadedJob:
    measures: pd.DataFrame
    programcost: pd.DataFrame
    source: str                          # a description for the window
    base_workbook: str | None = None     # the workbook the tables came from (its other sheets still apply to a run)


def load_tables(measures: str | None = None, programcost: str | None = None, workbook: str | None = None,
                max_rows: int = MAX_EDIT_ROWS) -> LoadedJob:
    """The job the Job tab points at. As in a run, a CET file given separately wins over the workbook's sheet of the same table."""
    from . import workbook as wb_mod
    m_df = pc_df = None
    src = []
    if workbook:
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            files = wb_mod.extract(workbook, tmp)
            if not measures:
                m_df = _read_cet(files["measures"])
                src.append(f"Measure sheet of {Path(workbook).name}")
            if not programcost and "programcost" in files:
                pc_df = _read_cet(files["programcost"])
                src.append(f"ProgramCost sheet of {Path(workbook).name}")
    if measures:
        m_df = _read_cet(measures)
        src.append(Path(measures).name)
    if programcost:
        pc_df = _read_cet(programcost)
        src.append(Path(programcost).name)
    if m_df is None:
        raise EditorError("choose a job workbook or a Measure.csv on the Job tab first")
    if len(m_df) > max_rows:
        raise EditorError(f"this job has {len(m_df):,} measures; the editor is for jobs of up to {max_rows:,}. Edit it offline (Excel) and run it from the Job tab.")
    return LoadedJob(conform(m_df, MEASURE_COLUMNS), conform(pc_df, PROGRAMCOST_COLUMNS), " + ".join(src), workbook)


def new_job() -> LoadedJob:
    """A new job holding the workbook template's example measure and program-cost row, to edit."""
    m = conform(pd.DataFrame([templates.EXAMPLE_MEASURE]), MEASURE_COLUMNS)
    pc = conform(pd.DataFrame([templates.EXAMPLE_PROGRAMCOST]), PROGRAMCOST_COLUMNS)
    return LoadedJob(m, pc, "new job (template example rows)")


def new_measure_row(like: dict | None = None, existing_ids=()) -> dict:
    """A blank measure with safe defaults; `like` (another row) supplies its program, IOU, zone, sectors and claim quarter."""
    row = {c: "" for c in MEASURE_COLUMNS}
    row.update(NEW_MEASURE_DEFAULTS)
    for c in ("PrgID", "ClaimYearQuarter", "IOU_AC_Territory", "AC_Climate_Zone", "AC_Elec_Target_Sector", "AC_Gas_Target_Sector", "Sector", "Version"):
        if like and like.get(c):
            row[c] = like[c]
    row["CEInputID"] = unique_id("NEW", existing_ids)
    return row


def unique_id(base: str, existing_ids) -> str:
    taken = {str(x) for x in existing_ids}
    stem = re.sub(r"_(copy|\d+)(_\d+)?$", "", str(base) or "NEW")
    for i in range(1, 100000):
        cand = f"{stem}_{i}"
        if cand not in taken:
            return cand
    raise EditorError("could not make a unique CEInputID")


def _pipe_lines(df: pd.DataFrame, what: str) -> list[str]:
    bad = sorted({c for c in df.columns if df[c].astype(str).str.contains("|", regex=False).any()})
    if bad:
        raise EditorError(f"{what}: CET files are '|'-delimited, so no cell may contain '|' (column(s) {bad})")
    return ["|".join(df.columns)] + ["|".join(map(str, r)) for r in df.itertuples(index=False, name=None)]


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    """Drop rows that are entirely blank (an empty row left in the grid) and strip stray whitespace around every value."""
    df = df.fillna("").astype(str).apply(lambda s: s.str.strip())
    return df[(df != "").any(axis=1)].reset_index(drop=True)


def kept_rows(df: pd.DataFrame | None) -> list[int]:
    """The positions of the rows write_cet_files keeps (all but the entirely blank ones): row i of the written file is row kept_rows(df)[i] of `df`."""
    if df is None or df.empty:
        return []
    df = df.fillna("").astype(str).apply(lambda s: s.str.strip())
    return [int(i) for i in np.flatnonzero((df != "").any(axis=1).to_numpy())]


def write_cet_files(measures: pd.DataFrame, programcost: pd.DataFrame | None, folder: str | Path) -> dict[str, Path | None]:
    """Measure.csv (and ProgramCost.csv if it has rows) in CET's pipe-delimited format. Returns {"measures": path, "programcost": path or None}."""
    out = Path(folder)
    out.mkdir(parents=True, exist_ok=True)
    m = _clean(conform(measures, MEASURE_COLUMNS))          # files always in CET's column order, whatever order the editor shows
    m = m.drop(columns=[c for c in templates.RECET_MEASURE_COLUMNS if c in m and (m[c] == "").all()])   # ReCET-only columns only when used
    from .fieldnames import to_cet
    m = to_cet(m)                                           # CET's own names (E3ClimateZone ...): the file is ready for the CET
    if m.empty:
        raise EditorError("the Measure table is empty")
    files: dict[str, Path | None] = {"measures": out / "Measure.csv", "programcost": None}
    files["measures"].write_text("\n".join(_pipe_lines(m, "Measure")) + "\n", encoding="utf-8")
    pc = _clean(conform(programcost, PROGRAMCOST_COLUMNS)) if programcost is not None else pd.DataFrame()
    if not pc.empty:
        files["programcost"] = out / "ProgramCost.csv"
        files["programcost"].write_text("\n".join(_pipe_lines(pc, "ProgramCost")) + "\n", encoding="utf-8")
    return files


def _cell_value(col: str, v: str, numeric: set[str]):
    """A number for a numeric column whose text is a plain number (so Excel shows numbers, not 'number stored as text'), else the text unchanged."""
    if col in numeric and v.strip():
        try:
            f = float(v)
        except ValueError:
            return v
        return int(f) if f.is_integer() and re.fullmatch(r"-?\d+(\.0*)?", v.strip()) else f
    return v


def write_workbook(measures: pd.DataFrame, programcost: pd.DataFrame | None, path: str | Path, base_workbook: str | None = None) -> Path:
    """A job workbook with the edited Measure and ProgramCost sheets. With `base_workbook`, every other sheet of it (RunParameters, Programs ...) is kept."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    out = Path(path)
    if base_workbook and Path(base_workbook).exists():
        wb = load_workbook(base_workbook)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    numeric = set(MEASURE_NUMERIC) | set(PROGRAMCOST_NUMERIC) | {"PrgYear"}
    for name, df in (("Measure", _clean(conform(measures, MEASURE_COLUMNS))),
                     ("ProgramCost", _clean(conform(programcost, PROGRAMCOST_COLUMNS)) if programcost is not None else None)):
        existing = next((s for s in wb.sheetnames if re.sub(r"[\s_]", "", s).lower() == name.lower()), None)
        idx = wb.sheetnames.index(existing) if existing else len(wb.sheetnames)
        if existing:
            wb.remove(wb[existing])
        ws = wb.create_sheet(name, idx)
        cols = list(df.columns) if df is not None else (MEASURE_COLUMNS if name == "Measure" else PROGRAMCOST_COLUMNS)
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        if df is not None:
            for r in df.itertuples(index=False, name=None):
                ws.append([_cell_value(c, v, numeric) for c, v in zip(cols, r)])
        ws.freeze_panes = "B2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def prepare_run_inputs(measures: pd.DataFrame, programcost: pd.DataFrame | None, out_dir: str | Path, base_workbook: str | None = None,
                       given: dict | None = None) -> tuple[dict, Path]:
    """Write the edited tables for a run into `<out_dir>/_edited_inputs/<timestamp>/` and return (RunConfig field updates, that folder).

    The run then reads the edited Measure.csv / ProgramCost.csv instead of the Job tab's. The other sheets of `base_workbook` (RunParameters, Programs, ...) are
    written out too and used for any input not chosen separately (`given`: RunConfig field -> value), so the run uses everything it would have used
    from the workbook except its Measure and ProgramCost sheets. With no ProgramCost rows the run has no program costs, even if the workbook had a sheet."""
    from . import workbook as wb_mod
    folder = Path(out_dir) / "_edited_inputs" / datetime.now().strftime("%Y%m%d_%H%M%S")
    files = write_cet_files(measures, programcost, folder)
    updates: dict = {"workbook": None, "measures": str(files["measures"]), "programcost": str(files["programcost"]) if files["programcost"] else None}
    if base_workbook and Path(base_workbook).exists():
        for fld, f in wb_mod.extract(base_workbook, folder / "workbook_sheets").items():
            if fld not in ("measures", "programcost") and not (given or {}).get(fld):
                updates[fld] = str(f)
    return updates, folder


# ------------------------------------------------------------------------------------------------------------------------- dropdown values
# Changes ReCET admins made to the shipped lists (the web app's Docs > Value lists page), layered over them by value_list():
# {list name: {code: {"kind": "added" | "retired" | "reinstated", "start_year", "end_year", "description"}}}. Empty in the desktop tool.
_VALUE_LIST_CHANGES: dict[str, dict[str, dict]] = {}


def set_value_list_changes(changes: dict[str, dict[str, dict]]) -> None:
    """Replace the admin changes layered over the shipped value lists (see `_VALUE_LIST_CHANGES`)."""
    global _VALUE_LIST_CHANGES
    _VALUE_LIST_CHANGES = {k: dict(v) for k, v in changes.items()}
    value_list.cache_clear()


@lru_cache(maxsize=None)
def cet_value_list(name: str) -> pd.DataFrame:
    """A CET value list as shipped: code, start_year, end_year (years as numbers, NaN when blank), description ("" when the list has none)."""
    p = paths.package_data("value_lists", f"{name}.csv")
    df = pd.read_csv(p, dtype=str, keep_default_na=False)
    for c in ("start_year", "end_year"):
        df[c] = pd.to_numeric(df[c], errors="coerce") if c in df.columns else float("nan")
    df["code"] = df["code"].str.strip()
    if "description" not in df.columns:
        df["description"] = ""
    return df[["code", "start_year", "end_year", "description"]]


@lru_cache(maxsize=None)
def value_list(name: str) -> pd.DataFrame:
    """The value list in use: the shipped CET list with any admin changes applied. Columns as `cet_value_list`, plus `source`
    ("CET", or the kind of admin change: "added", "retired", "reinstated")."""
    df = cet_value_list(name).copy()
    df["source"] = "CET"
    for code, ch in _VALUE_LIST_CHANGES.get(name, {}).items():
        hit = df["code"].str.lower() == str(code).strip().lower()
        kind = ch.get("kind")
        if kind == "added" and not hit.any():
            row = {"code": str(code).strip(), "start_year": ch.get("start_year"), "end_year": ch.get("end_year"),
                   "description": ch.get("description") or "", "source": "added"}
            df = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
            continue
        if kind in ("retired", "added"):
            df.loc[hit, "end_year"] = ch.get("end_year")
        elif kind == "reinstated":
            df.loc[hit, "end_year"] = float("nan")
        df.loc[hit, "source"] = kind
    for c in ("start_year", "end_year"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def active_codes(name: str, year: int | None = None) -> list[str]:
    """Codes active in `year` (start_year <= year, end_year blank or >= year); with no year, every code not yet retired (blank end_year)."""
    df = value_list(name)
    if year is None:
        keep = df["end_year"].isna()
    else:
        keep = (df["start_year"].isna() | (df["start_year"] <= year)) & (df["end_year"].isna() | (df["end_year"] >= year))
    return sorted(df.loc[keep, "code"].unique(), key=_natural)


def code_status(name: str, value: str, year: int | None) -> str | None:
    """None if `value` is an active code of the list for `year` (case-insensitive), else a short reason."""
    df = value_list(name)
    hit = df[df["code"].str.lower() == value.strip().lower()]
    if hit.empty:
        return "not in CET's value list"
    if year is None:
        return None
    ok = (hit["start_year"].isna() | (hit["start_year"] <= year)) & (hit["end_year"].isna() | (hit["end_year"] >= year))
    if ok.any():
        return None
    r = hit.iloc[0]
    whose = "the" if r["source"] == "added" else "CET's"
    if pd.notna(r["end_year"]) and r["end_year"] < year:
        if r["source"] == "retired":
            return f"retired by a ReCET admin after {int(r['end_year'])}"
        return f"retired in {whose} value list after {int(r['end_year'])}"
    return f"not valid in {whose} value list before {int(r['start_year'])}"


def _natural(s: str):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(s))]


def claim_year(cyq: str) -> int | None:
    s = str(cyq).strip()
    return int(s[:4]) if CYQ.match(s) else None


class ChoiceSource:
    """Dropdown values and profile lookups. Built once per window; `load_profiles()` (slow: reads the profile catalog) runs in a worker thread."""

    def __init__(self):
        self.catalog: pd.DataFrame | None = None          # end_use_name, iou_ac_territory, _sector_key: profiles WITH hourly data
        self.catalog_note = "profile library not loaded yet"
        self.user_8760 = None                             # overrides.UserProfiles
        self.gas_names: set[str] = set()
        self.user_gas: set[str] = set()
        self._by_iou_sector: dict[tuple[str, str], list[str]] = {}
        self._by_iou: dict[str, list[str]] = {}
        self.catalog_available = False

    @property
    def profiles_loaded(self) -> bool:
        return self.catalog is not None

    def load_profiles(self, profiles_8760_file: str | None = None, profiles_gas_file: str | None = None) -> None:
        from . import overrides
        from .pipeline import gas_savings_profiles, real_impact_profiles
        try:
            if real_impact_profiles.is_available():
                cat = real_impact_profiles._catalog_hourly()[["end_use_name", "iou_ac_territory", "_sector_key"]].copy()
                self.catalog = cat
                self.catalog_note = f"{cat['end_use_name'].nunique():,} end uses with 8,760 profiles in the library"
            else:
                self.catalog = pd.DataFrame(columns=["end_use_name", "iou_ac_territory", "_sector_key"])
                self.catalog_note = "profile library not found (Reference data tab): end-use choices are CET's value list and are not checked"
        except Exception as exc:                       # an unreadable library must not break editing
            self.catalog = pd.DataFrame(columns=["end_use_name", "iou_ac_territory", "_sector_key"])
            self.catalog_note = f"profile library unreadable ({exc}); end uses are not checked"
        self.catalog_available = not self.catalog.empty
        for (iou, sec), g in self.catalog.groupby(["iou_ac_territory", "_sector_key"]):
            self._by_iou_sector[(iou, sec)] = sorted(g["end_use_name"].unique(), key=_natural)
        for iou, g in self.catalog.groupby("iou_ac_territory"):
            self._by_iou[iou] = sorted(g["end_use_name"].unique(), key=_natural)
        try:
            if gas_savings_profiles.PROFILES_CSV.exists():
                self.gas_names = set(gas_savings_profiles._table()["Profile"].astype(str))
        except Exception:
            self.gas_names = set()
        self.user_8760 = None
        if profiles_8760_file and Path(profiles_8760_file).exists():
            try:
                self.user_8760 = overrides.load_user_profiles(profiles_8760_file)
            except Exception:
                self.user_8760 = None
        self.user_gas = set()
        if profiles_gas_file and Path(profiles_gas_file).exists():
            try:
                self.user_gas = set(overrides.load_user_gas_profiles(profiles_gas_file)[0])
            except Exception:
                self.user_gas = set()

    # --- dropdown values
    def column_choices(self, table: str, column: str, years: list[int] | None = None) -> list[str] | None:
        """Values offered for a whole column (None = free text). Year-limited value lists offer every code active in any of `years`."""
        if table == "ProgramCost":
            if column == "PA":
                return active_codes("pa")
            return None
        if column == "IOU_AC_Territory":
            return sorted(preflight.VALID_PA)
        if column in FLAG_COLUMNS:
            return ["0", "1"]
        if column == "AC_Gas_Impact_Profile":
            return sorted(self.gas_names | self.user_gas | set(active_codes("ac_gas_impact_profile")), key=_natural)
        if column in END_USE_COLUMNS:
            return None                                    # per row: see end_use_choices
        if column == "ACCVersion":
            from . import versions
            return list(versions.BUILTIN)
        name = VALUE_LISTS.get(column)
        if name is None:
            return None
        if not years:
            return active_codes(name)
        out: set[str] = set()
        for y in sorted(set(years)):
            out.update(active_codes(name, y))
        return sorted(out, key=_natural)

    def end_use_choices(self, iou: str, sector: str) -> list[str]:
        """End uses that have an 8,760 profile for this IOU and target sector (falling back to the IOU's, then to CET's value list), plus supplied profiles."""
        iou = str(iou).strip().upper()
        out = list(self._by_iou_sector.get((iou, str(sector).strip().upper()), []) or self._by_iou.get(iou, []))
        if not out and not getattr(self, "catalog_available", False):
            out = active_codes("ac_elec_impact_profile")
        if self.user_8760 is not None:
            out = sorted(set(out) | {e.end_use for e in self.user_8760.entries}, key=_natural)
        return out

    # --- profile checks
    def electric_profile_status(self, end_use: str, iou: str, sector: str) -> str | None:
        """None if the run will find an 8,760 profile; else why not. Not checked (None) while the library is not loaded or not present."""
        if not getattr(self, "catalog_available", False) and self.user_8760 is None:
            return None
        if self.user_8760 is not None and self.user_8760.match(end_use, iou, sector) is not None:
            return None
        if getattr(self, "catalog_available", False):
            from .pipeline import real_impact_profiles
            pid, exact = real_impact_profiles.match_profile_id(end_use, str(iou).strip().upper(), sector)
            if pid is not None:
                return None if exact else "INFO:matched the library on end use and IOU only (no profile for this target sector)"
            return "no 8,760 profile for this end use, IOU and target sector: the measure's electric savings will NOT be priced"
        return None

    def gas_profile_status(self, name: str) -> str | None:
        if not self.gas_names and not self.user_gas:
            return None
        if name in self.user_gas or name in self.gas_names:
            return None
        return "not a known gas monthly profile: the measure's gas savings will NOT be priced"


# ------------------------------------------------------------------------------------------------------------------------- checks
@dataclass(frozen=True)
class CellIssue:
    table: str          # "Measure" | "ProgramCost"
    row: int            # 0-based row in that table
    column: str
    severity: str       # "error" (the run would stop) | "warning" (it runs, but CET would reject it or the value is not priced) | "info"
    message: str


def _num(v: str):
    v = str(v).strip()
    if not v:
        return None
    try:
        return float(v)
    except ValueError:
        return "bad"


def check_tables(measures: pd.DataFrame, programcost: pd.DataFrame | None, choices: ChoiceSource | None = None) -> list[CellIssue]:
    """Every problem cell. Mirrors preflight.check_inputs's rules (same constants), cell by cell; see the module docstring."""
    out: list[CellIssue] = []
    m = measures.fillna("").astype(str)
    pc = programcost.fillna("").astype(str) if programcost is not None else pd.DataFrame(columns=PROGRAMCOST_COLUMNS)
    blank_row = (m.apply(lambda s: s.str.strip()) == "").all(axis=1) if len(m) else pd.Series(dtype=bool)
    pc_blank_row = (pc.apply(lambda s: s.str.strip()) == "").all(axis=1) if len(pc) else pd.Series(dtype=bool)

    def add(table, row, col, sev, msg):
        out.append(CellIssue(table, int(row), col, sev, msg))

    ranges = {c: (lo, hi, lo_open, hi_open) for c, lo, hi, lo_open, hi_open in preflight.CET_RANGES}
    ids = m["CEInputID"].str.strip()
    dup_ids = set(ids[ids.duplicated(keep=False) & (ids != "")])
    pc_ids = set(pc.loc[~pc_blank_row, "PrgID"].str.strip()) if len(pc) else set()
    m_prg = set(m.loc[~blank_row, "PrgID"].str.strip()) if len(m) else set()

    for i, r in m.iterrows():
        if blank_row.iloc[i]:
            continue
        g = lambda c: str(r.get(c, "")).strip()
        year = claim_year(g("ClaimYearQuarter"))
        if not g("CEInputID"):
            add("Measure", i, "CEInputID", "error", "CEInputID is blank")
        elif g("CEInputID") in dup_ids:
            add("Measure", i, "CEInputID", "error", "CEInputID must be unique (results are keyed on it)")
        if not g("PrgID"):
            add("Measure", i, "PrgID", "error", "PrgID is blank")
        elif len(pc_ids) and g("PrgID") not in pc_ids:
            add("Measure", i, "PrgID", "error", "no ProgramCost row has this PrgID: with program costs, every measure's program needs a row (else the run stops)")
        if year is None:
            add("Measure", i, "ClaimYearQuarter", "error", "ClaimYearQuarter must look like 2028Q1")
        if g("IOU_AC_Territory") not in preflight.VALID_PA:
            add("Measure", i, "IOU_AC_Territory", "error", f"IOU_AC_Territory must be one of {sorted(preflight.VALID_PA)}")
        vals = {}
        for c in MEASURE_NUMERIC:
            if c not in m.columns:
                continue
            v = _num(g(c))
            if v == "bad":
                add("Measure", i, c, "error", f"{c} must be a number")
                v = None
            vals[c] = v
            if v is not None and c in ranges:
                lo, hi, lo_open, hi_open = ranges[c]
                if (v <= lo if lo_open else v < lo) or (v >= hi if hi_open else v > hi):
                    rng = f"{'(' if lo_open else '['}{lo}, {hi if hi != float('inf') else 'no upper limit'}{')' if hi_open else ']'}"
                    add("Measure", i, c, "warning", f"outside CET's allowed range {rng}: CET would reject this upload (the POC still prices it)")
        for c in preflight.NEEDED_VALUES:
            if not g(c):
                add("Measure", i, c, "error", f"{c} is blank (CET requires a value; enter 0 if there is none)")
        rul, eul = vals.get("RUL_Yrs"), vals.get("EUL_Yrs")
        for c, v in (("EUL_Yrs", eul), ("RUL_Yrs", rul)):
            if v is not None and v < 0:
                add("Measure", i, c, "error", f"{c} is negative (CET allows 0 or more)")
        if rul is not None and eul is not None and rul > eul:
            add("Measure", i, "RUL_Yrs", "error", "RUL_Yrs is greater than EUL_Yrs: the remaining useful life of the existing equipment cannot exceed the measure's life")
        elif rul is not None and rul > 0 and eul is not None and rul == eul:
            add("Measure", i, "RUL_Yrs", "warning", "RUL_Yrs equals EUL_Yrs: the measure is treated as single-baseline for its whole life")
        for c in FLAG_COLUMNS:
            if g(c) and g(c).lower() not in ("0", "1", "true", "false"):
                add("Measure", i, c, "warning", f"{c} should be 0 or 1")
        for c, name in VALUE_LISTS.items():
            if c in ("AC_Elec_Impact_Profile", "AC_Gas_Impact_Profile") or not g(c):
                continue
            why = code_status(name, g(c), year)
            if why:
                add("Measure", i, c, "warning", f"'{g(c)}' is {why}: CET would reject this upload")
        has_elec = any((vals.get(c) or 0) != 0 for c in ("UnitkWh1stBaseline", "UnitkWh2ndBaseline", "UnitkW1stBaseline", "UnitkW2ndBaseline",
                                                          "UnitkWhIOUWater1stBaseline", "UnitkWhIOUWater2ndBaseline"))
        has_gas = any((vals.get(c) or 0) != 0 for c in ("UnitTherm1stBaseline", "UnitTherm2ndBaseline"))
        if has_elec:
            if not g("AC_Elec_Impact_Profile"):
                add("Measure", i, "AC_Elec_Impact_Profile", "warning", "blank end use on a measure with electric savings: its electric savings will NOT be priced")
            elif choices is not None:
                why = choices.electric_profile_status(g("AC_Elec_Impact_Profile"), g("IOU_AC_Territory"), g("AC_Elec_Target_Sector"))
                if why:
                    sev, msg = ("info", why[5:]) if why.startswith("INFO:") else ("warning", why)
                    add("Measure", i, "AC_Elec_Impact_Profile", sev, msg)
        has_kwh2 = rul is not None and rul > 0 and (eul is None or rul < eul) and (vals.get("UnitkWh2ndBaseline") or 0) != 0
        has_water = any((vals.get(c) or 0) != 0 for c in ("UnitkWhIOUWater1stBaseline", "UnitkWhIOUWater2ndBaseline"))
        for c, used, what in (("AC_Elec_Impact_Profile_2nd_Baseline", has_kwh2, "an early replacement's 2nd-baseline kWh"), ("AC_Elec_Impact_Profile_Water", has_water, "water energy"),
                              ("AC_Elec_Impact_Profile_Parent", True, "electric savings"),
                              ("AC_Elec_Impact_Profile_Parent_2nd_Baseline", has_kwh2, "an early replacement's 2nd-baseline kWh")):
            if not g(c):
                continue
            if not used:
                add("Measure", i, c, "info", f"not used: the measure has no {what}")
            elif choices is not None:
                why = choices.electric_profile_status(g(c), g("IOU_AC_Territory"), g("AC_Elec_Target_Sector"))
                if why:
                    sev, msg = ("info", why[5:]) if why.startswith("INFO:") else ("warning", why.replace("electric savings", what))
                    add("Measure", i, c, sev, msg)
        if g("ACCVersion"):
            from . import versions
            key = re.sub(r"(?i)\s*ACC$", "", g("ACCVersion"))
            if key not in versions.BUILTIN:
                add("Measure", i, "ACCVersion", "warning", f"'{g('ACCVersion')}' is not a built-in ACC version ({', '.join(versions.BUILTIN)}): "
                    "the run stops at its input check unless it is the run's own version")
        if has_gas:
            if g("AC_Gas_Target_Sector").upper() not in preflight.VALID_GAS_SECTORS:
                add("Measure", i, "AC_Gas_Target_Sector", "error", "AC_Gas_Target_Sector must be Residential, Commercial or Core on a measure with gas savings")
            if not g("AC_Gas_Impact_Profile"):
                add("Measure", i, "AC_Gas_Impact_Profile", "warning", "blank gas profile on a measure with gas savings: its gas savings will NOT be priced")
            elif choices is not None:
                why = choices.gas_profile_status(g("AC_Gas_Impact_Profile"))
                if why:
                    add("Measure", i, "AC_Gas_Impact_Profile", "warning", why)
        for c in m.columns:
            if "|" in g(c):
                add("Measure", i, c, "error", "a cell may not contain '|' (CET files are '|'-delimited)")

    if len(pc):
        keys = pc["PrgID"].str.strip() + "|" + pc["ClaimYearQuarter"].str.strip()
        dup_keys = set(keys[keys.duplicated(keep=False) & ~pc_blank_row])
        for i, r in pc.iterrows():
            if pc_blank_row.iloc[i]:
                continue
            g = lambda c: str(r.get(c, "")).strip()
            if not g("PrgID"):
                add("ProgramCost", i, "PrgID", "error", "PrgID is blank")
            elif g("PrgID") not in m_prg:
                add("ProgramCost", i, "PrgID", "error", "no measure has this PrgID: its costs would never be charged (the run stops on it)")
            if claim_year(g("ClaimYearQuarter")) is None:
                add("ProgramCost", i, "ClaimYearQuarter", "error", "ClaimYearQuarter must look like 2028Q4 (its year is when the cost is incurred)")
            elif keys.iloc[i] in dup_keys:
                add("ProgramCost", i, "ClaimYearQuarter", "warning", "another row has the same PrgID and ClaimYearQuarter (both are summed)")
            if g("PrgYear") and (_num(g("PrgYear")) == "bad"):
                add("ProgramCost", i, "PrgYear", "error", "PrgYear must be a year")
            elif g("PrgYear") and claim_year(g("ClaimYearQuarter")) and int(float(g("PrgYear"))) != claim_year(g("ClaimYearQuarter")):
                add("ProgramCost", i, "PrgYear", "info", "PrgYear differs from the ClaimYearQuarter's year (the run uses ClaimYearQuarter)")
            pa = g("PA").upper()
            if not pa:
                add("ProgramCost", i, "PA", "error", "PA is blank (it selects the discount rate for this row's costs)")
            elif pa not in preflight.VALID_PA and pa != "SW":
                add("ProgramCost", i, "PA", "warning", f"the POC discounts program costs at an IOU's rate: PA should be one of {sorted(preflight.VALID_PA)} or SW")
            for c in PROGRAMCOST_NUMERIC:
                if c in pc.columns and _num(g(c)) == "bad":
                    add("ProgramCost", i, c, "error", f"{c} must be a number (blank = 0)")
            for c in pc.columns:
                if "|" in g(c):
                    add("ProgramCost", i, c, "error", "a cell may not contain '|' (CET files are '|'-delimited)")
    return out


# ------------------------------------------------------------------------------------------------------------------------- field help
@lru_cache(maxsize=1)
def field_help() -> dict[str, str]:
    """{field: 'type - description'} from docs/INPUT_FIELDS.md's Measure and ProgramCost tables (shown under the grid for the selected column)."""
    p = Path(__file__).resolve().parent / "docs" / "INPUT_FIELDS.md"
    out: dict[str, str] = {}
    if not p.exists():
        return out
    section = None
    for line in p.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            section = line
        if not (section and ("Measure.csv" in section or "ProgramCost.csv" in section)):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")] if line.startswith("|") else []
        if len(cells) >= 3 and cells[0].startswith("`"):
            text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", cells[-1]).replace("`", "").replace("**", "")
            kind = cells[1].replace("`", "")
            for name in re.findall(r"`([^`]+)`", cells[0]):
                out.setdefault(name, f"{kind}  -  {text}")
    return out
