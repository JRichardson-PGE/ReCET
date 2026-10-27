"""Preflight check of a CET-format input job, before any pricing.

Reads Measure.csv (and ProgramCost.csv) exactly as the run would and reports, without pricing anything:

* format problems that would stop the run (no recognisable delimiter, missing columns, duplicate CEInputID, bad ClaimYearQuarter, unknown IOU, non-numeric values),
* what the run will and will not be able to price: for every measure with electric savings whether an 8,760 profile (user-supplied or CEDARS library) or a TOU
  H-factor profile exists, for every measure with gas savings whether its monthly profile exists,
* the ACC horizon: claim years before the first ACC year or measure lives that run past the last one,
* PrgID alignment between Measure.csv and ProgramCost.csv, and statewide-program rows,
* optionally (`filing_rules`), CEDARS' own upload checks for a claims or a budget filing (filing_rules.py): the records CEDARS would reject or warn on, each
  listed under the rule's ID, and kept on the report (`PreflightReport.filing`) for the run to flag in the measure-level table. They never stop the run.

Severity: "error" = the run cannot proceed as is; "warning" = the run proceeds but some measures are unpriced or priced on a fallback; "info" = worth knowing.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from .pipeline import delimited

REQUIRED_MEASURE_COLUMNS = [
    "CEInputID", "PrgID", "MeasDescription", "ClaimYearQuarter", "Sector", "AC_Climate_Zone", "AC_Gas_Impact_Profile", "AC_Gas_Target_Sector", "AC_Elec_Impact_Profile",
    "AC_Elec_Target_Sector", "IOU_AC_Territory", "EUL_Yrs", "RUL_Yrs", "NumUnits",
    "UnitkW1stBaseline", "UnitkWh1stBaseline", "UnitTherm1stBaseline", "UnitkW2ndBaseline", "UnitkWh2ndBaseline", "UnitTherm2ndBaseline",
    "NTGRkW", "NTGRkWh", "NTGRTherm", "NTGRCost", "RealizationRatekW", "RealizationRatekWh", "RealizationRateTherm",
    "InstallationRatekW", "InstallationRatekWh", "InstallationRateTherm",
    "UnitMeaCost1stBaseline", "UnitMeaCost2ndBaseline", "UnitIncentiveToOthers", "UnitEndUserRebate", "UnitDirectInstallLab", "UnitDirectInstallMat",
    "MeasInflation", "UnitGasInfraBens", "UnitRefrigCosts", "UnitRefrigBens", "UnitMiscCosts", "UnitMiscBens", "UnitTaxCredits",
]
OPTIONAL_MEASURE_COLUMNS = {
    "FuelID": "fuel-substitution measures are recognised by FuelID or MeasImpactType; without both, none are",
    "MeasImpactType": "fuel-substitution measures are recognised by FuelID or MeasImpactType; without both, none are",
    "CombustionType": "gas NOx emissions need it (they are 0 without)",
    "UnitkWhIOUWater1stBaseline": "water energy benefits/costs are 0 without it and UnitkWhIOUWater2ndBaseline",
    "UnitkWhIOUWater2ndBaseline": "water energy benefits/costs are 0 without it and UnitkWhIOUWater1stBaseline",
}
NUMERIC_COLUMNS = [c for c in REQUIRED_MEASURE_COLUMNS if c.startswith(("Unit", "NTGR", "RealizationRate", "InstallationRate"))
                   or c in ("EUL_Yrs", "RUL_Yrs", "NumUnits", "MeasInflation")]
VALID_PA = {"PGE", "SCE", "SDGE", "SCG"}
# CET's own QC ranges (cet_spec.sql, `CedarsNumRange`): a value outside them would make CET reject the whole upload. Here they are warnings -- the tool still prices the row.
# (column, low, high, low is exclusive, high is exclusive)
CET_RANGES = [("NTGRkW", 0, 1.5, True, False), ("NTGRkWh", 0, 1.5, True, False), ("NTGRTherm", 0, 1.5, True, False), ("NTGRCost", 0, 1.5, True, False),
              ("EUL_Yrs", 0, 40, False, False), ("RUL_Yrs", 0, 30, False, False),
              ("RealizationRatekW", 0, 2, True, False), ("RealizationRatekWh", 0, 2, True, False), ("RealizationRateTherm", 0, 2, True, False),
              ("InstallationRatekW", 0, 1, True, False), ("InstallationRatekWh", 0, 1, True, False), ("InstallationRateTherm", 0, 1, True, False),
              ("MeasInflation", 0, 1, False, False), ("UnitMeaCost1stBaseline", 0, float("inf"), False, False), ("UnitMeaCost2ndBaseline", 0, float("inf"), False, False)]
# Values the arithmetic cannot do without: a blank turns the measure's results blank (CET marks each of these NotEmpty)
NEEDED_VALUES = ("NumUnits", "EUL_Yrs", "NTGRkWh", "UnitkWh1stBaseline")
VALID_GAS_SECTORS = {"RESIDENTIAL", "COMMERCIAL", "CORE"}
MAX_CELLS = 5000          # cells listed per issue (the web editor highlights them)


@dataclass
class Issue:
    severity: str
    code: str
    message: str
    count: int = 0
    examples: list = field(default_factory=list)
    cells: list = field(default_factory=list)     # (table, row, column): where it is, row = 0-based data row of Measure.csv / ProgramCost.csv
    impact: dict = field(default_factory=dict)    # {quantity: total} of the records it names, e.g. {"first-year net kWh": 12345.6}; parts add up

    def impact_text(self) -> str:
        """'12,346 first-year net kWh; 2,100 first-year net therms' (blank when there is no impact)."""
        return "; ".join(f"{v:,.0f} {k}" for k, v in self.impact.items())


IMPACT_KWH, IMPACT_THERMS = "first-year net kWh", "first-year net therms"


@dataclass
class PreflightReport:
    issues: list[Issue] = field(default_factory=list)
    stats: dict = field(default_factory=dict)
    agg: dict = field(default_factory=dict)          # raw values behind some stats, for combining the reports of a file checked in parts (not reported)
    filing: pd.DataFrame | None = None               # with filing_rules: one row per failing record (table, id, FilingValidation, FilingRules)
    filing_rules: dict | None = None                 # with filing_rules: the rule set used (filing_rules.RuleSet.describe)

    def add(self, severity: str, code: str, message: str, count: int = 0, examples=None, cells=None, impact=None) -> None:
        self.issues.append(Issue(severity, code, message, count, list(examples or [])[:5], list(cells or [])[:MAX_CELLS],
                                 {k: float(v) for k, v in (impact or {}).items()}))

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.severity == "error"]

    @property
    def ok(self) -> bool:
        return not self.errors

    def text(self) -> str:
        lines = []
        for sev in ("error", "warning", "info"):
            for i in (x for x in self.issues if x.severity == sev):
                n = f" [{i.count:,}]" if i.count else ""
                ex = f"  e.g. {i.examples}" if i.examples else ""
                imp = f"  ({i.impact_text()})" if i.impact else ""
                lines.append(f"{sev.upper():8s}{i.code}{n}: {i.message}{imp}{ex}")
        head = ["PREFLIGHT: " + ("no blocking problems" if self.ok else f"{len(self.errors)} blocking problem(s)")]
        head += [f"  {k}: {v}" for k, v in self.stats.items()]
        return "\n".join(head + [""] + (lines or ["(nothing to report)"]))

    def to_dict(self) -> dict:
        """The report as plain JSON values (the web app's Validate inputs panel reads it from preflight.json)."""
        return {"ok": self.ok, "errors": len(self.errors), "stats": {str(k): str(v) for k, v in self.stats.items()}, "filing_rules": self.filing_rules,
                "issues": [{"severity": i.severity, "code": i.code, "message": i.message, "count": int(i.count or 0), "examples": [str(e) for e in i.examples],
                            "cells": [[str(tb), int(r), str(c)] for tb, r, c in i.cells], "impact": dict(i.impact)}
                           for i in self.issues]}

    def write_xlsx(self, path: Path, measure_ids=None, program_ids=None) -> Path:
        """The report as a workbook a reviewer can open and filter: Summary (the result and the statistics), Findings (one row per finding, errors first)
        and Cells (one row per flagged cell: its file, line, record ID and column). `measure_ids` / `program_ids` give each data row's CEInputID / PrgID + PrgYear,
        in file order (blank when not given)."""
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        order = {"error": 0, "warning": 1, "info": 2}
        issues = sorted(self.issues, key=lambda i: order.get(i.severity, 3))
        counts = {s: sum(i.severity == s for i in issues) for s in order}
        files = {"Measure": "Measure.csv", "ProgramCost": "ProgramCost.csv"}
        ids = {"Measure": measure_ids, "ProgramCost": program_ids}

        def rid(table, row):
            s = ids.get(table)
            return "" if s is None or not 0 <= row < len(s) else str(s[row])

        summary = [("Result", "no blocking problems: the run can price" if self.ok else f"{counts['error']} blocking problem(s): the run stops before pricing"),
                   ("Errors (stop the run)", counts["error"]), ("Warnings (the run goes on; some records are priced on a fallback or not at all)", counts["warning"]),
                   ("Information", counts["info"]), *((str(k).strip(), str(v)) for k, v in self.stats.items()),
                   ("Cells sheet", f"the flagged cells, up to {MAX_CELLS:,} per finding; Line is the line of the file (line 1 is the header)")]
        findings = [(i.severity.upper(), i.code, i.count or None, i.message, i.impact.get(IMPACT_KWH), i.impact.get(IMPACT_THERMS),
                     "; ".join(map(str, i.examples))) for i in issues]
        cells = [(i.severity.upper(), i.code, files.get(t, t), int(r) + 2, rid(t, int(r)), c) for i in issues for t, r, c in i.cells]
        wb = Workbook()
        head_font, head_fill = Font(name="Arial", bold=True, color="FFFFFF", size=10), PatternFill("solid", fgColor="1F3864")
        sheets = (("Summary", ("Item", "Value"), summary, (60, 90)),
                  ("Findings", ("Severity", "Code", "Records", "Message", "First-year net kWh", "First-year net therms", "Examples"), findings,
                   (10, 28, 10, 100, 16, 16, 60)),
                  ("Cells", ("Severity", "Code", "File", "Line", "Record", "Column"), cells, (10, 28, 16, 9, 40, 28)))
        for k, (name, header, rows, widths) in enumerate(sheets):
            ws = wb.active if k == 0 else wb.create_sheet()
            ws.title = name
            ws.append(list(header))
            for cell in ws[1]:
                cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(vertical="center")
            for row in rows:
                ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
            for j, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(j)].width = w
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
            if name == "Findings":
                for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
        path = Path(path)
        wb.save(path)
        return path

    def to_frame(self) -> pd.DataFrame:
        return pd.DataFrame([{"severity": i.severity, "code": i.code, "count": i.count, "message": i.message, "examples": "; ".join(map(str, i.examples))}
                             for i in self.issues])


def _nonblank_numeric_problems(df: pd.DataFrame, col: str) -> pd.Series:
    raw = df[col]
    try:                                            # the usual case, every value a number: one vectorised cast instead of to_numeric's slow path on Arrow strings
        import pyarrow as pa
        import pyarrow.compute as pc
        cast = pc.cast(pa.array(raw, type=pa.string(), from_pandas=True), pa.float64())
        if not pc.any(pc.is_nan(cast)).as_py():     # a literal "nan" casts to NaN; to_numeric counts it a problem, so leave it to the exact check
            return pd.Series(False, index=raw.index)
    except Exception:                               # some value did not cast (or the column is not text): find which, exactly as before
        pass
    parsed = pd.to_numeric(raw, errors="coerce")
    return raw.notna() & parsed.isna() & (raw.astype(str).str.strip() != "")


def _cells(table: str, rows, col: str) -> list[tuple[str, int, str]]:
    """(table, row, column) for the rows where a mask is true (or for a list of row numbers)."""
    rows = np.flatnonzero(np.asarray(rows)) if np.asarray(rows).dtype == bool else rows
    return [(table, int(r), col) for r in sorted(set(int(x) for x in rows))]


def _check_programs(rep: PreflightReport, measures: pd.DataFrame, programcost_csv, allocate_statewide: bool, apply_exclude_from_ce: bool,
                    mrows=lambda mask: []) -> None:
    """The measures' programs against the program metadata (programs.py): unknown PrgIDs, claims outside the program's years, closed programs, programs
    flagged Exclude_From_CE, and a ProgramCost PA that differs from the program list's PA. None of these stops the run. `mrows(mask)` turns a mask over
    `measures` into Measure.csv row numbers."""
    from . import outputs, programs
    meta = programs.lookup(measures["prg_id"])
    pid = measures["prg_id"].astype(str).str.strip().to_numpy()
    year = measures["claim_year"].to_numpy()
    missing = meta["Prg_MatchedOn"].to_numpy() == "(not found)"
    rep.agg["programs"] = (set(pid), set(pid[missing]))
    rep.stats["programs"] = f"{len(set(pid)):,} ({len(set(pid[missing])):,} not in the program metadata)"
    if missing.any():
        rep.add("warning", "program_not_in_metadata", "PrgID is not in the program metadata (the built-in CEDARS Programs list or the supplied one): its Prg_ columns are "
                "blank and its measures count as segment '(blank)'", int(missing.sum()), sorted(set(pid[missing])), _cells("Measure", mrows(missing), "PrgID"))
    start = pd.to_numeric(meta.get("Prg_StartYear"), errors="coerce").to_numpy() if "Prg_StartYear" in meta else None
    end = pd.to_numeric(meta.get("Prg_EndYear"), errors="coerce").to_numpy() if "Prg_EndYear" in meta else None
    if start is not None and end is not None:
        out = (~np.isnan(start) & (year < start)) | (~np.isnan(end) & (year > end))
        if out.any():
            rep.add("warning", "claim_outside_program_years", "claim year is before the program's StartYear or after its EndYear in the program metadata",
                    int(out.sum()), sorted(set(pid[out])), _cells("Measure", mrows(out), "ClaimYearQuarter"))
    if "Prg_ProgramStatus" in meta:
        closed = meta["Prg_ProgramStatus"].str.lower().isin({"closed", "closure_pending"}).to_numpy()
        if closed.any():
            rep.add("info", "program_closed", "the program metadata lists the program as closed or closure_pending", int(closed.sum()), sorted(set(pid[closed])),
                    _cells("Measure", mrows(closed), "PrgID"))
    if outputs.EXCLUDE_FLAG in meta:
        flag = programs.is_true(meta[outputs.EXCLUDE_FLAG]).to_numpy()
        if flag.any():
            what = ("left out of the portfolio summary (shown in its 'excluded' row)" if apply_exclude_from_ce
                    else "still in the portfolio summary; tick 'Leave Exclude_From_CE programs out of the portfolio summary' to leave them out")
            rep.add("info", "program_excluded_from_ce", f"the program metadata flags the program Exclude_From_CE: its measures are priced and {what}",
                    int(flag.sum()), sorted(set(pid[flag])), _cells("Measure", mrows(flag), "PrgID"))
    if programcost_csv:
        credited = outputs.credited_iou(measures, programcost_csv, allocate_statewide).to_numpy()
        listed = meta["Prg_PA"].to_numpy()
        differ = (credited != "") & ~missing & (credited != listed)
        if differ.any():
            rep.add("info", "program_pa_differs", "the program's PA in ProgramCost.csv (CreditedIOU) differs from its PA in the program metadata", int(differ.sum()),
                    sorted({f"{p}: {c} vs {l}" for p, c, l in zip(pid[differ], credited[differ], listed[differ])}), _cells("Measure", mrows(differ), "PrgID"))


@dataclass
class _Part:
    """How to check one part of a Measure.csv that is checked in parts: the whole file's column types, and the ProgramCost PrgIDs (as the run sees them) to
    mark this part's measures that have no program costs (None without a ProgramCost file)."""
    dtypes: dict
    program_ids: set | None


@dataclass
class _Filing:
    """CEDARS' upload checks to apply (filing_rules.py): the rule set and the filing year (None: each record's own claim year, and for ProgramCost the
    earliest claim year of the job)."""
    rules: object
    year: int | None


def check_inputs(measure_csv: str | Path, programcost_csv: str | Path | None = None, *, acc=None, allocate_statewide: bool = False,
                 base_year: int | None = None, apply_exclude_from_ce: bool = False, chunk_size: int | None = None,
                 parts: list | None = None, dtypes: dict | None = None, filing_rules: str | None = None,
                 filing_year: int | None = None) -> PreflightReport:
    """`acc` is a versions.ResolvedAcc (for the horizon and E3-version checks); the runtime must already be configured for it (runtime.apply) so that profile
    and TOU coverage reflect this run's user-supplied profiles.

    A Measure.csv with more than `chunk_size` rows is checked in parts of that size, so the check never holds the whole file: `parts` / `dtypes` are the
    file already split (delimited.split_file, delimited.part_dtypes; a run splits once and passes them), else it is split here into a temporary folder. The
    checks across rows (duplicate CEInputIDs, PrgID alignment with ProgramCost.csv) are made on the whole file, and the counts and stats are combined, so the
    report is the one the whole file gives, except that a part's checks still run when another part has a blocking problem.

    `filing_rules` ("claims" or "budget_filing") also applies CEDARS' upload checks for that kind of filing (filing_rules.py); `filing_year` is the filing's
    year for the rules that name it (default: each record's claim year, and the job's earliest claim year for ProgramCost)."""
    from . import filing_rules as fr
    filing = _Filing(fr.rule_set(filing_rules), filing_year) if filing_rules else None
    p = Path(measure_csv)
    if parts is None and chunk_size and p.exists() and delimited.detect_sep(p) and delimited.count_rows(p) > chunk_size:
        import tempfile
        with tempfile.TemporaryDirectory(prefix="recet_preflight_") as tmp:
            parts = delimited.split_file(p, chunk_size, tmp)
            dtypes = delimited.part_dtypes(parts)[0]
            rep = _check_parts(p, parts, dtypes, programcost_csv, acc, allocate_statewide, base_year, apply_exclude_from_ce, filing)
    elif parts is not None and len(parts) > 1:
        rep = _check_parts(p, parts, dtypes or {}, programcost_csv, acc, allocate_statewide, base_year, apply_exclude_from_ce, filing)
    else:
        rep = _check(p, programcost_csv, acc=acc, allocate_statewide=allocate_statewide, base_year=base_year, apply_exclude_from_ce=apply_exclude_from_ce,
                     filing=filing)
    if filing is not None:
        _finish_filing(rep, filing)
    return rep


def _filing_check(rep: PreflightReport, table: str, df: pd.DataFrame, ids: pd.Series, filing: _Filing, year: int | None) -> None:
    """CEDARS' upload checks on one table (or one part of Measure.csv): an issue per failing rule, and the failing records kept in rep.agg["filing"]."""
    from . import filing_rules as fr
    from .fieldnames import recet_name
    from .pipeline import real_impact_profiles
    results, notes = fr.check_frame(filing.rules, table, df, fr.Context(year, real_impact_profiles.is_available()))
    for res in results:
        reject = res.rule.severity == "reject"
        rows = np.flatnonzero(res.fail)
        col = {"combo_value_list": "AC_Elec_Impact_Profile"}.get(res.rule.spec.get("rule"), recet_name(res.rule.field.split("+")[0].strip()))
        eg = [f"{ids.iloc[i]}: {'(blank)' if res.actual is None or pd.isna(res.actual.iloc[i]) else res.actual.iloc[i]}" for i in rows[:5]]
        rep.add("warning", "cet_reject" if reject else "cet_warning",
                f"{res.message()} -- CEDARS {'rejects' if reject else 'warns on'} a {table} record that fails it", len(rows), eg, _cells(table, rows, col))
    for n in notes:
        rep.add("info", "cet_rule_not_checked", n)
    rep.agg.setdefault("filing", []).append(fr.flags(results, ids).assign(table=table))


def _filing_programcost(rep: PreflightReport, programcost_csv, filing: _Filing, first_claim_year: int | None) -> None:
    if programcost_csv is None or not Path(programcost_csv).exists():
        return
    pc = delimited.read_cet_csv(programcost_csv, dtype=str, keep_default_na=False)
    ids = (pc["PrgID"].astype(str) + " " + pc["PrgYear"].astype(str)) if {"PrgID", "PrgYear"} <= set(pc.columns) else pd.Series(range(len(pc))).astype(str)
    _filing_check(rep, "ProgramCost", pc, ids, filing, filing.year or first_claim_year)


def _finish_filing(rep: PreflightReport, filing: _Filing) -> None:
    frames = rep.agg.pop("filing", [])
    fl = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=["id", "FilingValidation", "FilingRules", "table"])
    rep.filing = fl[["table", "id", "FilingValidation", "FilingRules"]]
    rep.filing_rules = filing.rules.describe()
    m = fl[fl["table"] == "Measure"]
    rej, warn = int((m["FilingValidation"] == "reject").sum()), int((m["FilingValidation"] == "warning").sum())
    pcr = int(((fl["table"] == "ProgramCost") & (fl["FilingValidation"] == "reject")).sum())
    rep.stats["CET filing validation"] = (f"{filing.rules.label}: {rej:,} measures CEDARS would reject, {warn:,} with warnings only"
                                          + (f"; {pcr:,} ProgramCost rows it would reject" if pcr else ""))


def _check_parts(p: Path, parts: list, dtypes: dict, programcost_csv, acc, allocate_statewide: bool, base_year: int | None,
                 apply_exclude_from_ce: bool, filing: _Filing | None = None) -> PreflightReport:
    """check_inputs for a file split into `parts` ((file, first row, rows) each): each part is checked on its own, its Measure row numbers shifted to the whole
    file's, and the reports combined; duplicate CEInputIDs and the ProgramCost alignment are checked on the whole file."""
    from .pipeline import cost_tests
    rep = PreflightReport()
    sep = delimited.detect_sep(p)
    if sep != "|":
        rep.add("info", "not_pipe_delimited", f"Measure.csv is {delimited.describe(sep)}; it is read as such (the CET itself needs pipe-delimited files)")
    if "CEInputID" in delimited.header_columns(p):
        ids = delimited.read_cet_csv(p, dtype=str, keep_default_na=False, na_values=[""], usecols=lambda c: c.strip() == "CEInputID")["CEInputID"]
        dupmask = ids.duplicated(keep=False) & ids.notna()
        dup = ids[ids.duplicated()]
        if len(dup):
            rep.add("error", "duplicate_id", "CEInputID must be unique (results are keyed on it)", int(dup.nunique()), dup.unique(), _cells("Measure", dupmask, "CEInputID"))
        del ids, dupmask, dup
    program = None
    if programcost_csv is not None:
        try:
            program = cost_tests.load_program_costs(Path(programcost_csv), allocate_statewide=allocate_statewide)
        except Exception as exc:
            rep.add("error", "program_costs_unreadable", f"could not read ProgramCost.csv: {exc}")
    from . import progress
    reports = []
    for k, (f, start, _n) in enumerate(parts):
        progress.report("preflight", k, len(parts))
        r = _check(Path(f), programcost_csv, acc=acc, allocate_statewide=allocate_statewide, base_year=base_year, apply_exclude_from_ce=apply_exclude_from_ce,
                   part=_Part(dtypes, None if program is None else set(program["PrgID"])), filing=filing)
        for i in r.issues:
            i.cells = [(t, row + start if t == "Measure" else row, c) for t, row, c in i.cells]
        r.agg["no_cost_cells"] = [(t, row + start, c) for t, row, c in r.agg.get("no_cost_cells", [])]
        reports.append(r)
    merged: dict[tuple, Issue] = {}
    for r in reports:
        for i in r.issues:
            k = (i.severity, i.code, i.message)
            if k not in merged:
                merged[k] = Issue(i.severity, i.code, i.message, i.count, list(i.examples), list(i.cells), dict(i.impact))
            else:
                m = merged[k]
                m.count += i.count
                m.examples = list(dict.fromkeys([*m.examples, *i.examples]))[:5]
                m.cells = (m.cells + i.cells)[:MAX_CELLS]
                for q, v in i.impact.items():
                    m.impact[q] = m.impact.get(q, 0.0) + v
    rep.issues += list(merged.values())
    rep.stats.update(_merge_stats(reports, base_year))
    if filing is not None:
        rep.agg["filing"] = [f for r in reports for f in r.agg.get("filing", [])]
        years = [int(str(r.stats["claim years"]).split("-")[0]) for r in reports if "claim years" in r.stats]
        _filing_programcost(rep, programcost_csv, filing, min(years) if years else None)
    if programcost_csv is None:
        rep.add("info", "no_program_costs", "no ProgramCost.csv: NoAdmin and regular cost tests will be equal and WeightedProgramCost is 0")
    elif program is not None:
        pc = delimited.read_cet_csv(programcost_csv)
        swmask = (pc["PA"].astype(str).str.upper() == "SW") if "PA" in pc.columns else pd.Series(False, index=pc.index)
        sw_rows = int(swmask.sum())
        if sw_rows and not allocate_statewide:
            rep.add("info", "statewide_program_costs", "ProgramCost rows with PA = SW are given PA = SCE (what CET does with them); tick 'allocate statewide programs' to split them across the IOUs", sw_rows,
                    cells=_cells("ProgramCost", swmask, "PA"))
        m_ids = set().union(*(r.agg.get("measure_prg_ids", set()) for r in reports))
        p_ids = set(program["PrgID"])
        gap = {"program_costs_without_measures": sorted(p_ids - m_ids), "measure_programs_without_costs": sorted(m_ids - p_ids)}
        if gap["measure_programs_without_costs"] or gap["program_costs_without_measures"]:
            no_meas = set(gap["program_costs_without_measures"]) | {i.partition("_")[2] for i in gap["program_costs_without_measures"] if i.partition("_")[0] in VALID_PA}
            where = sorted(set(c for r in reports for c in r.agg["no_cost_cells"]))
            if "PrgID" in pc.columns:
                where += _cells("ProgramCost", pc["PrgID"].astype(str).str.strip().isin(no_meas), "PrgID")
            rep.add("error", "program_alignment", f"PrgIDs do not line up: {gap}", cells=where)
    return rep


def _merge_stats(reports: list[PreflightReport], base_year: int | None) -> dict:
    """The stats of the whole file from those of its parts, in the order a part lists them: counts add up, ranges widen, the program count is of distinct PrgIDs."""
    def as_int(v):
        try:
            return int(str(v).replace(",", ""))
        except ValueError:
            return None
    keys: list[str] = []
    for r in reports:                               # every part's keys in its own order: a key a part adds goes after the key it follows there
        prev = None
        for k in r.stats:
            if k not in keys:
                keys.insert(keys.index(prev) + 1 if prev is not None else 0, k)
            prev = k
    out = {}
    for k in keys:
        vals = [r.stats[k] for r in reports if k in r.stats]
        if k == "claim years":
            out[k] = f"{min(int(str(v).split('-')[0]) for v in vals)}-{max(int(str(v).split('-')[1]) for v in vals)}"
        elif k == "measures by ACC version":
            counts: dict[str, int] = {}
            for v in vals:
                for item in str(v).split(", "):
                    name, _, n = item.rpartition(": ")
                    counts[name] = counts.get(name, 0) + int(n.replace(",", ""))
            out[k] = ", ".join(f"{name}: {n:,}" for name, n in sorted(counts.items()))
        elif k == "First Year (base year) that will be used":
            out[k] = base_year or min(int(str(r.stats["claim years"]).split("-")[0]) for r in reports if "claim years" in r.stats)
        elif k == "programs":
            allp = set().union(*(r.agg["programs"][0] for r in reports if "programs" in r.agg))
            miss = set().union(*(r.agg["programs"][1] for r in reports if "programs" in r.agg))
            out[k] = f"{len(allp):,} ({len(miss):,} not in the program metadata)"
        elif all(as_int(v) is not None for v in vals):
            out[k] = f"{sum(as_int(v) for v in vals):,}"
        else:
            out[k] = vals[0]
    return out


def _check(p: Path, programcost_csv, *, acc=None, allocate_statewide: bool = False, base_year: int | None = None, apply_exclude_from_ce: bool = False,
           part: _Part | None = None, filing: _Filing | None = None) -> PreflightReport:
    """The check of one file. `part` = it is one part of a larger file (see _Part): the checks across the whole file are left to _check_parts."""
    rep = PreflightReport()
    p = Path(p)
    if not p.exists():
        rep.add("error", "file_missing", f"Measure file not found: {p}")
        return rep
    sep = delimited.detect_sep(p)
    if sep is None:
        rep.add("error", "no_delimiter", "The header has no '|', ',' or tab -- Measure.csv must be pipe-delimited (CET's format) or comma-delimited")
        return rep
    if sep != "|":
        rep.add("info", "not_pipe_delimited", f"Measure.csv is {delimited.describe(sep)}; it is read as such (the CET itself needs pipe-delimited files)")
    raw = delimited.read_cet_csv(p, dtype=str, keep_default_na=False, low_memory=False)   # the read outputs.read_raw makes: a part is parsed once for both
    df = raw.mask(raw == "")                                                            # blank -> NA (what na_values=[""] gave)
    del raw
    rep.stats["measures"] = f"{len(df):,}"
    missing = [c for c in REQUIRED_MEASURE_COLUMNS if c not in df.columns]
    if missing:
        rep.add("error", "missing_columns", f"Measure.csv lacks required column(s): {missing}")
        return rep
    for col, why in OPTIONAL_MEASURE_COLUMNS.items():
        if col not in df.columns:
            rep.add("warning", "missing_optional_column", f"no {col} column: {why}")
    if filing is not None:                          # CEDARS' own checks, on the file's text as given (filing_rules.py)
        _filing_check(rep, "Measure", df, df["CEInputID"], filing, filing.year)
        if part is None:                            # a file checked in parts: ProgramCost once, in _check_parts
            cy = pd.to_numeric(df["ClaimYearQuarter"].astype("string").str[:4], errors="coerce")
            _filing_programcost(rep, programcost_csv, filing, int(cy.min()) if cy.notna().any() else None)

    dupmask = df["CEInputID"].duplicated(keep=False) & df["CEInputID"].notna()
    dup = df["CEInputID"][df["CEInputID"].duplicated()]
    if df["CEInputID"].isna().any():
        rep.add("error", "blank_id", "blank CEInputID", int(df["CEInputID"].isna().sum()), cells=_cells("Measure", df["CEInputID"].isna(), "CEInputID"))
    if len(dup) and part is None:                   # a file checked in parts is checked for duplicates as a whole (_check_parts)
        rep.add("error", "duplicate_id", "CEInputID must be unique (results are keyed on it)", int(dup.nunique()), dup.unique(), _cells("Measure", dupmask, "CEInputID"))
    bad_cyq = ~df["ClaimYearQuarter"].astype(str).str.match(r"^\d{4}Q[1-4]$")
    if bad_cyq.any():
        rep.add("error", "bad_claim_year_quarter", "ClaimYearQuarter must look like 2028Q1", int(bad_cyq.sum()), df["ClaimYearQuarter"][bad_cyq].unique(),
                _cells("Measure", bad_cyq, "ClaimYearQuarter"))
    bad_pa = ~df["IOU_AC_Territory"].isin(VALID_PA)
    if bad_pa.any():
        rep.add("error", "bad_iou", f"IOU_AC_Territory must be one of {sorted(VALID_PA)} (the run prices each measure in that territory)",
                int(bad_pa.sum()), df["IOU_AC_Territory"][bad_pa].unique(), _cells("Measure", bad_pa, "IOU_AC_Territory"))
    for col in NUMERIC_COLUMNS + [c for c in ("UnitkWhIOUWater1stBaseline", "UnitkWhIOUWater2ndBaseline") if c in df.columns]:
        bad = _nonblank_numeric_problems(df, col)
        if bad.any():
            rep.add("error", "non_numeric", f"{col} has non-numeric values", int(bad.sum()), df[col][bad].unique(), _cells("Measure", bad, col))
    for col in NEEDED_VALUES:
        blank = df[col].isna() | (df[col].astype(str).str.strip() == "")
        if blank.any():
            rep.add("error", "blank_required_value", f"{col} is blank (CET requires a value; the measure's results would come out blank)", int(blank.sum()),
                    df["CEInputID"][blank].tolist(), _cells("Measure", blank, col))
    rul = pd.to_numeric(df["RUL_Yrs"], errors="coerce").fillna(0.0)
    blank2 = (rul > 0) & (df["UnitkWh2ndBaseline"].isna() | (df["UnitkWh2ndBaseline"].astype(str).str.strip() == ""))
    if blank2.any():
        rep.add("error", "blank_second_baseline", "UnitkWh2ndBaseline is blank on early-replacement measures (RUL_Yrs > 0), which are priced on it after the RUL", int(blank2.sum()),
                df["CEInputID"][blank2].tolist(), _cells("Measure", blank2, "UnitkWh2ndBaseline"))
    if rep.errors:
        return rep
    for col, lo, hi, lo_open, hi_open in ([] if filing is not None else CET_RANGES):     # the filing rules check the same ranges, under their IDs
        v = pd.to_numeric(df[col], errors="coerce")
        bad = v.notna() & ((v <= lo) if lo_open else (v < lo)) | v.notna() & ((v >= hi) if hi_open else (v > hi))
        if bad.any():
            rng = f"{'(' if lo_open else '['}{lo}, {hi if hi != float('inf') else 'no upper limit'}{')' if hi_open else ']'}"
            rep.add("warning", "outside_cet_range", f"{col} is outside CET's allowed range {rng}: CET would reject this upload", int(bad.sum()), df[col][bad].unique(),
                    _cells("Measure", bad, col))
    eul = pd.to_numeric(df["EUL_Yrs"], errors="coerce")
    if ((rul > 0) & (rul >= eul)).any():
        rep.add("warning", "rul_not_below_eul", "RUL_Yrs is not below EUL_Yrs; such a measure is treated as single-baseline for its whole life", int(((rul > 0) & (rul >= eul)).sum()),
                cells=_cells("Measure", (rul > 0) & (rul >= eul), "RUL_Yrs"))

    # measure_id (CEInputID, with "<IOU>_" in front after statewide allocation) -> Measure.csv row, for the checks on the loaded measures below
    raw_pos = {str(v).strip(): i for i, v in enumerate(df["CEInputID"])}

    def _raw_row(mid) -> int | None:
        mid = str(mid).strip()
        if mid in raw_pos:
            return raw_pos[mid]
        head, _, rest = mid.partition("_")
        return raw_pos.get(rest) if rest and head in VALID_PA else None

    def mrows(mask) -> list[int]:
        ids = measures["measure_id"].to_numpy()[np.asarray(mask, dtype=bool)] if not isinstance(mask, pd.DataFrame) else mask["measure_id"].to_numpy()
        return [r for r in (_raw_row(m) for m in ids) if r is not None]

    from . import run_poc
    from .pipeline import cost_tests, gas_savings_profiles, real_impact_profiles, system_zones
    try:
        if part is not None:                        # one part's measures never line up with the whole ProgramCost.csv: _check_parts checks the alignment
            measures, program = run_poc.load_job(p, None, allocate_statewide=allocate_statewide, dtypes=part.dtypes)
        else:
            measures, program = run_poc.load_job(p, programcost_csv, allocate_statewide=allocate_statewide)
    except ValueError as exc:
        where: list = []
        try:                                        # point at the PrgIDs that have no partner in the other file
            pc = delimited.read_cet_csv(programcost_csv, dtype=str, keep_default_na=False) if programcost_csv is not None else None
            if pc is not None and "PrgID" in pc.columns:
                mp, cp = df["PrgID"].astype(str).str.strip(), pc["PrgID"].astype(str).str.strip()
                where = _cells("Measure", ~mp.isin(set(cp)), "PrgID") + _cells("ProgramCost", ~cp.isin(set(mp)), "PrgID")
        except Exception:
            pass
        rep.add("error", "program_alignment", str(exc), cells=where)
        return rep
    rep.stats["claim years"] = f"{measures['claim_year'].min()}-{measures['claim_year'].max()}"
    named = measures.loc[measures["acc_version"] != "", "acc_version"]
    if len(named):                                  # the optional ACCVersion column (BRD BR-06.5): each measure may name its own ACC version
        from . import paths, versions
        run_key = acc.key if acc is not None else None
        rep.stats["measures by ACC version"] = ", ".join(f"{k}: {n:,}" for k, n in named.value_counts().sort_index().items())
        for k in sorted(named.unique()):
            rows = mrows(measures["acc_version"] == k)
            if k != run_key and k not in versions.BUILTIN:
                rep.add("error", "bad_acc_version", f"ACCVersion {k!r} is not a built-in ACC version ({', '.join(versions.BUILTIN)}) or this run's version",
                        len(rows), [k], _cells("Measure", rows, "ACCVersion"))
            elif k != run_key and paths.ref_path(versions.BUILTIN[k].electric_key) is None:
                rep.add("error", "acc_version_file_missing", f"the {versions.BUILTIN[k].label} electric file is not in the reference data folder, so the measures "
                        f"that name it cannot be priced", len(rows), [k], _cells("Measure", rows, "ACCVersion"))
    rep.stats["First Year (base year) that will be used"] = base_year or int(measures["claim_year"].min())
    if part is not None:
        prg = measures["prg_id"].astype(str).str.strip()
        rep.agg["measure_prg_ids"] = set(prg)
        if part.program_ids is not None:
            rep.agg["no_cost_cells"] = [("Measure", r, "PrgID") for r in sorted(set(mrows(~prg.isin(part.program_ids))))]
    elif programcost_csv is None:
        rep.add("info", "no_program_costs", "no ProgramCost.csv: NoAdmin and regular cost tests will be equal and WeightedProgramCost is 0")
    else:
        pc = delimited.read_cet_csv(programcost_csv)
        swmask = (pc["PA"].astype(str).str.upper() == "SW") if "PA" in pc.columns else pd.Series(False, index=pc.index)
        sw_rows = int(swmask.sum())
        if sw_rows and not allocate_statewide:
            rep.add("info", "statewide_program_costs", "ProgramCost rows with PA = SW are given PA = SCE (what CET does with them); tick 'allocate statewide programs' to split them across the IOUs", sw_rows,
                    cells=_cells("ProgramCost", swmask, "PA"))
        gap = cost_tests.check_program_alignment(measures, program)
        if gap["measure_programs_without_costs"] or gap["program_costs_without_measures"]:
            def bare(ids):                                       # with or without the "<IOU>_" that statewide allocation adds
                return set(ids) | {i.partition("_")[2] for i in ids if i.partition("_")[0] in VALID_PA}
            no_meas = bare(gap["program_costs_without_measures"])
            no_cost = measures["prg_id"].astype(str).str.strip().isin(gap["measure_programs_without_costs"])
            where = [("Measure", r, "PrgID") for r in sorted(set(mrows(no_cost)))]
            if "PrgID" in pc.columns:
                where += _cells("ProgramCost", pc["PrgID"].astype(str).str.strip().isin(no_meas), "PrgID")
            rep.add("error", "program_alignment", f"PrgIDs do not line up: {gap}", cells=where)

    _check_programs(rep, measures, programcost_csv, allocate_statewide, apply_exclude_from_ce, mrows)

    # --- electric profile coverage
    elec = measures[(measures["shape_scale"] != 0) | (measures["shape_scale_2nd"] != 0) | (measures["annual_net_kw_1st"] != 0)]
    cat = real_impact_profiles._catalog_hourly() if real_impact_profiles.is_available() else None
    user = real_impact_profiles.user_profiles()
    keys = elec.groupby(["shape_key", "PA", "target_sector"]).size().rename("n").reset_index()
    counts = {"user profile": 0, "library": 0, "library (sector-loose match)": 0, "NO PROFILE": 0}
    missing_examples: list = []
    missing_keys: list = []
    loose_examples: list = []
    loose_keys: list = []
    zone_gaps = 0
    for r in keys.itertuples():
        n = int(r.n)
        if user is not None and user.match(r.shape_key, r.PA, r.target_sector) is not None:
            counts["user profile"] += n
            continue
        pid, exact = real_impact_profiles.match_profile_id(r.shape_key, r.PA, r.target_sector) if cat is not None else (None, False)
        if pid is not None:
            counts["library" if exact else "library (sector-loose match)"] += n
            if not exact:
                loose_examples.append(f"{r.shape_key} | {r.PA} | {r.target_sector}")
                loose_keys.append((r.shape_key, r.PA, r.target_sector))
            continue
        counts["NO PROFILE"] += n
        missing_examples.append(f"{r.shape_key} | {r.PA} | {r.target_sector}")
        missing_keys.append((r.shape_key, r.PA, r.target_sector))
    rep.stats["electric-savings measures"] = f"{len(elec):,}"
    for k, v in counts.items():
        if v:
            rep.stats[f"  electric shape: {k}"] = f"{v:,}"
    def with_keys(keys_: list) -> pd.DataFrame:
        return elec[pd.MultiIndex.from_frame(elec[["shape_key", "PA", "target_sector"]]).isin(keys_)] if keys_ else elec.iloc[:0]

    if counts["NO PROFILE"]:
        hit = with_keys(missing_keys)
        rep.add("warning", "no_electric_profile", "electric measures with no 8,760 profile are NOT priced (ElecPricingStatus = UNPRICED); "
                "supply a profile for them or check the AC_Elec_Impact_Profile spelling", counts["NO PROFILE"], missing_examples,
                _cells("Measure", mrows(hit), "AC_Elec_Impact_Profile"), {IMPACT_KWH: float(hit["shape_scale"].sum())})
    if counts["library (sector-loose match)"]:
        # which records, and how much energy rides on a profile matched without its sector (examples: end use | IOU | AC_Elec_Target_Sector given)
        hit = with_keys(loose_keys)
        rep.add("info", "loose_profile_match", "matched the library profile on (end use, IOU) only, because AC_Elec_Target_Sector has no exact match",
                counts["library (sector-loose match)"], loose_examples, _cells("Measure", mrows(hit), "AC_Elec_Target_Sector"),
                {IMPACT_KWH: float(hit["shape_scale"].sum())})
    if cat is None and not counts["user profile"]:
        rep.add("warning", "no_profile_library", "CEDARS impact-profile library not found (reference data folder): measures cannot be matched to 8,760 profiles")

    # --- kWh-per-unit profiles (BRD BR-06.6): the measure's electric kWh comes from the profile, not from UnitkWh1stBaseline / UnitkWh2ndBaseline
    if user is not None and any(e.scaled for e in user.entries):
        zones = [system_zones.priced_zone(k, pa, ts, cz) for k, pa, ts, cz in zip(measures["shape_key"], measures["PA"], measures["target_sector"], measures["climate_zone"])]
        scaled = pd.Series([real_impact_profiles.scaled_profile(k, pa, ts, z) is not None for k, pa, ts, z in
                            zip(measures["shape_key"], measures["PA"], measures["target_sector"], zones)], index=measures.index)
        if scaled.any():
            rep.add("info", "kwh_per_unit_profile", "priced on a kWh-per-unit 8,760 profile: the measure's annual kWh is the profile's (its UnitkWh1stBaseline is "
                    "not used; a 2nd baseline keeps the file's UnitkWh2ndBaseline / UnitkWh1stBaseline ratio)", int(scaled.sum()),
                    measures.loc[scaled, "measure_id"].head(5).tolist(), _cells("Measure", mrows(scaled), "AC_Elec_Impact_Profile"))
            zero_net = measures[scaled & (measures["unit_kwh_1st"] == 0) & ((measures["kwh_water_1st"] != 0) | (measures["kwh_water_2nd"] != 0))]
            if len(zero_net):
                rep.add("warning", "zero_net_profile_water", "water energy on a load-shifting (zero-net) kWh-per-unit profile is priced on that profile, which nets to "
                        "zero, so its WaterEnergyBen is about 0; give water-energy savings on a separate measure", len(zero_net), zero_net["measure_id"].head(5).tolist(),
                        _cells("Measure", mrows(zero_net), "UnitkWhIOUWater1stBaseline"))

    # --- a separate 2nd-baseline shape (AC_Elec_Impact_Profile_2nd_Baseline) on an early replacement
    second = measures[(measures["shape_key_2nd"] != measures["shape_key"]) & (measures["rul_years_exact"] > 0)
                      & (measures["rul_years_exact"] < measures["eul_years_exact"]) & (measures["shape_scale_2nd"] != 0)]
    if len(second):
        rep.stats["measures with a separate 2nd-baseline shape"] = f"{len(second):,}"
        missing2, eg2, where2 = 0, [], []
        for key, g in second.groupby(["shape_key_2nd", "PA", "target_sector"]):
            has_user = user is not None and user.match(*key) is not None
            has_lib = cat is not None and real_impact_profiles.match_profile_id(*key)[0] is not None
            if not (has_user or has_lib):
                missing2 += len(g)
                eg2.append(" | ".join(map(str, key)))
                where2 += _cells("Measure", mrows(g), "AC_Elec_Impact_Profile_2nd_Baseline")
        if missing2:
            rep.add("warning", "no_electric_profile_2nd", "measures whose AC_Elec_Impact_Profile_2nd_Baseline has no 8,760 profile: their 2nd-baseline years are NOT priced "
                    "(ElecPricingStatus = PARTLY PRICED)", missing2, eg2, where2)

    # --- water energy (priced on the measure's electric profile: a water measure with no electric profile cannot be priced)
    water = measures[(measures["kwh_water_1st"] != 0) | (measures["kwh_water_2nd"] != 0)]
    rep.stats["water-energy measures"] = f"{len(water):,}"
    if len(water):
        no_prof = 0
        eg = []
        where: list = []
        for key, g in water.groupby(["shape_key_water", "PA", "target_sector"]):          # AC_Elec_Impact_Profile_Water, or the measure's own shape
            has_user = user is not None and user.match(*key) is not None
            has_lib = cat is not None and real_impact_profiles.match_profile_id(*key)[0] is not None
            if not (has_user or has_lib):
                no_prof += len(g)
                eg.append(" | ".join(map(str, key)))
                own = g["shape_key_water"] != g["shape_key"]
                where += _cells("Measure", mrows(g[own]), "AC_Elec_Impact_Profile_Water") + _cells("Measure", mrows(g[~own]), "AC_Elec_Impact_Profile")
        if no_prof:
            rep.add("warning", "no_water_profile", "water-energy measures with no electric profile (the water energy is priced on the measure's electric profile) are NOT priced "
                    "(WaterPricingStatus = UNPRICED)", no_prof, eg, where)

    # --- gas profile coverage
    gas = measures[(measures["annual_net_therms_1st"] != 0) | (measures["annual_net_therms_2nd"] != 0)]
    rep.stats["gas-savings measures"] = f"{len(gas):,}"
    no_gas = gas[gas["gas_profile_key"].isna() | ~gas["gas_profile_key"].map(lambda k: not pd.isna(k) and gas_savings_profiles.get_shape(str(k)) is not None)]
    if len(no_gas):
        rep.add("warning", "no_gas_profile", "gas measures whose AC_Gas_Impact_Profile is not a known monthly profile are NOT priced (GasPricingStatus = UNPRICED); "
                "supply the profile or fix the name", len(no_gas), no_gas["gas_profile_key"].fillna("(blank)").unique(), _cells("Measure", mrows(no_gas), "AC_Gas_Impact_Profile"),
                {IMPACT_THERMS: float(no_gas["annual_net_therms_1st"].sum())})
    bad_gs = gas[~gas["gas_sector"].astype(str).str.upper().isin(VALID_GAS_SECTORS)]
    if len(bad_gs):
        rep.add("error", "bad_gas_sector", f"AC_Gas_Target_Sector must be one of Residential/Commercial/Core for measures with gas savings", len(bad_gs), bad_gs["gas_sector"].unique(),
                _cells("Measure", mrows(bad_gs), "AC_Gas_Target_Sector"))

    # --- ACC horizon and run-parameter version
    if acc is not None:
        from .pipeline import discount_settings, gas_avoided_cost, real_avoided_cost_v2
        if acc.settings_acc not in discount_settings.known_versions():
            rep.add("error", "unknown_settings_version", f"no built-in run parameters / E3 tables for the {acc.settings_acc} ACC "
                    f"(known: {discount_settings.known_versions()})")
        try:
            import pyarrow.parquet as pq
            yrs = pq.read_table(str(real_avoided_cost_v2._pq()), columns=["year"]).column("year").to_numpy()
            e0, e1 = int(yrs.min()), int(yrs.max())
            g0, g1 = gas_avoided_cost.year_range()
            last = (measures["claim_year"] + np.ceil(measures["eul_years_exact"])).astype(int)
            rep.stats["ACC years"] = f"electric {e0}-{e1}, gas {g0}-{g1}"
            early_m, long_m = (measures["claim_year"] < e0).to_numpy(), (last > e1 + 1).to_numpy()
            early, long = int(early_m.sum()), int(long_m.sum())
            if early:
                rep.add("warning", "before_acc_horizon", f"claim year is before the first ACC year ({e0}): those years are priced with the {e0} values", early,
                        cells=_cells("Measure", mrows(early_m), "ClaimYearQuarter"))
            if long:
                rep.add("info", "beyond_acc_horizon", f"measure life runs past the last ACC year ({e1}): later years are priced with the {e1} values", long,
                        cells=_cells("Measure", mrows(long_m), "EUL_Yrs"))
        except Exception as exc:                                    # ACC unreadable: the run itself will report it
            rep.add("warning", "acc_unreadable", f"could not read the ACC year range: {exc}")
    return rep
